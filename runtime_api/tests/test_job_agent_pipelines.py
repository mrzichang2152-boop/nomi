import os
import sys
import base64
import json
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("DATABASE_URL", "postgresql://test")
os.environ.setdefault("REDIS_URL", "redis://test")
os.environ.setdefault("APP_PASSWORD", "secret")


class Cursor:
    rowcount = 1

    def __init__(self, rows=None):
        self.rows = rows or []

    def fetchone(self):
        return self.rows[0] if self.rows else None

    def fetchall(self):
        return self.rows


def install_fake_db(monkeypatch, main, handler):
    executed = []

    class Conn:
        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return None

        def execute(self, sql, params=()):
            normalized = " ".join(sql.split())
            executed.append((normalized, params))
            return handler(normalized, params)

    monkeypatch.setattr(main, "db", lambda: Conn())
    return executed


def sample_job():
    return {
        "job_id": "job_pm_ai_1",
        "source": "linkedin_browser",
        "title": "AI Product Manager",
        "company": "Example AI",
        "location": "Shanghai",
        "url": "https://www.linkedin.com/jobs/view/job_pm_ai_1",
        "jd_text": (
            "We need an AI Product Manager with LLM product experience, "
            "workflow automation, data analysis, and cross-functional collaboration. "
            "Experience shipping B2B SaaS products is preferred."
        ),
    }


def sample_resume():
    return {
        "resume_id": "resume_base_1",
        "profile_name": "Zhang",
        "headline": "Product manager building AI workflow products",
        "summary": "Built AI workflow tools for private data, memory, and automation.",
        "skills": ["LLM product", "workflow automation", "data analysis", "B2B SaaS"],
        "experience": [
            {
                "company": "Private Cloud AI",
                "role": "Product Lead",
                "evidence_id": "resume_exp_1",
                "summary": "Led an AI assistant project with memory, pipelines, and browser automation.",
            }
        ],
    }


def test_job_request_routes_to_job_discovery_pipeline(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    route = main.route_tool_request("帮我找 AI 产品经理的远程工作", {"connected_adapters": {"composio": ["gmail"]}})

    assert route["route_type"] == "core_pipeline"
    assert route["pipeline"]["id"] == "job_discovery_pipeline"
    assert route["capability"]["id"] == "career.job.discover"
    assert route["execution_guard"]["permission"] == "read_only"


def test_job_recommendation_request_routes_to_recommendation_pipeline(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    route = main.route_tool_request("根据我的简历推荐适合我的岗位")

    assert route["route_type"] == "core_pipeline"
    assert route["pipeline"]["id"] == "job_recommendation_pipeline"
    assert route["capability"]["id"] == "career.job.recommend"
    assert route["execution_guard"]["permission"] == "read_only"


def test_job_opportunity_request_routes_to_recommendation_pipeline(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    route = main.route_tool_request("帮我找找看有没有适合我的工作机会")

    assert route["route_type"] == "core_pipeline"
    assert route["pipeline"]["id"] == "job_recommendation_pipeline"
    assert route["capability"]["id"] == "career.job.recommend"
    assert route["execution_guard"]["permission"] == "read_only"


def test_job_discovery_pipeline_plans_linkedin_jobs_search_when_no_candidates(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我找适合我的 Backend Engineer 工作机会",
        {
            "pipeline_id": "job_discovery_pipeline",
            "career_profile": {
                "career_profile_id": "career_profile_backend",
                "target_roles": ["Backend Engineer"],
                "target_locations": ["Singapore"],
                "skills": ["Python", "PostgreSQL", "distributed systems"],
            },
        },
    )

    assert result["pipeline_id"] == "job_discovery_pipeline"
    assert result["status"] == "ready_for_browser_navigation"
    assert result["risk"]["permission"] == "read_only"
    assert result["output"]["job_opportunities"] == []
    search_request = result["output"]["linkedin_job_search"]
    assert search_request["expected_event_type"] == "linkedin_job_search_results"
    assert search_request["search_terms"] == {"query": "Backend Engineer", "location": "Singapore"}
    command = search_request["command"]
    assert command["action"] == "open_linkedin_job_search"
    assert "Backend+Engineer" in command["url"]
    assert "location=Singapore" in command["url"]
    assert command["expected_event_type"] == "linkedin_job_search_results"
    assert result["provider_calls"][0]["action"] == "browser.open_linkedin_job_search"
    assert result["provider_calls"][0]["external_side_effect"] is False
    assert result["writeback_plan"][0]["operation"] == "append_browser_navigation_plan"


def test_job_recommendation_pipeline_plans_linkedin_jobs_search_when_candidates_missing(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我找找看有没有适合我的工作机会",
        {
            "pipeline_id": "job_recommendation_pipeline",
            "career_profile": {
                "career_profile_id": "career_profile_pm",
                "target_roles": ["AI Product Manager"],
                "target_locations": ["Shanghai"],
                "skills": ["LLM product", "workflow automation"],
            },
        },
    )

    assert result["pipeline_id"] == "job_recommendation_pipeline"
    assert result["status"] == "needs_candidate_jobs"
    assert result["risk"]["permission"] == "read_only"
    search_request = result["output"]["linkedin_job_search"]
    assert search_request["search_terms"] == {"query": "AI Product Manager", "location": "Shanghai"}
    assert search_request["command"]["action"] == "open_linkedin_job_search"
    assert result["provider_calls"][0]["action"] == "browser.open_linkedin_job_search"
    assert result["provider_calls"][0]["status"] == "planned"
    assert result["external_effects"] == []


def test_linkedin_contact_search_request_routes_to_contact_search_pipeline(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    route = main.route_tool_request(
        "帮我在 LinkedIn 搜索 Example AI 的 Backend Engineer recruiter",
        {"connected_adapters": {"browser": ["linkedin_browser"]}},
    )

    assert route["route_type"] == "core_pipeline"
    assert route["pipeline"]["id"] == "linkedin_contact_search_pipeline"
    assert route["capability"]["id"] == "career.linkedin.contact_search"
    assert route["execution_guard"]["permission"] == "read_only"


def test_linkedin_contact_search_pipeline_outputs_read_only_browser_plan(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我找 Example AI 的 Backend Engineer 招聘负责人",
        {
            "pipeline_id": "linkedin_contact_search_pipeline",
            "company": "Example AI",
            "job_title": "Backend Engineer",
            "location": "Singapore",
            "source_event_ids": ["job_evt_1", "resume_evt_1"],
        },
    )

    assert result["pipeline_id"] == "linkedin_contact_search_pipeline"
    assert result["status"] == "ready_for_browser_navigation"
    assert result["risk"]["permission"] == "read_only"
    assert result["external_effects"] == []
    assert result["resolved_slots"] == {
        "company": "Example AI",
        "job_title": "Backend Engineer",
        "location": "Singapore",
    }
    assert [step["name"] for step in result["steps"]] == [
        "解析公司和岗位",
        "生成 LinkedIn people search",
        "打开候选搜索页",
        "采样 recruiter / hiring manager 主页",
    ]
    search_request = result["output"]["linkedin_contact_search"]
    assert search_request["source"] == "linkedin"
    assert search_request["expected_event_type"] == "linkedin_contact_snapshot"
    assert search_request["command"]["action"] == "open_linkedin_contact_search"
    assert "Example+AI+Backend+Engineer+Singapore" in search_request["command"]["url"]
    assert "recruiter" in search_request["keywords"]
    assert "hiring manager" in search_request["keywords"]
    assert result["provider_calls"][0]["action"] == "browser.open_linkedin_contact_search"
    assert result["writeback_plan"][0]["target"] == "task_trace"


def test_job_discovery_pipeline_normalizes_public_ats_pages_without_external_actions(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "整理这些招聘页里的岗位",
        {
            "pipeline_id": "job_discovery_pipeline",
            "job_pages": [
                {
                    "url": "https://boards.greenhouse.io/exampleai/jobs/123",
                    "title": "AI Product Manager - Example AI",
                    "text": (
                            "AI Product Manager\n"
                            "Example AI\n"
                            "Location: Shanghai\n"
                            "We need LLM product experience, workflow automation, data analysis."
                    ),
                    "source_event_id": "ats_evt_1",
                },
                {
                    "url": "https://jobs.lever.co/acme/abc",
                    "title": "Growth Product Manager",
                    "text": (
                            "Growth Product Manager\n"
                            "Acme\n"
                            "Remote\n"
                            "Experience with B2B SaaS and cross-functional collaboration preferred."
                    ),
                    "source_event_id": "ats_evt_2",
                },
            ],
        },
    )

    assert result["pipeline_id"] == "job_discovery_pipeline"
    assert result["status"] == "completed_read_only"
    assert result["external_effects"] == []
    jobs = result["output"]["job_opportunities"]
    assert len(jobs) == 2
    assert jobs[0]["source"] == "greenhouse_public"
    assert jobs[0]["title"] == "AI Product Manager"
    assert jobs[0]["company"] == "Example AI"
    assert jobs[0]["location"] == "Shanghai"
    assert "LLM product" in jobs[0]["requirements"]
    assert jobs[1]["source"] == "lever_public"
    assert jobs[1]["title"] == "Growth Product Manager"
    assert jobs[1]["company"] == "Acme"
    assert jobs[1]["location"] == "Remote"
    assert "B2B SaaS" in jobs[1]["requirements"]
    assert result["writeback_plan"][0]["payload"]["source_event_ids"] == ["ats_evt_1"]
    assert result["writeback_plan"][1]["payload"]["source_event_ids"] == ["ats_evt_2"]


def test_job_discovery_pipeline_preserves_linkedin_structured_location(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "整理当前 LinkedIn 页面里的岗位机会",
        {
            "pipeline_id": "job_discovery_pipeline",
            "job_pages": [
                {
                    "job_id": "linkedin_search_c1b707edb411",
                    "source": "linkedin_browser_observation",
                    "title": "交易产品经理（跟单交易 & 量化策略）",
                    "company": "Confidential",
                    "location": "Shenzhen, Guangdong, China (Remote)",
                    "url": "https://www.linkedin.com/jobs/search/?currentJobId=4431606283",
                    "text": (
                        "交易产品经理（跟单交易 & 量化策略）\n"
                        "Confidential\n"
                        "Shenzhen, Guangdong, China (Remote)\n"
                        "About the job\n"
                        "推动经典机器人策略优化、AI 策略探索、策略表现与归因展示等功能落地"
                    ),
                    "source_event_ids": ["linkedin_evt_1"],
                }
            ],
        },
    )

    job = result["output"]["job_opportunities"][0]
    assert job["source"] == "linkedin_browser_observation"
    assert job["title"] == "交易产品经理（跟单交易 & 量化策略）"
    assert job["company"] == "Confidential"
    assert job["location"] == "Shenzhen, Guangdong, China (Remote)"
    assert job["source_event_ids"] == ["linkedin_evt_1"]


def test_job_fit_scoring_pipeline_outputs_grounded_score_and_gaps(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我判断这个岗位和我的简历匹配吗",
        {
            "pipeline_id": "job_fit_scoring_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "career_profile": {"target_roles": ["AI Product Manager"], "target_locations": ["Shanghai", "remote"]},
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )

    assert result["pipeline_id"] == "job_fit_scoring_pipeline"
    assert result["status"] == "completed_read_only"
    assert result["risk"]["permission"] == "read_only"
    assert result["resolved_slots"]["job_id"] == "job_pm_ai_1"
    assert result["resolved_slots"]["resume_id"] == "resume_base_1"

    output = result["output"]
    assert 0.0 <= output["fit_score"] <= 1.0
    assert output["fit_score"] >= 0.65
    assert "LLM product" in output["matched_requirements"]
    assert "workflow automation" in output["matched_requirements"]
    assert output["unsupported_claims"] == []
    assert set(output["evidence_ids"]) >= {"jd_evt_1", "resume_evt_1", "resume_exp_1"}


def test_job_recommendation_pipeline_pushes_ranked_jobs_to_dialog_card(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    weak_job = {
        "job_id": "job_growth_pm_low",
        "source": "linkedin_browser_observation",
        "title": "Growth Product Manager",
        "company": "Ad Funnel Co",
        "location": "Remote",
        "url": "https://www.linkedin.com/jobs/view/job_growth_pm_low",
        "jd_text": "Own paid acquisition, SEO, lifecycle campaigns, and demand generation.",
    }
    strong_job = {
        "job_id": "job_ai_workflow_pm_high",
        "source": "linkedin_browser_observation",
        "title": "Senior AI Workflow Product Manager",
        "company": "LinkedIn Observed AI Team",
        "location": "Shanghai / Remote",
        "url": "https://www.linkedin.com/jobs/view/4431606283",
        "jd_text": (
            "We need a Senior AI Workflow Product Manager with LLM product experience, "
            "workflow automation, data analysis, B2B SaaS shipping experience, and "
            "cross-functional collaboration with engineering and design."
        ),
    }

    result = main.run_core_pipeline(
        "根据我的简历筛选适合我的岗位，并把推荐结果推送给我",
        {
            "pipeline_id": "job_recommendation_pipeline",
            "jobs": [weak_job, strong_job],
            "resume": sample_resume(),
            "career_profile": {
                "career_profile_id": "career_profile_default",
                "target_roles": ["AI Product Manager"],
                "target_locations": ["Shanghai", "Remote"],
                "skills": ["LLM product", "workflow automation", "data analysis", "B2B SaaS"],
            },
            "source_event_ids": ["resume_evt_1", "jd_evt_high", "jd_evt_low"],
        },
    )

    assert result["pipeline_id"] == "job_recommendation_pipeline"
    assert result["status"] == "completed_read_only"
    assert result["risk"]["permission"] == "read_only"
    assert result["external_effects"] == []

    recommendations = result["output"]["job_recommendations"]
    assert [item["job_id"] for item in recommendations] == ["job_ai_workflow_pm_high"]
    top = recommendations[0]
    assert top["fit_score"] >= 0.75
    assert top["recommendation_tier"] == "strong_fit"
    assert top["title"] == "Senior AI Workflow Product Manager"
    assert top["company"] == "LinkedIn Observed AI Team"
    assert top["url"] == "https://www.linkedin.com/jobs/view/4431606283"
    assert "LLM product" in top["recommendation_reasons"]
    assert "workflow automation" in top["recommendation_reasons"]
    assert "JD 摘要" not in top["summary"]
    assert "paid acquisition" not in top["summary"].lower()

    suggestion_plan = next(item for item in result["writeback_plan"] if item["target"] == "proactive_suggestions")
    suggestion = suggestion_plan["payload"]
    assert suggestion["metadata"]["suggestion_type"] == "career_job_recommendation"
    assert suggestion["metadata"]["recommended_jobs"][0]["job_id"] == "job_ai_workflow_pm_high"
    assert "Senior AI Workflow Product Manager" in suggestion["body"]
    assert "LinkedIn Observed AI Team" in suggestion["body"]
    assert "LLM product" in suggestion["body"]
    assert "https://www.linkedin.com/jobs/view/4431606283" in suggestion["body"]

    realtime = main.suggestion_to_realtime_message(
        {
            "id": "suggestion-job-1",
            "source_event_id": "jd_evt_high",
            "title": suggestion["title"],
            "body": suggestion["body"],
            "priority": suggestion["priority"],
            "metadata": suggestion["metadata"],
            "created_at": "2026-06-23T18:00:00+08:00",
        }
    )
    assert realtime["type"] == "proactive_message"
    assert realtime["open_view"] == "chat"
    assert "https://www.linkedin.com/jobs/view/4431606283" in realtime["body"]
    assert [action["label"] for action in realtime["actions"]] == ["打开岗位链接", "加入求职看板", "改简历", "联系HR", "忽略"]
    assert realtime["actions"][0]["url"] == "https://www.linkedin.com/jobs/view/4431606283"


def test_job_recommendation_pipeline_does_not_push_placeholder_job_links(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    placeholder_job = {
        "job_id": "job_placeholder_ai_pm",
        "source": "manual",
        "title": "Senior AI Workflow Product Manager",
        "company": "Example AI",
        "location": "Shanghai / Remote",
        "url": "https://boards.greenhouse.io/exampleai/jobs/ai-workflow-pm",
        "jd_text": (
            "We need a Senior AI Workflow Product Manager with LLM product experience, "
            "workflow automation, data analysis, B2B SaaS shipping experience, and "
            "cross-functional collaboration with engineering and design."
        ),
    }

    result = main.run_core_pipeline(
        "根据我的简历筛选适合我的岗位，并把推荐结果推送给我",
        {
            "pipeline_id": "job_recommendation_pipeline",
            "jobs": [placeholder_job],
            "resume": sample_resume(),
            "career_profile": {
                "career_profile_id": "career_profile_default",
                "target_roles": ["AI Product Manager"],
                "target_locations": ["Shanghai", "Remote"],
                "skills": ["LLM product", "workflow automation", "data analysis", "B2B SaaS"],
            },
            "source_event_ids": ["resume_evt_1", "placeholder_jd"],
        },
    )

    recommendation = result["output"]["job_recommendations"][0]
    assert recommendation["url"] == ""
    assert recommendation["url_validation_status"] == "placeholder_or_test_url"
    assert [action["label"] for action in recommendation["actions"]] == ["加入求职看板", "改简历", "联系HR", "忽略"]

    suggestion = next(item for item in result["writeback_plan"] if item["target"] == "proactive_suggestions")["payload"]
    assert "https://boards.greenhouse.io/exampleai/jobs/ai-workflow-pm" not in suggestion["body"]
    assert "链接待验证" in suggestion["body"]
    assert [action["label"] for action in suggestion["metadata"]["actions"]] == ["加入求职看板", "改简历", "联系HR", "忽略"]


def test_job_recommendation_writeback_persists_flat_suggestion_metadata(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    suggestion_payload = {
        "title": "发现高匹配岗位",
        "body": "推荐岗位：Senior AI Workflow Product Manager\n链接：https://boards.greenhouse.io/nomilabs/jobs/ai-workflow-pm",
        "priority": 0.82,
        "status": "open",
        "source_event_id": "acceptance_jd_high",
        "metadata": {
            "suggestion_type": "career_job_recommendation",
            "source": "career_job_recommendation_pipeline",
            "dedupe_key": "career_job_recommendation:job_ai_workflow_pm_high",
            "actions": [
                {
                    "id": "open_job_url",
                    "label": "打开岗位链接",
                    "risk": "read_only",
                        "url": "https://boards.greenhouse.io/nomilabs/jobs/ai-workflow-pm",
                }
            ],
        },
    }
    result = {
        "task_trace_id": "trace-job-recommendation",
        "pipeline_id": "job_recommendation_pipeline",
        "writeback_plan": [
            {"target": "proactive_suggestions", "operation": "create", "payload": suggestion_payload}
        ],
    }
    captured_metadata = []
    published_events = []
    monkeypatch.setattr(main, "publish_realtime_message_safely", lambda event, redis_obj=None: published_events.append(event) or True)

    class Conn:
        def execute(self, sql, params=()):
            normalized = " ".join(sql.split())
            if "INSERT INTO proactive_suggestions" in normalized:
                captured_metadata.append(params[6])
            return Cursor()

    summary = main.apply_pipeline_writeback_plan(Conn(), result)

    assert summary["applied"] is True
    assert summary["failed_count"] == 0
    assert captured_metadata
    persisted_metadata = json.loads(captured_metadata[0])
    assert persisted_metadata["suggestion_type"] == "career_job_recommendation"
    assert persisted_metadata["actions"][0]["label"] == "打开岗位链接"
    assert "metadata" not in persisted_metadata
    assert published_events
    assert published_events[0]["type"] == "proactive_message"
    assert published_events[0]["title"] == "发现高匹配岗位"
    assert published_events[0]["actions"][0]["url"] == "https://boards.greenhouse.io/nomilabs/jobs/ai-workflow-pm"


def test_job_recommendation_writeback_skips_existing_dedupe_without_realtime_publish(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    suggestion_payload = {
        "title": "发现高匹配岗位",
        "body": "推荐岗位：Senior AI Workflow Product Manager",
        "priority": 0.82,
        "status": "open",
        "metadata": {
            "suggestion_type": "career_job_recommendation",
            "source": "career_job_recommendation_pipeline",
            "dedupe_key": "career_job_recommendation:job_ai_workflow_pm_high",
        },
    }
    result = {
        "task_trace_id": "trace-job-recommendation",
        "pipeline_id": "job_recommendation_pipeline",
        "writeback_plan": [
            {"target": "proactive_suggestions", "operation": "create", "payload": suggestion_payload}
        ],
    }
    published_events = []
    executed_sql = []
    monkeypatch.setattr(main, "publish_realtime_message_safely", lambda event, redis_obj=None: published_events.append(event) or True)

    class ExistingCursor:
        def fetchone(self):
            return ("existing-suggestion-id",)

    class Cursor:
        def fetchone(self):
            return None

    class Conn:
        def execute(self, sql, params=()):
            normalized = " ".join(sql.split())
            executed_sql.append(normalized)
            if "FROM proactive_suggestions" in normalized and "metadata->>'dedupe_key'" in normalized:
                assert params[0] == "career_job_recommendation:job_ai_workflow_pm_high"
                return ExistingCursor()
            if "INSERT INTO proactive_suggestions" in normalized:
                raise AssertionError("duplicate writeback suggestions must not be inserted")
            return Cursor()

    summary = main.apply_pipeline_writeback_plan(Conn(), result)

    assert summary["applied"] is True
    assert summary["failed_count"] == 0
    assert any("FROM proactive_suggestions" in sql for sql in executed_sql)
    assert published_events == []


def test_cover_letter_pipeline_generates_grounded_draft_without_sending(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "根据这个 JD 和我的简历写一封 cover letter",
        {
            "pipeline_id": "cover_letter_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "recipient": "Maya",
            "channel": "gmail",
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )

    assert result["pipeline_id"] == "cover_letter_pipeline"
    assert result["status"] == "draft_ready"
    assert result["risk"]["permission"] == "draft"
    assert result["external_effects"] == []
    assert result["output"]["draft"]["channel"] == "gmail"
    assert "AI Product Manager" in result["output"]["draft"]["subject"]
    assert "LLM product" in result["output"]["draft"]["body"]
    assert "workflow automation" in result["output"]["draft"]["body"]
    assert result["output"]["unsupported_claims"] == []
    assert result["output"]["send_blocked"] is True


def test_cover_letter_pipeline_uses_short_resume_highlights_without_sensitive_dump(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    resume = sample_resume()
    resume["summary"] = (
        "基本信息: 姓名: 范小刚 联系方式: 15510261379 邮箱: xiaogangfan1228@gmail.com "
        "技术能力: 精通 Java、熟悉 Go。具备 AI Agent 项目经验，负责深度自主研究模块。"
        "工作经历: 2024.04-2026.04 AElf 职位: Leader。"
    )
    result = main.run_core_pipeline(
        "生成一版简短求职信，不要发送",
        {
            "pipeline_id": "cover_letter_pipeline",
            "job": {
                **sample_job(),
                "title": "AI Native 全栈工程师（LLM / Agent / AI应用开发）",
                "company": "QuantGroup",
            },
            "resume": resume,
            "recipient": "Hiring Team",
            "channel": "gmail",
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )

    body = result["output"]["draft"]["body"]
    assert result["status"] == "draft_ready"
    assert "QuantGroup" in body
    assert "AI Native 全栈工程师" in body
    assert "AI Agent" in body
    assert "15510261379" not in body
    assert "xiaogangfan1228@gmail.com" not in body
    assert "基本信息:" not in body
    assert len(body) < 900


def test_job_application_pipeline_blocks_apply_submit_without_delegated_grant(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我自动点击 LinkedIn Apply 并提交",
        {
            "pipeline_id": "job_application_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "application_action": "submit_application",
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )

    assert result["pipeline_id"] == "job_application_pipeline"
    assert result["status"] == "blocked_until_delegated_grant"
    assert result["risk"]["permission"] == "external_execution"
    assert result["risk"]["confirmation_required"] is True
    assert "submit_application" in result["external_effects"]
    assert result["output"]["apply_submit_blocked"] is True
    assert result["output"]["automation_requirements"]["requires_target_manifest"] is True
    assert result["output"]["automation_requirements"]["requires_grounded_content"] is True
    assert result["output"]["suggested_grant"]["platform"] == "linkedin"
    assert result["output"]["suggested_grant"]["daily_limit"] <= 20
    assert result["output"]["suggested_manifest"]["targets"][0]["related_job_id"] == "job_pm_ai_1"


def test_linkedin_outreach_pipeline_creates_confirmation_only_draft(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "帮我在 LinkedIn 给招聘负责人 Maya 发一段自我介绍",
        {
            "pipeline_id": "outreach_message_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "contact": {
                "contact_id": "maya_linkedin",
                "name": "Maya",
                "role": "Recruiter",
                "company": "Example AI",
                "channel": "linkedin",
            },
            "source_event_ids": ["jd_evt_1", "resume_evt_1", "contact_evt_1"],
        },
    )

    assert result["pipeline_id"] == "outreach_message_pipeline"
    assert result["status"] == "draft_ready"
    assert result["risk"]["permission"] == "external_message"
    assert result["risk"]["confirmation_required"] is True
    assert result["external_effects"] == ["send_message"]
    assert result["output"]["draft"]["channel"] == "linkedin"
    assert result["output"]["draft"]["recipient"] == "Maya"
    assert "AI Product Manager" in result["output"]["draft"]["body"]
    assert "发送前需要用户确认" in result["output"]["confirmation_card"]["message"]


def test_career_schema_bootstrap_creates_job_agent_state_tables(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        return Cursor()

    executed = install_fake_db(monkeypatch, main, handler)
    main.ensure_task_routing_schema()
    combined = "\n".join(sql for sql, _ in executed)

    assert "CREATE TABLE IF NOT EXISTS career_profiles" in combined
    assert "CREATE TABLE IF NOT EXISTS career_resumes" in combined
    assert "CREATE TABLE IF NOT EXISTS job_opportunities" in combined
    assert "CREATE TABLE IF NOT EXISTS resume_versions" in combined
    assert "CREATE TABLE IF NOT EXISTS job_applications" in combined
    assert "career_resumes_updated_idx" in combined
    assert "job_opportunities_status_idx" in combined
    assert "job_applications_stage_idx" in combined


def test_career_writeback_materializes_opportunities_resumes_and_applications(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        return Cursor()

    executed = install_fake_db(monkeypatch, main, handler)

    discovery = main.run_core_pipeline(
        "帮我找 AI 产品经理岗位",
        {"pipeline_id": "job_discovery_pipeline", "jobs": [sample_job()], "source_event_ids": ["job_evt_1"]},
    )
    fit = main.run_core_pipeline(
        "这个岗位和我的简历匹配吗",
        {
            "pipeline_id": "job_fit_scoring_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )
    tailored_resume = main.run_core_pipeline(
        "针对这个岗位改简历",
        {
            "pipeline_id": "resume_tailoring_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )
    application = main.run_core_pipeline(
        "帮我自动点击 LinkedIn Apply 并提交",
        {
            "pipeline_id": "job_application_pipeline",
            "job": sample_job(),
            "resume": sample_resume(),
            "application_action": "submit_application",
            "source_event_ids": ["jd_evt_1", "resume_evt_1"],
        },
    )

    for result in [discovery, fit, tailored_resume, application]:
        summary = main.apply_pipeline_writeback_plan(main.db(), result)
        assert summary["failed_count"] == 0

    assert any("INSERT INTO job_opportunities" in sql for sql, _ in executed)
    assert any("ON CONFLICT (id) DO UPDATE" in sql and "fit_score" in sql for sql, _ in executed)
    assert any("INSERT INTO resume_versions" in sql for sql, _ in executed)
    assert any("INSERT INTO job_applications" in sql for sql, _ in executed)

    opportunity_params = next(params for sql, params in executed if "INSERT INTO job_opportunities" in sql)
    assert opportunity_params[0] == "job_pm_ai_1"
    assert opportunity_params[2] == "AI Product Manager"
    assert opportunity_params[3] == "Example AI"

    application_params = next(params for sql, params in executed if "INSERT INTO job_applications" in sql)
    assert application_params[1] == "job_pm_ai_1"
    assert application_params[2] in {"blocked_until_delegated_grant", "ready_for_confirmation", "interested"}


def test_career_board_endpoint_returns_job_agent_state_for_ui(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        if "FROM career_profiles" in sql:
            return Cursor(
                [
                    (
                        "career_profile_default",
                        "AI workflow product manager",
                        ["AI Product Manager"],
                        ["Shanghai"],
                        ["LLM product", "workflow automation"],
                        ["resume_evt_1"],
                        {"profile_name": "Zhang"},
                        "2026-06-08T09:00:00+08:00",
                    )
                ]
            )
        if "FROM job_opportunities" in sql:
            return Cursor(
                [
                    (
                        "job_pm_ai_1",
                        "linkedin_browser",
                        "AI Product Manager",
                        "Example AI",
                        "Shanghai",
                        "https://www.linkedin.com/jobs/view/job_pm_ai_1",
                        "tracked",
                        0.82,
                        ["LLM product", "workflow automation"],
                        ["job_evt_1"],
                        {"matched_requirements": ["LLM product"]},
                        "2026-06-08T09:10:00+08:00",
                        "2026-06-08T09:20:00+08:00",
                    )
                ]
            )
        if "FROM resume_versions" in sql:
            return Cursor(
                [
                    (
                        "resume_version_resume_base_1_job_pm_ai_1",
                        "resume_base_1",
                        "job_pm_ai_1",
                        "draft",
                        ["resume_evt_1"],
                        {"changes": [{"section": "summary", "change": "突出 LLM product"}]},
                        "2026-06-08T09:30:00+08:00",
                        "2026-06-08T09:30:00+08:00",
                    )
                ]
            )
        if "FROM job_applications" in sql:
            return Cursor(
                [
                    (
                        "application_job_pm_ai_1_submit_application",
                        "job_pm_ai_1",
                        "blocked_until_delegated_grant",
                        "apply_submit_blocked",
                        "request_delegated_grant_and_target_manifest",
                        "submit_application",
                        "linkedin",
                        ["jd_evt_1"],
                        {"reason": "missing grant"},
                        "2026-06-08T09:40:00+08:00",
                        "2026-06-08T09:40:00+08:00",
                    )
                ]
            )
        return Cursor()

    install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).get("/api/career/board", headers={"x-par-password": "secret"})

    assert response.status_code == 200
    body = response.json()
    assert body["profiles"][0]["headline"] == "AI workflow product manager"
    assert body["opportunities"][0]["title"] == "AI Product Manager"
    assert body["opportunities"][0]["fit_score"] == 0.82
    assert body["resume_versions"][0]["target_job_id"] == "job_pm_ai_1"
    assert body["applications"][0]["status"] == "blocked_until_delegated_grant"
    assert body["applications"][0]["next_step"] == "request_delegated_grant_and_target_manifest"


def test_career_board_endpoint_filters_manual_demo_opportunities_by_default(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        if "FROM career_profiles" in sql:
            return Cursor([])
        if "FROM job_opportunities" in sql:
            return Cursor(
                [
                    (
                        "job_manual_example",
                        "manual",
                        "AI Product Manager",
                        "Example AI",
                        "Shanghai",
                        "https://example.com/jobs/ai-pm",
                        "tracked",
                        0.82,
                        ["LLM product"],
                        [],
                        {"demo": True, "source_label": "manual_test_data"},
                        "2026-06-08T09:10:00+08:00",
                        "2026-06-08T09:20:00+08:00",
                    ),
                    (
                        "job_real_linkedin",
                        "linkedin_browser",
                        "Backend Engineer",
                        "ByteDance",
                        "Singapore",
                        "https://www.linkedin.com/jobs/view/1234567890",
                        "tracked",
                        0.76,
                        ["Python", "distributed systems"],
                        ["linkedin_job_evt_1"],
                        {"jd_text": "Build backend systems for real-time products."},
                        "2026-06-09T09:10:00+08:00",
                        "2026-06-09T09:20:00+08:00",
                    ),
                    (
                        "job_placeholder_guard_panel_1782285040",
                        "manual",
                        "Senior AI Workflow Product Manager",
                        "Example AI",
                        "Shanghai",
                        "",
                        "tracked",
                        0.76,
                        ["LLM product"],
                        ["placeholder_guard_panel_evt"],
                        {"source_label": "old_guard_panel"},
                        "2026-06-09T10:10:00+08:00",
                        "2026-06-09T10:20:00+08:00",
                    ),
                ]
            )
        if "FROM resume_versions" in sql:
            return Cursor([])
        if "FROM career_resumes" in sql:
            return Cursor([])
        if "FROM job_applications" in sql:
            return Cursor([])
        return Cursor()

    install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).get("/api/career/board", headers={"x-par-password": "secret"})

    assert response.status_code == 200
    body = response.json()
    assert [item["id"] for item in body["opportunities"]] == ["job_real_linkedin"]

    response_with_demo = TestClient(main.app).get(
        "/api/career/board?include_demo=true",
        headers={"x-par-password": "secret"},
    )
    assert response_with_demo.status_code == 200
    assert [item["id"] for item in response_with_demo.json()["opportunities"]] == [
        "job_manual_example",
        "job_real_linkedin",
        "job_placeholder_guard_panel_1782285040",
    ]


def test_career_opportunity_detail_endpoint_builds_grounded_drafts_and_interview_brief(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        if "FROM job_opportunities" in sql:
            return Cursor(
                [
                    (
                        "job_pm_ai_1",
                        "linkedin_browser",
                        "AI Product Manager",
                        "Example AI",
                        "Shanghai",
                        "https://www.linkedin.com/jobs/view/job_pm_ai_1",
                        "tracked",
                        0.82,
                        ["LLM product", "workflow automation", "B2B SaaS"],
                        ["job_evt_1"],
                        {
                            "jd_text": sample_job()["jd_text"],
                            "matched_requirements": ["LLM product", "workflow automation"],
                            "gap_requirements": ["B2B SaaS"],
                            "recruiter_name": "Maya",
                        },
                        "2026-06-08T09:10:00+08:00",
                        "2026-06-08T09:20:00+08:00",
                    )
                ]
            )
        if "FROM career_profiles" in sql:
            return Cursor(
                [
                    (
                        "career_profile_default",
                        "Product manager building AI workflow products",
                        ["AI Product Manager"],
                        ["Shanghai"],
                        ["LLM product", "workflow automation", "data analysis"],
                        ["resume_evt_1"],
                        {"summary": "Built AI workflow tools for memory, pipelines and private data."},
                        "2026-06-08T09:00:00+08:00",
                    )
                ]
            )
        if "FROM resume_versions" in sql:
            return Cursor(
                [
                    (
                        "resume_version_resume_base_1_job_pm_ai_1",
                        "resume_base_1",
                        "job_pm_ai_1",
                        "draft",
                        ["resume_evt_1"],
                        {"changes": [{"section": "summary", "change": "突出 LLM product 与 workflow automation。"}]},
                        "2026-06-08T09:30:00+08:00",
                        "2026-06-08T09:30:00+08:00",
                    )
                ]
            )
        if "FROM job_applications" in sql:
            return Cursor(
                [
                    (
                        "application_job_pm_ai_1_submit_application",
                        "job_pm_ai_1",
                        "blocked_until_delegated_grant",
                        "apply_submit_blocked",
                        "request_delegated_grant_and_target_manifest",
                        "submit_application",
                        "linkedin",
                        ["jd_evt_1"],
                        {"reason": "missing grant"},
                        "2026-06-08T09:40:00+08:00",
                        "2026-06-08T09:40:00+08:00",
                    )
                ]
            )
        return Cursor()

    install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).get(
        "/api/career/opportunities/job_pm_ai_1",
        headers={"x-par-password": "secret"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["opportunity"]["title"] == "AI Product Manager"
    assert body["fit_summary"]["fit_score"] == 0.82
    assert body["fit_summary"]["matched_requirements"] == ["LLM product", "workflow automation"]
    assert body["fit_summary"]["gap_requirements"] == ["B2B SaaS"]
    assert body["resume_draft"]["payload"]["changes"][0]["section"] == "summary"
    assert body["cover_letter_draft"]["send_blocked"] is True
    assert "LLM product" in body["cover_letter_draft"]["draft"]["body"]
    assert "workflow automation" in body["cover_letter_draft"]["draft"]["body"]
    assert body["outreach_draft"]["confirmation_card"]["final_user_confirmation"] is True
    assert body["outreach_draft"]["draft"]["recipient"] == "Maya"
    assert "AI Product Manager" in body["outreach_draft"]["draft"]["body"]
    assert "LLM product" in body["interview_prep"]["prep_brief"]["talking_points"]
    assert body["application_history"][0]["status"] == "blocked_until_delegated_grant"
    assert body["application_history"][0]["next_step"] == "request_delegated_grant_and_target_manifest"


def test_career_ats_preview_endpoint_fetches_public_page_read_only(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        raise AssertionError(f"ATS preview must not write or read local career tables, got SQL: {sql}")

    install_fake_db(monkeypatch, main, handler)

    fetched_urls = []

    def fake_fetch(url):
        fetched_urls.append(url)
        return {
            "url": url,
            "title": "AI Product Manager - Example AI",
            "text": (
                "AI Product Manager\n"
                "Example AI\n"
                "Location: Shanghai\n"
                "We need LLM product experience, workflow automation, data analysis, "
                "and cross-functional collaboration."
            ),
        }

    monkeypatch.setattr(main, "fetch_public_ats_page", fake_fetch)

    response = TestClient(main.app).post(
        "/api/career/ats/preview",
        headers={"x-par-password": "secret"},
        json={"url": "https://boards.greenhouse.io/exampleai/jobs/123"},
    )

    assert response.status_code == 200
    body = response.json()
    assert fetched_urls == ["https://boards.greenhouse.io/exampleai/jobs/123"]
    assert body["status"] == "completed_read_only"
    assert body["source"] == "greenhouse_public"
    assert body["external_effects"] == []
    assert body["writeback_ready"] is True
    assert body["writeback_performed"] is False
    assert body["next_actions"] == ["匹配简历", "生成外联草稿", "加入机会跟踪"]
    job = body["job_opportunities"][0]
    assert job["title"] == "AI Product Manager"
    assert job["company"] == "Example AI"
    assert job["location"] == "Shanghai"
    assert job["source"] == "greenhouse_public"
    assert job["source_event_ids"] == ["ats_preview_https_boards_greenhouse_io_exampleai_jobs_123"]
    assert {"LLM product", "workflow automation", "data analysis"}.issubset(set(job["requirements"]))


def test_career_ats_preview_endpoint_accepts_supplied_html_without_network(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def should_not_fetch(url):
        raise AssertionError("html_text preview should not fetch the network")

    monkeypatch.setattr(main, "fetch_public_ats_page", should_not_fetch)

    response = TestClient(main.app).post(
        "/api/career/ats/preview",
        headers={"x-par-password": "secret"},
        json={
            "url": "https://jobs.lever.co/acme/abc",
            "html_text": """
                <html><head><title>Growth Product Manager - Acme</title></head>
                <body><h1>Growth Product Manager</h1><p>Acme</p><p>Remote</p>
                <section>We need B2B SaaS and cross-functional collaboration.</section></body></html>
            """,
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["source"] == "lever_public"
    assert body["external_effects"] == []
    job = body["job_opportunities"][0]
    assert job["title"] == "Growth Product Manager"
    assert job["company"] == "Acme"
    assert job["location"] == "Remote"
    assert "B2B SaaS" in job["requirements"]
    assert "cross-functional collaboration" in job["requirements"]


def test_career_ats_preview_endpoint_rejects_unsupported_urls(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    response = TestClient(main.app).post(
        "/api/career/ats/preview",
        headers={"x-par-password": "secret"},
        json={"url": "https://example.com/jobs/123", "text": "AI Product Manager\nExample\nRemote"},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "unsupported public ATS URL"


def test_career_ats_list_preview_reads_greenhouse_board_jobs_read_only(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        raise AssertionError(f"ATS list preview must not touch local career tables, got SQL: {sql}")

    install_fake_db(monkeypatch, main, handler)
    fetched_urls = []

    def fake_fetch_json(url):
        fetched_urls.append(url)
        return {
            "jobs": [
                {
                    "id": 123,
                    "title": "AI Product Manager",
                    "location": {"name": "Shanghai"},
                    "absolute_url": "https://boards.greenhouse.io/exampleai/jobs/123",
                    "content": "Need LLM product, workflow automation, and data analysis.",
                },
                {
                    "id": 456,
                    "title": "Growth PM",
                    "location": {"name": "Remote"},
                    "absolute_url": "https://boards.greenhouse.io/exampleai/jobs/456",
                    "content": "Need B2B SaaS and cross-functional collaboration.",
                },
            ]
        }

    monkeypatch.setattr(main, "fetch_public_ats_json", fake_fetch_json)

    response = TestClient(main.app).post(
        "/api/career/ats/list-preview",
        headers={"x-par-password": "secret"},
        json={"url": "https://boards.greenhouse.io/exampleai"},
    )

    assert response.status_code == 200
    body = response.json()
    assert fetched_urls == ["https://boards-api.greenhouse.io/v1/boards/exampleai/jobs?content=true"]
    assert body["source"] == "greenhouse_public"
    assert body["status"] == "completed_read_only"
    assert body["writeback_performed"] is False
    assert body["external_effects"] == []
    assert len(body["job_opportunities"]) == 2
    assert body["job_opportunities"][0]["title"] == "AI Product Manager"
    assert body["job_opportunities"][0]["company"] == "Exampleai"
    assert body["job_opportunities"][0]["location"] == "Shanghai"
    assert "LLM product" in body["job_opportunities"][0]["requirements"]
    assert body["job_opportunities"][1]["title"] == "Growth PM"
    assert "B2B SaaS" in body["job_opportunities"][1]["requirements"]
    assert body["job_opportunities"][0]["source_event_ids"] == ["ats_list_greenhouse_public_123"]


def test_career_ats_list_preview_reads_workable_board_jobs_read_only(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        raise AssertionError(f"ATS list preview must not touch local career tables, got SQL: {sql}")

    install_fake_db(monkeypatch, main, handler)
    fetched_urls = []

    def fake_fetch_json(url):
        fetched_urls.append(url)
        return {
            "jobs": [
                {
                    "shortcode": "AI-PM-1",
                    "title": "AI Product Manager",
                    "location": {"location_str": "Remote"},
                    "url": "https://apply.workable.com/exampleai/j/AI-PM-1/",
                    "description": "Need LLM product, workflow automation, and data analysis.",
                }
            ]
        }

    monkeypatch.setattr(main, "fetch_public_ats_json", fake_fetch_json)

    response = TestClient(main.app).post(
        "/api/career/ats/list-preview",
        headers={"x-par-password": "secret"},
        json={"url": "https://apply.workable.com/exampleai/"},
    )

    assert response.status_code == 200
    body = response.json()
    assert fetched_urls == ["https://www.workable.com/api/accounts/exampleai?details=true"]
    assert body["source"] == "workable_public"
    assert body["writeback_performed"] is False
    assert body["job_opportunities"][0]["title"] == "AI Product Manager"
    assert body["job_opportunities"][0]["location"] == "Remote"
    assert "workflow automation" in body["job_opportunities"][0]["requirements"]
    assert body["job_opportunities"][0]["source_event_ids"] == ["ats_list_workable_public_AI-PM-1"]


def test_career_ats_list_preview_reads_smartrecruiters_board_jobs_read_only(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        raise AssertionError(f"ATS list preview must not touch local career tables, got SQL: {sql}")

    install_fake_db(monkeypatch, main, handler)
    fetched_urls = []

    def fake_fetch_json(url):
        fetched_urls.append(url)
        return {
            "content": [
                {
                    "id": "sr-123",
                    "name": "Growth PM",
                    "location": {"city": "Shanghai", "country": "China"},
                    "ref": "https://jobs.smartrecruiters.com/ExampleAI/sr-123-growth-pm",
                    "jobAd": {"sections": {"jobDescription": {"text": "Need B2B SaaS and cross-functional collaboration."}}},
                }
            ]
        }

    monkeypatch.setattr(main, "fetch_public_ats_json", fake_fetch_json)

    response = TestClient(main.app).post(
        "/api/career/ats/list-preview",
        headers={"x-par-password": "secret"},
        json={"url": "https://careers.smartrecruiters.com/ExampleAI"},
    )

    assert response.status_code == 200
    body = response.json()
    assert fetched_urls == ["https://api.smartrecruiters.com/v1/companies/ExampleAI/postings?limit=100"]
    assert body["source"] == "smartrecruiters_public"
    assert body["writeback_performed"] is False
    assert body["job_opportunities"][0]["title"] == "Growth PM"
    assert body["job_opportunities"][0]["location"] == "Shanghai, China"
    assert "B2B SaaS" in body["job_opportunities"][0]["requirements"]
    assert body["job_opportunities"][0]["source_event_ids"] == ["ats_list_smartrecruiters_public_sr-123"]


def test_career_profile_ingest_endpoint_builds_and_persists_profile_from_resume_text(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        return Cursor()

    executed = install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).post(
        "/api/career/profile/ingest",
        headers={"x-par-password": "secret"},
        json={
            "resume_text": (
                "AI Product Manager\n"
                "Built AI workflow tools for private data, memory, LLM product experiences, "
                "workflow automation and data analysis."
            ),
            "target_roles": ["AI Product Manager"],
            "target_locations": ["Shanghai", "Remote"],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed_read_only"
    assert body["writeback_performed"] is True
    assert body["writeback"]["failed_count"] == 0
    profile = body["career_profile"]
    assert profile["career_profile_id"] == "career_profile_default"
    assert profile["target_roles"] == ["AI Product Manager"]
    assert profile["target_locations"] == ["Shanghai", "Remote"]
    assert "LLM product" in profile["skills"]
    assert "workflow automation" in profile["skills"]
    assert "data analysis" in profile["skills"]
    assert profile["evidence_ids"] == ["career_resume_text_default"]
    insert_params = next(params for sql, params in executed if "INSERT INTO career_profiles" in sql)
    assert insert_params[0] == "career_profile_default"
    assert "AI Product Manager" in insert_params[1]
    assert insert_params[2] == ["AI Product Manager"]
    assert insert_params[3] == ["Shanghai", "Remote"]


def test_career_profile_pipeline_normalizes_target_role_from_resume_headline(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "整理我的求职画像",
        {
            "pipeline_id": "career_profile_pipeline",
            "resume": sample_resume(),
            "source_event_ids": ["resume_evt_1"],
        },
    )

    profile = result["output"]["career_profile"]
    assert result["pipeline_id"] == "career_profile_pipeline"
    assert result["status"] == "completed_read_only"
    assert profile["target_roles"] == ["Product Manager"]
    assert "Product manager building AI workflow products" not in profile["target_roles"]
    assert "LLM product" in profile["skills"]
    assert "workflow automation" in profile["skills"]
    assert "resume_evt_1" in profile["evidence_ids"]


def test_career_resume_file_import_parses_docx_and_tracks_base_resume(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    doc = Document()
    doc.add_heading("AI Product Manager", level=1)
    doc.add_paragraph("Built LLM product experiences, workflow automation, and data analysis tools.")
    buffer = BytesIO()
    doc.save(buffer)

    def handler(sql, params):
        return Cursor()

    executed = install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).post(
        "/api/career/resumes/import",
        headers={"x-par-password": "secret"},
        json={
            "filename": "resume.docx",
            "content_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
            "target_roles": ["AI Product Manager"],
            "target_locations": ["Shanghai"],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed_read_only"
    assert body["file_type"] == "docx"
    assert body["base_resume"]["filename"] == "resume.docx"
    assert "AI Product Manager" in body["parsed_text"]
    assert "LLM product" in body["career_profile"]["skills"]
    assert "workflow automation" in body["career_profile"]["skills"]
    assert body["writeback"]["failed_count"] == 0
    assert any("INSERT INTO career_resumes" in sql for sql, _ in executed)
    assert any("INSERT INTO career_profiles" in sql for sql, _ in executed)
    resume_params = next(params for sql, params in executed if "INSERT INTO career_resumes" in sql)
    assert resume_params[1] == "resume.docx"
    assert resume_params[2] == "docx"


def test_career_profile_ingest_extracts_real_backend_resume_skills_and_ignores_pdf_noise(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        return Cursor()

    install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).post(
        "/api/career/profile/ingest",
        headers={"x-par-password": "secret"},
        json={
            "resume_text": (
                "~ ~\n"
                "g\n"
                "M 7600c6fe9bae0b3e1HF82dS4F1RUwI-_UP-bWOWim_PSMg~S~M\n"
                "基本信息:\n"
                "姓 名 ：范小刚 性 别 ：男 工作年限 ：13 年\n"
                "技术能力:\n"
                "精通 Java、熟悉 Go。熟悉 JVM 原理、内存模型与垃圾回收机制，具备 JVM 调优经验。\n"
                "具备高并发、高可用架构实战经验，主导过 10 万+tps 系统建设与优化。\n"
                "熟悉 Kafka、RocketMQ、Redis、Elasticsearch、MySQL、分库分表、DDD 与微服务架构。\n"
                "有 AI Agent 项目经验，熟练 Claude Code、Cursor，理解 harness 工程。\n"
                "2020.04-2024.04 京东 职位：Leader &&架构\n"
            ),
            "target_roles": ["Java 后端架构师"],
            "target_locations": ["北京", "远程"],
        },
    )

    assert response.status_code == 200
    body = response.json()
    profile = body["career_profile"]
    assert profile["profile_name"] == "范小刚"
    assert profile["headline"] == "Java 后端架构 / 技术 Leader"
    assert "Java" in profile["skills"]
    assert "Go" in profile["skills"]
    assert "JVM tuning" in profile["skills"]
    assert "high concurrency" in profile["skills"]
    assert "Kafka" in profile["skills"]
    assert "RocketMQ" in profile["skills"]
    assert "Redis" in profile["skills"]
    assert "Elasticsearch" in profile["skills"]
    assert "MySQL" in profile["skills"]
    assert "DDD" in profile["skills"]
    assert "AI Agent" in profile["skills"]
    assert all("7600c6fe" not in skill for skill in profile["skills"])


def test_backend_job_fit_scoring_uses_real_backend_resume_requirements(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    result = main.run_core_pipeline(
        "评估这个 Java 后端架构岗位匹配度",
        {
            "pipeline_id": "job_fit_scoring_pipeline",
            "job": {
                "job_id": "linkedin_job_backend_architect_1",
                "source": "linkedin_browser_observation",
                "title": "Senior Java Backend Architect",
                "company": "Example Fintech",
                "location": "Beijing, China",
                "jd_text": (
                    "We need a senior backend architect with Java, JVM tuning, high concurrency, "
                    "distributed systems, Kafka, Redis, Elasticsearch, MySQL, microservices, DDD, "
                    "and team leadership experience."
                ),
            },
            "resume": {
                "resume_id": "resume_fan_xiaogang",
                "headline": "Java 后端架构 / 技术 Leader",
                "summary": (
                    "13 年 Java 后端与架构经验，熟悉 Go、JVM 调优、高并发高可用、Kafka、RocketMQ、"
                    "Redis、Elasticsearch、MySQL、分库分表、DDD、微服务和技术团队管理。"
                ),
                "skills": ["Java", "Go", "JVM tuning", "high concurrency", "Kafka", "Redis", "Elasticsearch", "MySQL", "DDD"],
                "experience": [{"evidence_id": "resume_real_backend_1"}],
            },
            "source_event_ids": ["linkedin_jd_backend_1", "resume_real_backend_1"],
        },
    )

    assert result["status"] == "completed_read_only"
    output = result["output"]
    assert output["fit_score"] >= 0.8
    assert "Java" in output["matched_requirements"]
    assert "JVM tuning" in output["matched_requirements"]
    assert "high concurrency" in output["matched_requirements"]
    assert "Kafka" in output["matched_requirements"]
    assert "Redis" in output["matched_requirements"]
    assert "Elasticsearch" in output["matched_requirements"]
    assert "MySQL" in output["matched_requirements"]
    assert "unsupported_claims" in output and output["unsupported_claims"] == []
    assert {"linkedin_jd_backend_1", "resume_real_backend_1"}.issubset(set(output["evidence_ids"]))


def test_career_resume_export_generates_local_docx_and_pdf_artifacts(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    payload = {
        "filename": "ai-product-manager-resume",
        "headline": "AI Product Manager",
        "sections": [
            {"title": "Summary", "body": "Built LLM product and workflow automation systems."},
            {"title": "Evidence", "body": "Evidence: resume_evt_1"},
        ],
        "source_event_ids": ["resume_evt_1"],
    }
    client = TestClient(main.app)
    docx_response = client.post(
        "/api/career/resumes/export",
        headers={"x-par-password": "secret"},
        json={**payload, "format": "docx"},
    )
    pdf_response = client.post(
        "/api/career/resumes/export",
        headers={"x-par-password": "secret"},
        json={**payload, "format": "pdf"},
    )

    assert docx_response.status_code == 200
    docx_body = docx_response.json()
    assert docx_body["filename"] == "ai-product-manager-resume.docx"
    assert docx_body["mime_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    exported_docx = Document(BytesIO(base64.b64decode(docx_body["content_base64"])))
    exported_text = "\n".join(paragraph.text for paragraph in exported_docx.paragraphs)
    assert "AI Product Manager" in exported_text
    assert "Built LLM product" in exported_text
    assert "Evidence: resume_evt_1" in exported_text

    assert pdf_response.status_code == 200
    pdf_body = pdf_response.json()
    assert pdf_body["filename"] == "ai-product-manager-resume.pdf"
    assert pdf_body["mime_type"] == "application/pdf"
    assert base64.b64decode(pdf_body["content_base64"]).startswith(b"%PDF-")
    assert pdf_body["source_event_ids"] == ["resume_evt_1"]


def test_career_offers_endpoint_returns_interview_and_offer_tracking(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        if "FROM job_applications" in sql:
            return Cursor(
                [
                    (
                        "application_job_pm_ai_1_offer",
                        "job_pm_ai_1",
                        "offer",
                        "offer",
                        "review_offer_terms",
                        "submit_application",
                        "greenhouse",
                        ["offer_evt_1"],
                        {"company": "Example AI", "title": "AI Product Manager", "compensation": "待确认"},
                        "2026-06-08T09:40:00+08:00",
                        "2026-06-08T10:00:00+08:00",
                    ),
                    (
                        "application_job_growth_2_interview",
                        "job_growth_2",
                        "interviewing",
                        "interviewing",
                        "prepare_interview_if_replied",
                        "click_apply",
                        "lever",
                        ["interview_evt_1"],
                        {"company": "Acme", "title": "Growth PM", "interview_time": "2026-06-12T10:00:00+08:00"},
                        "2026-06-08T09:40:00+08:00",
                        "2026-06-08T10:00:00+08:00",
                    ),
                ]
            )
        return Cursor()

    install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).get("/api/career/offers", headers={"x-par-password": "secret"})

    assert response.status_code == 200
    body = response.json()
    assert body["filters"]["stages"] == ["interviewing", "offer"]
    assert len(body["offers"]) == 2
    assert body["offers"][0]["status"] == "offer"
    assert body["offers"][0]["payload"]["company"] == "Example AI"
    assert body["offers"][0]["next_step"] == "review_offer_terms"
    assert body["offers"][1]["stage"] == "interviewing"
    assert body["offers"][1]["source_event_ids"] == ["interview_evt_1"]


def test_career_application_patch_updates_local_tracking_state(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    from app import main

    def handler(sql, params):
        if "UPDATE job_applications" in sql:
            return Cursor(
                [
                    (
                        "application_job_pm_ai_1_submit_application",
                        "job_pm_ai_1",
                        "submitted",
                        "submitted",
                        "prepare_interview_if_replied",
                        "submit_application",
                        "linkedin",
                        ["jd_evt_1"],
                        {"user_note": "用户确认已投递"},
                        "2026-06-08T09:40:00+08:00",
                        "2026-06-08T10:00:00+08:00",
                    )
                ]
            )
        return Cursor()

    executed = install_fake_db(monkeypatch, main, handler)
    response = TestClient(main.app).patch(
        "/api/career/applications/application_job_pm_ai_1_submit_application",
        headers={"x-par-password": "secret"},
        json={
            "status": "submitted",
            "stage": "submitted",
            "next_step": "prepare_interview_if_replied",
            "user_note": "用户确认已投递",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == "application_job_pm_ai_1_submit_application"
    assert body["status"] == "submitted"
    assert body["stage"] == "submitted"
    assert body["next_step"] == "prepare_interview_if_replied"
    update_params = next(params for sql, params in executed if "UPDATE job_applications" in sql)
    assert update_params[0] == "submitted"
    assert update_params[1] == "submitted"
    assert update_params[2] == "prepare_interview_if_replied"
    assert update_params[-1] == "application_job_pm_ai_1_submit_application"
