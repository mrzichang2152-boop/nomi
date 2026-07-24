import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("DATABASE_URL", "postgresql://test")
os.environ.setdefault("REDIS_URL", "redis://test")


def _spec() -> dict:
    return {
        "title": "项目报价成本与利润分析",
        "purpose": "核对每个项目的成本、收入和利润率",
        "locale": "zh-CN",
        "currency": "CNY",
        "sheets": [
            {
                "name": "明细",
                "description": "每个报价项目的原始数据与计算结果",
                "columns": [
                    {"key": "item", "label": "项目", "type": "text"},
                    {"key": "revenue", "label": "收入", "type": "currency"},
                    {"key": "cost", "label": "成本", "type": "currency"},
                    {"key": "profit", "label": "利润", "type": "formula", "formula": "=B{row}-C{row}"},
                    {"key": "margin", "label": "利润率", "type": "formula", "formula": "=IFERROR(D{row}/B{row},0)"},
                ],
                "rows": [
                    {"item": "方案设计", "revenue": 120000, "cost": 72000},
                    {"item": "实施服务", "revenue": 80000, "cost": 56000},
                ],
                "source_evidence_ids": ["evt_quote_001"],
            },
            {
                "name": "摘要",
                "description": "管理层快速查看总收入、总成本和整体利润率",
                "columns": [
                    {"key": "metric", "label": "指标", "type": "text"},
                    {"key": "value", "label": "数值", "type": "text"},
                ],
                "rows": [
                    {"metric": "总收入", "value": "=SUM(明细!B2:B3)"},
                    {"metric": "总成本", "value": "=SUM(明细!C2:C3)"},
                    {"metric": "总利润", "value": "=SUM(明细!D2:D3)"},
                    {"metric": "整体利润率", "value": "=IFERROR(B4/B2,0)"},
                ],
                "source_evidence_ids": ["evt_quote_001"],
            },
        ],
    }


def test_render_spreadsheet_creates_typed_formula_driven_traceable_workbook(tmp_path):
    from app.font_support import preferred_cjk_font_name
    from app.spreadsheet_capability import render_spreadsheet

    output = tmp_path / "quote-analysis.xlsx"
    manifest_path = tmp_path / "manifest.json"
    manifest = render_spreadsheet(_spec(), output, manifest_path=manifest_path)

    assert output.is_file()
    assert manifest["artifact_type"] == "xlsx"
    assert manifest["capability_pack_id"] == "spreadsheet"
    assert manifest["sheet_names"] == ["明细", "摘要"]
    assert manifest["source_evidence_ids"] == ["evt_quote_001"]
    assert manifest["evidence_to_content_map"]["evt_quote_001"] == ["sheet:明细", "sheet:摘要"]
    assert manifest["quality_report"]["status"] == "passed"
    assert manifest["quality_report"]["formula_count"] == 8

    workbook = load_workbook(output, data_only=False)
    detail = workbook["明细"]
    summary = workbook["摘要"]
    assert detail.freeze_panes == "A2"
    assert detail.auto_filter.ref == "A1:E3"
    assert detail["D2"].value == "=B2-C2"
    assert detail["E3"].value == "=IFERROR(D3/B3,0)"
    assert detail["B2"].number_format != "General"
    assert detail["E2"].number_format == "0.0%"
    assert summary["B2"].value == "=SUM(明细!B2:B3)"
    assert summary["B5"].value == "=IFERROR(B4/B2,0)"
    assert all(cell.fill.fill_type == "solid" for cell in detail[1])
    assert all(cell.font.bold for cell in detail[1])
    assert all(cell.font.name == preferred_cjk_font_name() for cell in detail[1])
    assert detail["A2"].font.name == preferred_cjk_font_name()


def test_render_spreadsheet_formats_formula_results_from_business_semantics(tmp_path):
    from app.spreadsheet_capability import render_spreadsheet

    spec = {
        "title": "Q2 项目利润分析",
        "purpose": "核对利润与利润率",
        "currency": "CNY",
        "sheets": [
            {
                "name": "Q2利润分析",
                "columns": [
                    {"key": "project", "label": "项目", "type": "text"},
                    {"key": "revenue", "label": "收入", "type": "currency"},
                    {"key": "cost", "label": "成本", "type": "currency"},
                    {"key": "profit", "label": "利润", "type": "formula", "formula": "=B{row}-C{row}"},
                    {"key": "margin", "label": "利润率", "type": "formula", "formula": "=D{row}/B{row}"},
                ],
                "rows": [{"project": "Alpha", "revenue": 128000, "cost": 83000}],
                "source_evidence_ids": ["evt_q2_profit"],
            }
        ],
    }
    output = tmp_path / "q2-profit.xlsx"

    manifest = render_spreadsheet(spec, output)

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Q2利润分析"]
    assert sheet["D2"].number_format != "General"
    assert "%" in sheet["E2"].number_format
    assert manifest["quality_report"]["checks"]["semantic_number_formats"] is True
    assert manifest["quality_report"]["format_issues"] == []


def test_render_spreadsheet_truncates_long_office_core_metadata(tmp_path):
    from app.spreadsheet_capability import render_spreadsheet

    spec = _spec()
    spec["purpose"] = "用于验证长指令不会破坏 Excel 工作簿交付。" * 40
    output = tmp_path / "long-metadata.xlsx"

    render_spreadsheet(spec, output)

    workbook = load_workbook(output, data_only=False)
    assert len(workbook.properties.subject) <= 255
    assert workbook.properties.subject.startswith("用于验证长指令不会破坏 Excel 工作簿交付")


def test_validate_spreadsheet_rejects_duplicate_columns_and_unsafe_formulas():
    from app.spreadsheet_capability import SpreadsheetSpecError, validate_spreadsheet_spec

    spec = _spec()
    spec["sheets"][0]["columns"][1]["key"] = "item"
    with pytest.raises(SpreadsheetSpecError, match="duplicate column key"):
        validate_spreadsheet_spec(spec)

    spec = _spec()
    spec["sheets"][1]["rows"][0]["value"] = "=WEBSERVICE(\"https://example.com\")"
    with pytest.raises(SpreadsheetSpecError, match="unsafe formula function"):
        validate_spreadsheet_spec(spec)


def test_validate_spreadsheet_rejects_rows_with_unknown_fields():
    from app.spreadsheet_capability import SpreadsheetSpecError, validate_spreadsheet_spec

    spec = _spec()
    spec["sheets"][0]["rows"][0]["invented"] = 99
    with pytest.raises(SpreadsheetSpecError, match="unknown fields"):
        validate_spreadsheet_spec(spec)


def test_spreadsheet_cli_writes_verified_manifest(tmp_path):
    spec_path = tmp_path / "spec.json"
    output = tmp_path / "analysis.xlsx"
    manifest = tmp_path / "manifest.json"
    spec_path.write_text(json.dumps(_spec(), ensure_ascii=False), encoding="utf-8")
    script = Path(__file__).resolve().parents[1] / "scripts" / "nomi_spreadsheet_tool.py"

    completed = subprocess.run(
        [sys.executable, str(script), "--spec", str(spec_path), "--output", str(output), "--manifest", str(manifest)],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["status"] == "passed"
    assert payload["manifest"]["quality_report"]["checks"]["formulas_are_safe"] is True
    assert json.loads(manifest.read_text(encoding="utf-8"))["filename"] == "analysis.xlsx"
