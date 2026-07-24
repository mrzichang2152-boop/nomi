import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("DATABASE_URL", "postgresql://test")
os.environ.setdefault("REDIS_URL", "redis://test")


def _spec() -> dict:
    return {
        "title": "客户报价复核报告",
        "subtitle": "成本、利润率与下一步建议",
        "audience": "项目负责人和财务团队",
        "purpose": "基于客户报价资料给出可执行的复核结论",
        "author": "Nomi",
        "sections": [
            {
                "kind": "executive_summary",
                "heading": "执行摘要",
                "level": 1,
                "paragraphs": [
                    "本次报价总收入为20万元，总成本为12.8万元，整体毛利率为36%。",
                    "实施服务利润率低于方案设计，应优先复核人力投入假设。",
                ],
                "source_evidence_ids": ["evt_quote_001"],
            },
            {
                "kind": "analysis",
                "heading": "关键分析",
                "level": 1,
                "paragraphs": ["报价结构由方案设计和实施服务两部分组成。"],
                "bullets": ["方案设计收入12万元，成本7.2万元", "实施服务收入8万元，成本5.6万元"],
                "source_evidence_ids": ["evt_quote_001"],
            },
            {
                "kind": "table",
                "heading": "报价明细",
                "level": 2,
                "table": {
                    "headers": ["项目", "收入", "成本", "利润率"],
                    "rows": [["方案设计", "¥120,000", "¥72,000", "40%"], ["实施服务", "¥80,000", "¥56,000", "30%"]],
                },
                "source_evidence_ids": ["evt_quote_001"],
            },
            {
                "kind": "recommendations",
                "heading": "建议与行动",
                "level": 1,
                "numbered_items": ["复核实施服务工时", "确认报价有效期", "向负责人提交最终版本"],
                "source_evidence_ids": ["evt_quote_001"],
            },
        ],
    }


def test_render_document_creates_structured_evidence_grounded_docx(tmp_path):
    from app.document_capability import render_document
    from app.font_support import preferred_cjk_font_name

    output = tmp_path / "quote-review.docx"
    manifest_path = tmp_path / "manifest.json"
    manifest = render_document(_spec(), output, manifest_path=manifest_path)

    assert output.is_file()
    assert manifest["artifact_type"] == "docx"
    assert manifest["capability_pack_id"] == "document"
    assert manifest["section_headings"] == ["执行摘要", "关键分析", "报价明细", "建议与行动"]
    assert manifest["source_evidence_ids"] == ["evt_quote_001"]
    assert manifest["evidence_to_content_map"]["evt_quote_001"] == [
        "section:执行摘要",
        "section:关键分析",
        "section:报价明细",
        "section:建议与行动",
    ]
    assert manifest["quality_report"]["status"] == "passed"
    assert manifest["quality_report"]["table_count"] == 1
    assert manifest["quality_report"]["action_item_count"] == 3

    document = Document(output)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["执行摘要", "关键分析", "报价明细", "建议与行动"]
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "方案设计"
    assert document.sections[0].header.paragraphs[0].text == "客户报价复核报告"
    assert "Nomi" in document.sections[0].footer.paragraphs[0].text
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "整体毛利率为36%" in all_text
    assert "复核实施服务工时" in all_text
    cjk_font = preferred_cjk_font_name()
    assert document.styles["Normal"]._element.rPr.rFonts.get(qn("w:eastAsia")) == cjk_font
    assert document.styles["Title"]._element.rPr.rFonts.get(qn("w:eastAsia")) == cjk_font
    default_language = document.styles.element.xpath(".//w:docDefaults//w:lang")[0]
    assert default_language.get(qn("w:eastAsia")) == "zh-CN"
    theme_language = document.settings.element.xpath(".//w:themeFontLang")[0]
    assert theme_language.get(qn("w:eastAsia")) == "zh-CN"
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    text_runs = [run for paragraph in paragraphs for run in paragraph.runs if run.text]
    assert text_runs
    assert all(
        run._element.rPr is not None
        and run._element.rPr.rFonts is not None
        and run._element.rPr.rFonts.get(qn("w:eastAsia")) == cjk_font
        for run in text_runs
    )


def test_render_document_truncates_long_office_core_metadata(tmp_path):
    from app.document_capability import render_document

    spec = _spec()
    spec["purpose"] = "用于验证长指令不会破坏 Word 文档交付。" * 40
    output = tmp_path / "long-metadata.docx"

    render_document(spec, output)

    document = Document(output)
    assert len(document.core_properties.subject) <= 255
    assert document.core_properties.subject.startswith("用于验证长指令不会破坏 Word 文档交付")


def test_validate_document_rejects_duplicate_headings_and_placeholders():
    from app.document_capability import DocumentSpecError, validate_document_spec

    spec = _spec()
    spec["sections"][1]["heading"] = "执行摘要"
    with pytest.raises(DocumentSpecError, match="duplicate section heading"):
        validate_document_spec(spec)

    spec = _spec()
    spec["sections"][1]["paragraphs"] = ["TODO: 后续补充"]
    with pytest.raises(DocumentSpecError, match="placeholder content"):
        validate_document_spec(spec)


def test_validate_document_rejects_broken_table_shape():
    from app.document_capability import DocumentSpecError, validate_document_spec

    spec = _spec()
    spec["sections"][2]["table"]["rows"][0] = ["方案设计", "¥120,000"]
    with pytest.raises(DocumentSpecError, match="table row width"):
        validate_document_spec(spec)


@pytest.mark.parametrize("field", ["paragraphs", "bullets", "numbered_items"])
def test_validate_document_rejects_non_string_section_content(field):
    from app.document_capability import DocumentSpecError, validate_document_spec

    spec = _spec()
    spec["sections"][1][field] = [
        {
            "heading": "不允许的嵌套结构",
            "items": ["这类对象不能被静默转换成正文"],
        }
    ]

    with pytest.raises(DocumentSpecError, match=rf"{field} items must be strings"):
        validate_document_spec(spec)


def test_render_document_rejects_observed_metric_recast_as_target(tmp_path):
    from app.document_capability import DocumentSpecError, render_document

    spec = _spec()
    spec["sections"][-1]["numbered_items"] = ["持续监控错误率，确保 0.9% 水平保持稳定。"]
    packet = {
        "plan_input": {
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": "evt_quote_001",
                        "content": "API 发布后错误率从 2.8% 降至 0.9%。",
                    }
                ]
            }
        }
    }

    with pytest.raises(DocumentSpecError, match="observed metric 0.9% is recast as a target"):
        render_document(spec, tmp_path / "report.docx", task_packet=packet)


def test_render_document_rejects_unsupported_causal_inference(tmp_path):
    from app.document_capability import DocumentSpecError, render_document

    spec = _spec()
    spec["sections"][1]["paragraphs"] = ["周一上午工单集中，可能与周末积压有关。"]
    packet = {
        "plan_input": {
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": "evt_quote_001",
                        "content": "周一上午工单集中。",
                    }
                ]
            }
        }
    }

    with pytest.raises(DocumentSpecError, match="unsupported causal inference"):
        render_document(spec, tmp_path / "report.docx", task_packet=packet)


def test_render_document_rejects_missing_required_contract_section(tmp_path):
    from app.document_capability import DocumentSpecError, render_document

    spec = _spec()
    packet = {
        "plan_input": {
            "requirements_contract": {
                "sections": ["执行摘要", "事实与分析", "建议与行动"],
            }
        }
    }

    with pytest.raises(DocumentSpecError, match="missing required section.*事实与分析"):
        render_document(spec, tmp_path / "missing-section.docx", task_packet=packet)


def test_document_cli_writes_verified_manifest(tmp_path):
    spec_path = tmp_path / "spec.json"
    output = tmp_path / "report.docx"
    manifest = tmp_path / "manifest.json"
    spec_path.write_text(json.dumps(_spec(), ensure_ascii=False), encoding="utf-8")
    script = Path(__file__).resolve().parents[1] / "scripts" / "nomi_document_tool.py"

    completed = subprocess.run(
        [sys.executable, str(script), "--spec", str(spec_path), "--output", str(output), "--manifest", str(manifest)],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["manifest"]["quality_report"]["checks"]["heading_hierarchy_valid"] is True
    assert json.loads(manifest.read_text(encoding="utf-8"))["filename"] == "report.docx"
