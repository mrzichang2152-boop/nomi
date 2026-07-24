import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("DATABASE_URL", "postgresql://test")
os.environ.setdefault("REDIS_URL", "redis://test")


def _rich_spec() -> dict:
    return {
        "title": "普通人也能理解大语言模型",
        "subtitle": "从文字切片到逐步生成",
        "audience": "没有机器学习背景的产品团队",
        "purpose": "让听众理解 LLM 的核心工作流程和能力边界",
        "desired_action": "能用准确但不晦涩的语言向别人解释 LLM",
        "theme": "industrial-light",
        "slides": [
            {
                "role": "title",
                "title": "大语言模型，本质上是一个文字预测引擎",
                "subtitle": "它很强，但不是数据库，也没有人的意识",
                "source_evidence_ids": ["general_knowledge"],
            },
            {
                "role": "concept",
                "title": "第一步：文字先被转换成模型能计算的表示",
                "takeaway": "模型看到的不是完整句子，而是一串 token 和向量。",
                "points": ["文字被拆成 token", "token 映射为数字向量", "相近概念在向量空间更接近"],
                "visual": {
                    "type": "flow",
                    "nodes": ["一段文字", "Token", "向量"],
                },
                "source_evidence_ids": ["evt_llm_notes"],
            },
            {
                "role": "process",
                "title": "生成答案是一个不断重复的预测循环",
                "takeaway": "每生成一个 token，模型都会把它放回上下文再预测下一步。",
                "steps": ["读取上下文", "计算相关线索", "预测下一个 token", "追加并继续"],
                "source_evidence_ids": ["evt_llm_notes"],
            },
            {
                "role": "comparison",
                "title": "理解它的能力，也要理解它不是什么",
                "left": {"label": "擅长", "items": ["语言模式", "归纳与改写", "基于上下文生成"]},
                "right": {"label": "不等于", "items": ["事实数据库", "人的意识", "永远正确的答案"]},
                "source_evidence_ids": ["general_knowledge"],
            },
            {
                "role": "closing",
                "title": "最实用的判断：上下文越清楚，输出越可靠",
                "takeaway": "给模型目标、资料和边界，比只说“帮我写一下”更有效。",
                "actions": ["说清目标", "提供证据", "检查重要事实"],
                "source_evidence_ids": ["evt_llm_notes"],
            },
        ],
    }


def test_render_presentation_creates_editable_varied_evidence_grounded_deck(tmp_path):
    from app.presentation_capability import render_presentation

    output = tmp_path / "llm-explained.pptx"
    manifest_path = tmp_path / "manifest.json"
    result = render_presentation(_rich_spec(), output, manifest_path=manifest_path)

    assert output.is_file()
    assert manifest_path.is_file()
    assert result["artifact_type"] == "pptx"
    assert result["capability_pack_id"] == "presentation"
    assert result["capability_pack_version"] == "1.0.0"
    assert result["slide_count"] == 5
    assert result["slide_roles"] == ["title", "concept", "process", "comparison", "closing"]
    assert result["source_evidence_ids"] == ["general_knowledge", "evt_llm_notes"]
    assert result["evidence_to_content_map"]["evt_llm_notes"] == ["slide_2", "slide_3", "slide_5"]
    assert result["quality_report"]["status"] == "passed"
    assert result["quality_report"]["layout_kind_count"] >= 4
    assert result["quality_report"]["visual_slide_count"] >= 2
    assert result["quality_report"]["maximum_slide_text_chars"] <= 700

    deck = Presentation(str(output))
    assert len(deck.slides) == 5
    assert all(len(slide.shapes) >= 3 for slide in deck.slides)
    assert any(shape.shape_type == 1 for slide in deck.slides for shape in slide.shapes)
    extracted = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "文字先被转换" in extracted
    assert "读取上下文" in extracted
    assert "永远正确" in extracted
    assert "evt_llm_notes" not in extracted
    assert "general_knowledge" not in extracted
    assert "用户提供的证据" in extracted
    assert "通用知识" in extracted

    from app.opencode_artifact_worker import verify_artifact_file

    verification = verify_artifact_file(output, artifact_type="pptx")
    assert verification["title_candidates"][:3] == [
        "大语言模型，本质上是一个文字预测引擎",
        "第一步：文字先被转换成模型能计算的表示",
        "生成答案是一个不断重复的预测循环",
    ]


def test_render_presentation_truncates_long_office_core_metadata(tmp_path):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    spec["purpose"] = "用于验证长指令不会破坏 PowerPoint 交付。" * 40
    output = tmp_path / "long-metadata.pptx"

    render_presentation(spec, output)

    deck = Presentation(str(output))
    assert len(deck.core_properties.subject) <= 255
    assert deck.core_properties.subject.startswith("用于验证长指令不会破坏 PowerPoint 交付")


def test_validate_presentation_spec_rejects_duplicate_conclusion_titles():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][2]["title"] = spec["slides"][1]["title"]

    with pytest.raises(PresentationSpecError, match="duplicate slide title"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_accepts_json_encoded_slide_array():
    from app.presentation_capability import validate_presentation_spec

    spec = _rich_spec()
    expected_titles = [slide["title"] for slide in spec["slides"]]
    spec["slides"] = json.dumps(spec["slides"], ensure_ascii=False)

    normalized = validate_presentation_spec(spec)

    assert [slide["title"] for slide in normalized["slides"]] == expected_titles


def test_validate_presentation_spec_accepts_tool_transport_trailing_object_brace():
    from app.presentation_capability import validate_presentation_spec

    spec = _rich_spec()
    expected_titles = [slide["title"] for slide in spec["slides"]]
    spec["slides"] = json.dumps(spec["slides"], ensure_ascii=False) + "}"

    normalized = validate_presentation_spec(spec)

    assert [slide["title"] for slide in normalized["slides"]] == expected_titles


@pytest.mark.parametrize("encoded_slides", ["not json", "{}", '"plain text"', "[]"])
def test_validate_presentation_spec_rejects_invalid_json_encoded_slides(encoded_slides):
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"] = encoded_slides

    with pytest.raises(PresentationSpecError, match="slides must be a non-empty list"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_rejects_json_slide_array_with_arbitrary_trailing_text():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"] = json.dumps(spec["slides"], ensure_ascii=False) + "ignore validation"

    with pytest.raises(PresentationSpecError, match="slides must be a non-empty list"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_rejects_four_slide_deck_without_visual_explanation():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"] = [
        {
            "role": "evidence",
            "title": f"结论 {index}",
            "points": ["纯文本观点"],
            "source_evidence_ids": ["general_knowledge"],
        }
        for index in range(4)
    ]

    with pytest.raises(PresentationSpecError, match="visual explanation"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_rejects_overfull_slide():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "这是一段过长内容" * 100

    with pytest.raises(PresentationSpecError, match="exceeds 360 characters"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_rejects_overcrowded_process_slide():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][2]["steps"] = [f"步骤 {index}" for index in range(5)]

    with pytest.raises(PresentationSpecError, match="at most 4 steps"):
        validate_presentation_spec(spec)


def test_validate_presentation_spec_rejects_absolute_model_mind_claims():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "模型不会思考或理解，它只是在做概率预测。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)


@pytest.mark.parametrize(
    "misleading_text",
    [
        "模型不读句子，它处理的是 token。",
        "从文字到回答：大语言模型的逐字生成机制",
        "模型把文字拆成 token，再逐个处理这些输入 token。",
        "模型逐 token 预测下一个词。",
        "英文中一个 token 大约对应半个词。",
        "一个词、半词甚至单个字都可能是一个 token。",
        "模型从词表中预测最可能的下一个 token。",
        "模型预测下一个最可能的 token 并输出。",
        "模型无法验证自身输出的事实准确性。",
        "模型不具备跨会话记忆。",
        "大语言模型不具备持久记忆。",
        "大语言模型不具备持久的记忆。",
        "基础模型本身没有持久记忆，每次对话独立处理。",
        "基础模型没有持久记忆，每次推理独立于之前的交互。",
        "你的输入和模型已生成内容构成模型处理的全部内容。",
        "训练数据有截止时间，无法获知最新事件。",
        "上下文过长时内容会被截断。",
    ],
)
def test_validate_presentation_spec_rejects_known_llm_explainer_hazards(misleading_text):
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = misleading_text

    with pytest.raises(PresentationSpecError, match="misleading technical explanation"):
        validate_presentation_spec(spec)


def test_concept_takeaway_uses_responsive_readable_layout(tmp_path):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    takeaway = (
        "大语言模型把输入和输出都表示为 token，"
        "并基于当前上下文计算生成阶段的后续输出。"
    )
    spec["slides"][1]["takeaway"] = takeaway
    output = tmp_path / "responsive-concept.pptx"

    render_presentation(spec, output)

    deck = Presentation(str(output))
    matching_shape = next(
        shape
        for shape in deck.slides[1].shapes
        if getattr(shape, "text", "") == takeaway
    )
    paragraph = matching_shape.text_frame.paragraphs[0]
    assert matching_shape.width.inches >= 5.7
    assert matching_shape.height.inches >= 1.15
    assert paragraph.font.size.pt <= 17


def test_validate_presentation_spec_rejects_absolute_llm_claim_variants():
    from app.presentation_capability import PresentationSpecError, validate_presentation_spec

    spec = _rich_spec()
    spec["slides"][0]["subtitle"] = "它不靠思考，只靠预测。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "训练数据之外的新知识无法获取，除非接入工具。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = (
        "大语言模型擅长模式匹配与文字生成，但并非真正的理解或推理。"
    )

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "大语言模型并非真正的理解或推理，只是在逐 token 预测。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "模型没有真正的因果推理。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)

    spec = _rich_spec()
    spec["slides"][1]["takeaway"] = "LLM 没有真正的理解、记忆或推理能力。"

    with pytest.raises(PresentationSpecError, match="overstated technical claim"):
        validate_presentation_spec(spec)


def test_presentation_cli_writes_manifest_and_quality_report(tmp_path):
    spec_path = tmp_path / "spec.json"
    output = tmp_path / "deck.pptx"
    manifest = tmp_path / "manifest.json"
    spec_path.write_text(json.dumps(_rich_spec(), ensure_ascii=False), encoding="utf-8")
    script = Path(__file__).resolve().parents[1] / "scripts" / "nomi_presentation_tool.py"

    completed = subprocess.run(
        [
            sys.executable,
            str(script),
            "--spec",
            str(spec_path),
            "--output",
            str(output),
            "--manifest",
            str(manifest),
        ],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["status"] == "passed"
    assert payload["manifest"]["quality_report"]["status"] == "passed"
    assert json.loads(manifest.read_text(encoding="utf-8"))["filename"] == "deck.pptx"


def test_presentation_cli_rejects_deck_that_ignores_task_contract(tmp_path):
    spec_path = tmp_path / "spec.json"
    packet_path = tmp_path / "packet.json"
    output = tmp_path / "deck.pptx"
    manifest = tmp_path / "manifest.json"
    spec_path.write_text(
        json.dumps(
            {
                "title": "Test",
                "audience": "General",
                "purpose": "Test",
                "slides": [
                    {
                        "role": "title",
                        "title": "Test Slide",
                        "subtitle": "Testing",
                        "source_evidence_ids": ["evt_contract"],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    packet_path.write_text(
        json.dumps(
            {
                "original_goal": "制作一份4页PPT，让普通人理解LLM如何逐步生成答案",
                "plan_input": {
                    "requirements_contract": {
                        "audience": "普通人",
                        "page_count": 4,
                    }
                },
                "context": {
                    "evidence_pack": {
                        "items": [
                            {
                                "evidence_id": "evt_contract",
                                "content": "必须解释token、上下文、下一个token预测和模型能力边界。",
                            }
                        ]
                    }
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    script = Path(__file__).resolve().parents[1] / "scripts" / "nomi_presentation_tool.py"

    completed = subprocess.run(
        [
            sys.executable,
            str(script),
            "--spec",
            str(spec_path),
            "--output",
            str(output),
            "--manifest",
            str(manifest),
            "--packet",
            str(packet_path),
        ],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 1
    assert "expected 4 slides, got 1" in completed.stderr
    assert output.exists() is False
    assert manifest.exists() is False


def test_render_presentation_reports_task_contract_coverage(tmp_path):
    from app.presentation_capability import render_presentation

    packet = {
        "original_goal": "制作5页PPT，让产品团队理解LLM如何生成答案",
        "plan_input": {
            "requirements_contract": {
                "audience": "没有机器学习背景的产品团队",
                "page_count": 5,
            }
        },
        "context": {
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": "evt_llm_notes",
                        "content": "必须解释token、上下文、下一个token预测和模型能力边界。",
                    }
                ]
            }
        },
    }

    result = render_presentation(_rich_spec(), tmp_path / "contract.pptx", task_packet=packet)

    report = result["quality_report"]
    assert report["checks"]["request_contract_matched"] is True
    assert report["request_contract"]["expected_slide_count"] == 5
    assert report["request_contract"]["audience_matched"] is True
    assert report["request_contract"]["goal_terms_missing"] == []
    assert report["request_contract"]["required_topics_missing"] == []


def test_presentation_contract_accepts_semantically_equivalent_general_audience(tmp_path):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    spec["audience"] = "普通大众（无 AI 技术背景）"
    packet = {
        "original_goal": "制作5页PPT，让普通人理解LLM如何生成答案",
        "plan_input": {
            "requirements_contract": {
                "audience": "普通人",
                "page_count": 5,
            }
        },
    }

    result = render_presentation(spec, tmp_path / "audience.pptx", task_packet=packet)

    assert result["quality_report"]["request_contract"]["audience_matched"] is True


def test_presentation_contract_matches_llm_goal_to_chinese_name(tmp_path):
    from app.presentation_capability import render_presentation

    spec = json.loads(
        json.dumps(_rich_spec(), ensure_ascii=False).replace("LLM", "大语言模型")
    )
    packet = {
        "original_goal": "制作5页PPT，让普通人理解LLM如何生成答案",
        "plan_input": {"requirements_contract": {"page_count": 5}},
    }

    result = render_presentation(spec, tmp_path / "llm-alias.pptx", task_packet=packet)

    assert result["quality_report"]["request_contract"]["goal_terms_missing"] == []


def test_presentation_contract_accepts_semantically_equivalent_multi_clause_topic(tmp_path):
    from app.presentation_capability import render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "四项运营指标",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "展示四项指标（含原因未知标注）",
                "takeaway": "夜间响应变慢的原因未知，必须作为信息缺口展示。",
                "points": [
                    "首次响应时间：12分钟",
                    "平均解决时长：4.2小时",
                    "升级率：0.9%",
                    "夜间响应更慢，原因未知（已明确标注未知）",
                ],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "concept",
                "title": "现状与证据",
                "takeaway": "仅呈现已确认事实",
                "points": ["四项指标来自已提供记录"],
                "visual": {"type": "nodes", "nodes": ["指标", "证据"]},
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "行动仍待评估",
                "actions": ["补充证据", "评估行动"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "制作4页Q2客户支持复盘PPT",
        "plan_input": {
            "requirements_contract": {
                "page_count": 4,
                "audience": "管理层",
                "must_include": ["展示上述四项指标，其中原因未知必须明确标注未知"],
            }
        },
    }

    result = render_presentation(spec, tmp_path / "semantic-required-topic.pptx", task_packet=packet)

    assert result["quality_report"]["request_contract"]["required_topics_missing"] == []


def test_presentation_contract_accepts_real_opencode_q2_structure():
    from app.presentation_capability import validate_presentation_request_contract

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘汇报-修复验证",
        "audience": "管理层",
        "purpose": "帮助管理层了解Q2客户支持关键指标现状、证据基础及待评估改进方向",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘汇报-修复验证",
                "takeaway": "基于四项可验证指标，呈现Q2客户支持现状与待评估改进方向",
                "preview_points": ["首次响应时间：12分钟", "平均解决时长：4.2小时"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "展示四项指标数据",
                "subtitle": "Q2核心指标一览，原因未知项已明确标注",
                "takeaway": "夜间响应偏慢的原因明确标注为未知",
                "points": [
                    "首次响应时间：12分钟",
                    "平均解决时长：4.2小时",
                    "升级率：0.9%",
                    "夜间响应更慢——原因：未知",
                ],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "concept",
                "title": "总结现状与证据",
                "takeaway": "四项指标数据均已确认；夜间响应偏慢的具体原因尚未查明",
                "points": ["已确认事实：首次响应时间12分钟、平均解决时长4.2小时、升级率0.9%、夜间响应更慢"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "待评估行动",
                "subtitle": "三项候选改进方向",
                "takeaway": "以下三项行动均为待评估候选，尚未确认实施计划",
                "actions": ["优化夜间排班", "建立升级预警", "每周复盘"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "制作4页Q2客户支持复盘PPT",
        "requirements_contract": {
            "page_count": 4,
            "audience": "管理层",
            "must_include": [
                "标题与结论摘要",
                "展示上述四项指标，其中原因未知必须明确标注未知",
                "总结现状与证据",
                "列出三项待评估行动：优化夜间排班、建立升级预警、每周复盘",
            ],
        },
    }

    report = validate_presentation_request_contract(spec, packet)

    assert report["required_topics_missing"] == []


def test_presentation_contract_accepts_strict_q2_instruction_clauses_without_repeating_forbidden_text():
    from app.presentation_capability import validate_presentation_request_contract

    source_id = "evt_q2_support_strict"
    spec = {
        "title": "Q2客户支持复盘汇报-最终恢复验收",
        "audience": "管理层",
        "purpose": "15分钟复盘",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘汇报-最终恢复验收",
                "takeaway": "本报告呈现四项核心指标现状与三项待评估行动。",
                "preview_points": [
                    "首次响应时间：12分钟",
                    "平均解决时长：4.2小时",
                    "升级率：0.9%",
                    "夜间响应更慢（原因未知）",
                ],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "展示四项指标，原因未知已明确标注",
                "takeaway": "以下数据来自用户提供的证据。",
                "points": [
                    "首次响应时间：12分钟",
                    "平均解决时长：4.2小时",
                    "升级率：0.9%",
                    "夜间响应更慢——原因未知",
                ],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "concept",
                "title": "现状与证据总结",
                "takeaway": "以上为Q2客户支持的已提供指标。",
                "points": [
                    "首次响应时间为12分钟",
                    "平均解决时长为4.2小时",
                    "升级率为0.9%",
                    "夜间响应更慢，其原因未知",
                ],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "待评估行动",
                "takeaway": "以下三项行动待管理层评估。",
                "actions": ["优化夜间排班", "建立升级预警", "每周复盘"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "制作4页Q2客户支持复盘PPT",
        "requirements_contract": {
            "page_count": 4,
            "audience": "管理层",
            "must_include": [
                "标题与结论摘要",
                "展示上述四项指标，其中原因未知必须明确标注未知",
                "只总结上述现状与证据，不添加信息缺口或评价",
                "只列出三项待评估行动：优化夜间排班、建立升级预警、每周复盘，不要增加前提、缺口或行动细节",
            ],
        },
    }

    report = validate_presentation_request_contract(spec, packet)

    assert report["required_topics_missing"] == []


def test_presentation_contract_rejects_known_reason_even_when_unknown_appears_elsewhere():
    from app.presentation_capability import PresentationSpecError, validate_presentation_request_contract

    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘",
        "slides": [
            {
                "role": "evidence",
                "title": "展示四项指标数据",
                "takeaway": "夜间响应偏慢的原因是排班不足",
                "points": [
                    "首次响应时间：12分钟",
                    "平均解决时长：4.2小时",
                    "升级率：0.9%",
                    "夜间响应更慢——原因：排班不足",
                ],
            },
            {
                "role": "concept",
                "title": "其他待补信息",
                "takeaway": "客户满意度数据未知",
                "points": ["该数据仍待补充"],
            },
        ],
    }
    packet = {
        "requirements_contract": {
            "audience": "管理层",
            "must_include": ["展示上述四项指标，其中原因未知必须明确标注未知"],
        }
    }

    with pytest.raises(PresentationSpecError, match="required topics"):
        validate_presentation_request_contract(spec, packet)


def test_presentation_contract_rejects_missing_named_action_from_real_structure():
    from app.presentation_capability import PresentationSpecError, validate_presentation_request_contract

    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘",
        "slides": [
            {"role": "title", "title": "Q2客户支持复盘", "takeaway": "结论摘要"},
            {
                "role": "closing",
                "title": "待评估行动",
                "takeaway": "三项行动待评估",
                "actions": ["优化夜间排班", "每周复盘"],
            },
        ],
    }
    packet = {
        "requirements_contract": {
            "audience": "管理层",
            "must_include": ["列出三项待评估行动：优化夜间排班、建立升级预警、每周复盘"],
        }
    }

    with pytest.raises(PresentationSpecError, match="required topics"):
        validate_presentation_request_contract(spec, packet)


def test_presentation_contract_rejects_wrong_per_slide_outline():
    from app.presentation_capability import PresentationSpecError, validate_presentation_request_contract

    packet = {
        "requirements_contract": {
            "page_count": 2,
            "audience": "技术团队",
            "topic": "Auto Delivery Verification",
            "slide_outline": [
                {
                    "slide_number": 1,
                    "title": "Auto Delivery Verification",
                    "subtitle": "No status follow-up",
                    "must_include": [],
                },
                {
                    "slide_number": 2,
                    "title": "Acceptance Result",
                    "subtitle": "",
                    "must_include": [
                        "Server-side unified delivery",
                        "Android real-device received",
                    ],
                },
            ],
        }
    }
    spec = {
        "audience": "技术团队",
        "slides": [
            {
                "role": "concept",
                "title": "Auto Delivery Verification",
                "subtitle": "No status follow-up",
                "points": ["Server-side unified delivery"],
            },
            {
                "role": "comparison",
                "title": "Acceptance Result",
                "points": ["generic acceptance note"],
            },
        ],
    }

    with pytest.raises(PresentationSpecError, match="slide 2.*Android real-device received"):
        validate_presentation_request_contract(spec, packet)


def test_presentation_contract_does_not_count_metadata_as_required_topic_content(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    spec = _rich_spec()
    spec["purpose"] = "讲清太阳能板如何工作以及关键组件"
    packet = {
        "original_goal": "制作5页太阳能入门PPT",
        "plan_input": {
            "requirements_contract": {
                "page_count": 5,
                "required_topics": ["太阳能板如何工作", "关键组件"],
            }
        },
    }

    with pytest.raises(PresentationSpecError, match="required topics"):
        render_presentation(spec, tmp_path / "metadata-only.pptx", task_packet=packet)


def test_render_presentation_preserves_nested_agent_content_in_the_actual_deck(tmp_path):
    from app.presentation_capability import render_presentation

    spec = {
        "title": "LLM 如何逐步生成答案",
        "audience": "普通人",
        "purpose": "解释从输入到输出的核心过程",
        "slides": [
            {
                "role": "title",
                "title": "大语言模型如何逐步生成答案",
                "subtitle": "从输入到输出的四步拆解",
                "content": {"deck_tagline": "不神秘、不玄学"},
                "evidence_ids": ["evt_agent"],
            },
            {
                "role": "concept",
                "title": "Token 与上下文：模型看到的是编号",
                "visual": {
                    "type": "flow",
                    "nodes": [
                        {"label": "输入文字", "shape": "rect"},
                        {"label": "Token 编号", "shape": "chips"},
                        {"label": "上下文窗口", "shape": "frame"},
                    ],
                },
                "content": {
                    "explanation": "模型先把文字拆成 Token，再把 Token 转成可计算的编号。",
                    "key_points": [
                        "Token 数量会影响计算量",
                        "上下文窗口限定一次可处理的信息量",
                    ],
                    "analogy": "像把一句话拆成编号卡片。",
                },
                "evidence_ids": ["evt_agent"],
            },
            {
                "role": "process",
                "title": "答案由一次次下一 Token 预测组成",
                "content": {
                    "explanation": "模型每次预测一个 Token，并把结果追加回上下文。",
                    "process_steps": [
                        {"label": "接收输入", "description": "把问题转成 Token 序列"},
                        {"label": "预测", "description": "选择可能的下一个 Token"},
                        {"label": "继续", "description": "追加结果并再次预测"},
                    ],
                },
                "evidence_ids": ["evt_agent"],
            },
            {
                "role": "closing",
                "title": "能力边界：流畅不等于事实正确",
                "content": {
                    "explanation": "理解生成机制，才能更可靠地使用模型。",
                    "strengths": ["语言表达流畅", "擅长归纳与改写"],
                    "limits": ["可能生成错误事实", "重要结论仍需核验"],
                    "takeaway": "把模型当作高效草稿工具，而不是真理裁判。",
                },
                "evidence_ids": ["evt_agent"],
            },
        ],
    }

    output = tmp_path / "nested-agent-content.pptx"
    result = render_presentation(spec, output)
    deck = Presentation(str(output))
    extracted = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert "模型先把文字拆成 Token" in extracted
    assert "输入文字" in extracted
    assert "{'label'" not in extracted
    assert "上下文窗口限定" in extracted
    assert "接收输入" in extracted
    assert "把问题转成 Token 序列" in extracted
    assert "语言表达流畅" in extracted
    assert "可能生成错误事实" in extracted
    assert result["source_evidence_ids"] == ["evt_agent"]
    assert result["quality_report"]["checks"]["semantic_content_rendered"] is True
    assert result["quality_report"]["rendered_content_coverage"] >= 0.9


def test_render_presentation_preserves_visual_node_detail_and_counts_it_as_semantic_content(
    tmp_path,
):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    concept = next(slide for slide in spec["slides"] if slide["role"] == "concept")
    concept["visual"] = {
        "type": "flow",
        "nodes": [
            {"label": "原始文本", "content": "我喜欢机器学习"},
            {"label": "Token 序列", "content": "拆成模型可处理的片段"},
            {"label": "数字表示", "content": "映射为向量后参与计算"},
        ],
    }

    output = tmp_path / "visual-node-detail.pptx"
    result = render_presentation(spec, output)
    deck = Presentation(str(output))
    extracted = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert "原始文本" in extracted
    assert "我喜欢机器学习" in extracted
    assert "Token 序列" in extracted
    assert "拆成模型可处理的片段" in extracted
    assert "数字表示" in extracted
    assert "映射为向量后参与计算" in extracted
    assert result["quality_report"]["checks"]["semantic_content_rendered"] is True
    assert result["quality_report"]["semantic_segments_rendered"] == result["quality_report"][
        "semantic_segments_total"
    ]


def test_generated_concept_visual_preserves_complete_long_labels(tmp_path):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    concept = next(slide for slide in spec["slides"] if slide["role"] == "concept")
    concept.pop("visual", None)
    concept["points"] = [
        "Server-side unified delivery: single service routes all messages.",
        "Android real-device received: confirmed on actual hardware.",
    ]

    output = tmp_path / "complete-visual-labels.pptx"
    render_presentation(spec, output)
    deck = Presentation(str(output))
    concept_slide = deck.slides[1]
    visual_labels = [
        shape.text
        for shape in concept_slide.shapes
        if hasattr(shape, "text") and shape.left.inches >= 7.8
    ]

    assert "Server-side unified delivery" in visual_labels
    assert "Android real-device received" in visual_labels


def test_render_presentation_preserves_mixed_layout_content_from_agent_spec(tmp_path):
    from app.presentation_capability import render_presentation

    spec = {
        "title": "Q2 客户支持复盘汇报",
        "audience": "管理层",
        "purpose": "复盘 Q2 客户支持事实并给出待评估行动",
        "slides": [
            {
                "role": "title",
                "title": "Q2 客户支持复盘汇报",
                "subtitle": "基于已知事实，不推断未知原因",
                "source_evidence_ids": ["evt_q2_support"],
            },
            {
                "role": "section",
                "title": "整体平稳，夜间响应待改进",
                "takeaway": "关键指标保持稳定，但夜间响应更慢。",
                "left": {
                    "label": "关键指标",
                    "items": ["首次响应 12 分钟", "解决时长 4.2 小时", "升级率 0.9%"],
                },
                "right": {"label": "已知观察", "items": ["夜间响应更慢，原因未知"]},
                "source_evidence_ids": ["evt_q2_support"],
            },
            {
                "role": "evidence",
                "title": "四项关键事实",
                "takeaway": "这些数据是当前可追溯结论的边界。",
                "visual": {
                    "type": "comparison",
                    "nodes": [
                        {"label": "首次响应", "content": "12 分钟"},
                        {"label": "解决时长", "content": "4.2 小时"},
                        {"label": "升级率", "content": "0.9%"},
                        {"label": "夜间响应", "content": "较慢，原因未知"},
                    ],
                },
                "source_evidence_ids": ["evt_q2_support"],
            },
            {
                "role": "closing",
                "title": "待评估行动",
                "takeaway": "确认原因后再决定改进措施。",
                "visual": {
                    "type": "sequence",
                    "nodes": [
                        {"label": "核对排班", "content": "确认夜间覆盖情况"},
                        {"label": "抽样工单", "content": "区分排队与处理耗时"},
                        {"label": "复盘决策", "content": "基于证据决定下一步"},
                    ],
                },
                "source_evidence_ids": ["evt_q2_support"],
            },
        ],
    }

    output = tmp_path / "mixed-agent-layout.pptx"
    result = render_presentation(spec, output)
    deck = Presentation(str(output))
    extracted = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    for expected in (
        "首次响应 12 分钟",
        "解决时长 4.2 小时",
        "升级率 0.9%",
        "夜间响应更慢，原因未知",
        "较慢，原因未知",
        "确认夜间覆盖情况",
        "区分排队与处理耗时",
        "基于证据决定下一步",
    ):
        assert expected in extracted
    assert "待补充内容" not in extracted
    assert result["quality_report"]["checks"]["semantic_content_rendered"] is True
    assert result["quality_report"]["rendered_content_coverage"] >= 0.9


def test_render_presentation_quality_ignores_inactive_role_alternative_fields(tmp_path):
    from app.presentation_capability import render_presentation

    spec = _rich_spec()
    closing = spec["slides"][-1]
    closing["points"] = [
        "备用字段一",
        "备用字段二",
        "备用字段三",
        "备用字段四",
    ]

    result = render_presentation(spec, tmp_path / "role-alternatives.pptx")

    assert result["quality_report"]["status"] == "passed"
    assert result["quality_report"]["checks"]["semantic_content_rendered"] is True
    assert result["quality_report"]["semantic_segments_rendered"] == result["quality_report"][
        "semantic_segments_total"
    ]


def test_presentation_contract_rejects_required_topic_hidden_in_inactive_role_field(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    spec = _rich_spec()
    spec["slides"][-1]["points"] = ["只藏在备用字段的秘密行动"]
    packet = {
        "original_goal": "制作一份介绍大语言模型的PPT",
        "plan_input": {
            "requirements_contract": {
                "audience": "没有机器学习背景的产品团队",
                "page_count": 5,
                "must_include": ["只藏在备用字段的秘密行动"],
            }
        },
    }

    with pytest.raises(PresentationSpecError, match="does not cover required topics"):
        render_presentation(spec, tmp_path / "hidden-required-topic.pptx", task_packet=packet)


def test_presentation_rejects_unsupported_metric_interpretation_for_private_evidence(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "已知数据如下",
                "points": ["首次响应时间12分钟，反映团队日常响应效率良好"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "comparison",
                "title": "昼夜观察",
                "left": {"label": "日间", "items": ["未提供更多数据"]},
                "right": {"label": "夜间", "items": ["响应更慢，原因未知"]},
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "先补充证据再判断原因",
                "actions": ["核对夜间数据"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "基于已知事实制作Q2客户支持复盘",
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": "首次响应时间12分钟；夜间响应更慢但原因未知。",
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "unsupported-inference.pptx", task_packet=packet)


def test_presentation_rejects_unsupported_business_baseline_for_private_evidence(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "已知数据如下",
                "points": ["首次响应时间12分钟", "夜间响应更慢，原因未知"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "concept",
                "title": "现状与证据总结",
                "takeaway": "现有数据清晰呈现了Q2支持效率的基本面",
                "points": ["升级率0.9%"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "候选行动仍待评估",
                "actions": ["核对夜间数据"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "基于已知事实制作Q2客户支持复盘",
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": "首次响应时间12分钟；升级率0.9%；夜间响应更慢但原因未知。",
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "unsupported-business-baseline.pptx", task_packet=packet)


def test_presentation_rejects_invented_missing_analysis_for_private_evidence(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "已知数据如下",
                "points": ["首次响应时间12分钟", "升级事件类别分布尚未统计"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "comparison",
                "title": "昼夜观察",
                "left": {"label": "日间", "items": ["未提供更多数据"]},
                "right": {"label": "夜间", "items": ["响应更慢，原因未知"]},
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "先补充证据再判断原因",
                "actions": ["核对夜间数据"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "基于已知事实制作Q2客户支持复盘",
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": "首次响应时间12分钟；夜间响应更慢但原因未知。",
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "invented-missing-analysis.pptx", task_packet=packet)


def test_presentation_rejects_invented_missing_target_comparison_for_private_evidence(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "已知数据如下",
                "points": ["首次响应时间12分钟", "夜间响应更慢，原因未知"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "comparison",
                "title": "现状与证据总结",
                "left": {"label": "事实", "items": ["升级率0.9%"]},
                "right": {
                    "label": "未知",
                    "items": ["夜间响应更慢的原因未知", "未提供原因分析或目标对比"],
                },
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "候选行动仍待评估",
                "actions": ["核对夜间数据"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "基于已知事实制作Q2客户支持复盘",
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": "首次响应时间12分钟；升级率0.9%；夜间响应更慢但原因未知。",
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "invented-missing-target-comparison.pptx", task_packet=packet)


@pytest.mark.parametrize(
    "invented_gap",
    [
        "现有数据不足以对整体服务质量做出综合评价",
        "需进一步分析夜间响应数据以确认方案",
        "需明确预警阈值与触发机制",
        "需确定复盘频率、参与方与产出格式",
        "三项行动尚未经过数据验证",
        "夜间响应更慢，原因未知，暂无进一步解释",
        "以下行动为待评估选项，尚未实施",
    ],
)
def test_presentation_rejects_invented_information_gaps_when_request_forbids_them(
    tmp_path,
    invented_gap,
):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "已知数据如下",
                "points": ["首次响应时间12分钟", "夜间响应更慢，原因未知"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "concept",
                "title": "现状与证据总结",
                "takeaway": invented_gap,
                "points": ["升级率0.9%"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "待评估行动",
                "takeaway": "三项行动均为待评估状态",
                "actions": ["优化夜间排班", "建立升级预警", "每周复盘"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": (
            "基于已知事实制作Q2客户支持复盘；"
            "列出优化夜间排班、建立升级预警、每周复盘三项待评估行动；"
            "不得编造其他数据、信息缺口或无证据评价。"
        ),
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": (
                            "首次响应时间12分钟；平均解决时长4.2小时；"
                            "升级率0.9%；夜间响应更慢但原因未知。"
                        ),
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "invented-information-gap.pptx", task_packet=packet)


def test_presentation_rejects_invented_private_evidence_provenance(tmp_path):
    from app.presentation_capability import PresentationSpecError, render_presentation

    source_id = "evt_q2_support"
    spec = {
        "title": "Q2客户支持复盘",
        "audience": "管理层",
        "purpose": "复盘已知事实",
        "slides": [
            {
                "role": "title",
                "title": "Q2客户支持复盘",
                "subtitle": "仅使用已知事实",
                "source_evidence_ids": [source_id],
            },
            {
                "role": "evidence",
                "title": "当前指标",
                "takeaway": "所有数据均来自Q2实际运营记录",
                "points": ["首次响应时间12分钟", "夜间响应更慢，原因未知"],
                "source_evidence_ids": [source_id],
            },
            {
                "role": "comparison",
                "title": "昼夜观察",
                "left": {"label": "日间", "items": ["首次响应时间12分钟"]},
                "right": {"label": "夜间", "items": ["响应更慢，原因未知"]},
                "source_evidence_ids": [source_id],
            },
            {
                "role": "closing",
                "title": "下一步",
                "takeaway": "候选行动仍待评估",
                "actions": ["核对夜间数据"],
                "source_evidence_ids": [source_id],
            },
        ],
    }
    packet = {
        "original_goal": "基于已知事实制作Q2客户支持复盘",
        "plan_input": {
            "requirements_contract": {"source_policy": "must_use_private_evidence"},
            "evidence_pack": {
                "items": [
                    {
                        "evidence_id": source_id,
                        "content": "首次响应时间12分钟；夜间响应更慢但原因未知。",
                    }
                ]
            },
        },
    }

    with pytest.raises(PresentationSpecError, match="unsupported private-evidence claim"):
        render_presentation(spec, tmp_path / "invented-provenance.pptx", task_packet=packet)
