#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path
from typing import Any
from urllib.parse import unquote

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from app.presentation_capability import render_presentation
from app.document_capability import render_document
from app.image_capability import render_image
from app.spreadsheet_capability import render_spreadsheet


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

CHECKPOINT_UNSUPPORTED_PROVENANCE = re.compile(
    r"(?:均)?来自(?:已验证(?:的)?数据源|实际运营记录|内部系统|调研数据)"
)
CHECKPOINT_INFORMATION_GAP_FORBIDDANCE = re.compile(
    r"(?:不得|不要|禁止)(?:编造|添加|补充)[^。；]{0,48}信息缺口"
)
CHECKPOINT_INFORMATION_GAP_PATTERNS = (
    re.compile(
        r"(?:现有|当前|以上)?(?:数据|证据|指标)?(?:不足以|无法)"
        r"[^。；]{0,40}(?:评价|判断|说明|支持|确认)"
    ),
    re.compile(
        r"(?:需|需要|尚需|还需)(?:进一步)?"
        r"(?:分析|补充|核对|明确|确定|收集|获取|验证)[^。；]{0,40}"
    ),
    re.compile(r"(?:尚未|未)(?:经过|完成)?(?:数据)?验证"),
    re.compile(r"(?:暂无|尚无)[^。；]{0,24}(?:进一步)?(?:解释|说明|分析)"),
    re.compile(r"(?:行动|事项|方案)?[^。；]{0,16}(?:尚未|未)(?:实施|执行|开始|落地)"),
)


def main() -> int:
    packet_path = Path(required_env("NOMI_OPENCODE_STEP_PACKET"))
    workspace = Path(required_env("NOMI_ARTIFACT_WORKSPACE"))
    manifest_path = Path(required_env("NOMI_ARTIFACT_MANIFEST"))
    workspace.mkdir(parents=True, exist_ok=True)
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    artifact_type = infer_artifact_type(packet)
    builders = {
        "pptx": build_pptx,
        "xlsx": build_xlsx,
        "docx": build_docx,
        "image": build_image,
        "png": build_image,
    }
    builder = builders.get(artifact_type)
    if builder is None:
        raise SystemExit(f"Unsupported default artifact type: {artifact_type}")
    manifest = builder(packet, workspace)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": "passed", "artifact_manifest": manifest}, ensure_ascii=False))
    return 0


def required_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise SystemExit(f"{name} is required")
    return value


def infer_artifact_type(packet: dict[str, Any]) -> str:
    route = packet.get("route_decision") if isinstance(packet.get("route_decision"), dict) else {}
    plan_input = packet.get("plan_input") if isinstance(packet.get("plan_input"), dict) else {}
    route_input = plan_input.get("route") if isinstance(plan_input.get("route"), dict) else {}
    raw = str(
        route.get("artifact_type")
        or plan_input.get("artifact_type")
        or route_input.get("artifact_type")
        or packet.get("artifact_type")
        or ""
    )
    if raw:
        return raw.lower()
    goal = str(packet.get("original_goal") or packet.get("goal") or "")
    if any(token in goal.lower() for token in ["ppt", "pptx", "slides", "deck", "幻灯片"]):
        return "pptx"
    return "pptx"


def goal_from_packet(packet: dict[str, Any]) -> str:
    plan_input = packet.get("plan_input") if isinstance(packet.get("plan_input"), dict) else {}
    candidates = [
        plan_input.get("user_request"),
        plan_input.get("original_goal"),
        packet.get("original_goal"),
        packet.get("original_goal_summary"),
        packet.get("user_request"),
        packet.get("goal"),
    ]
    for candidate in candidates:
        text = normalize_goal_text(str(candidate or ""))
        if text:
            return text
    return "生成 PPT"


def normalize_goal_text(text: str) -> str:
    stripped = str(text or "").strip()
    if not stripped:
        return ""
    if re.search(r"%[0-9A-Fa-f]{2}", stripped):
        return unquote(stripped).strip()
    return stripped


def build_pptx(packet: dict[str, Any], workspace: Path) -> dict[str, Any]:
    checkpoint = load_checkpoint_spec(workspace, "nomi_presentation_spec.json")
    if checkpoint is not None:
        checkpoint, repair_changes = repair_presentation_checkpoint(checkpoint, packet)
        if repair_changes:
            safe_workspace_path(workspace, "nomi_presentation_spec.json").write_text(
                json.dumps(checkpoint, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        title = str(checkpoint.get("title") or requirements_contract_from_packet(packet).get("topic") or "presentation")
        output = safe_workspace_path(workspace, safe_filename(title) + ".pptx")
        manifest = render_presentation(checkpoint, output, task_packet=packet)
        if repair_changes:
            manifest["checkpoint_repair"] = {
                "status": "applied",
                "changes": repair_changes,
            }
        return manifest
    goal = goal_from_packet(packet)
    evidence_items = extract_evidence_items(packet)
    source_evidence_ids = [
        str(item.get("evidence_id") or item.get("event_id") or "").strip()
        for item in evidence_items
        if str(item.get("evidence_id") or item.get("event_id") or "").strip()
    ]
    requirements_contract = requirements_contract_from_packet(packet)
    requested_slide_count = requested_slide_count_from_contract(requirements_contract)
    if requested_slide_count is None:
        requested_slide_count = requested_slide_count_from_goal(goal)
    slides = outline_for_goal(goal, evidence_items, requested_slide_count=requested_slide_count)
    contract_topic = str(requirements_contract.get("topic") or "").strip()
    if contract_topic and slides:
        slides[0] = {**slides[0], "title": contract_topic}
    artifact_title = contract_topic or title_from_goal(goal)
    filename = safe_filename(artifact_title) + ".pptx"
    file_path = (workspace / filename).resolve()
    file_path.relative_to(workspace.resolve())
    specification = presentation_specification_from_outline(
        goal,
        slides,
        source_evidence_ids=source_evidence_ids,
        task_packet=packet,
    )
    specification["title"] = artifact_title
    return render_presentation(specification, file_path, task_packet=packet)


def build_xlsx(packet: dict[str, Any], workspace: Path) -> dict[str, Any]:
    checkpoint = load_checkpoint_spec(workspace, "nomi_spreadsheet_spec.json")
    if checkpoint is not None:
        title = str(checkpoint.get("title") or requirements_contract_from_packet(packet).get("topic") or "spreadsheet")
        output = safe_workspace_path(workspace, safe_filename(title) + ".xlsx")
        return render_spreadsheet(checkpoint, output, task_packet=packet)
    goal = goal_from_packet(packet)
    evidence_items = extract_evidence_items(packet)
    source_ids = evidence_ids(evidence_items)
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(evidence_items, start=1):
        content = evidence_content(item)
        revenue = extract_named_number(content, ("报价", "收入", "营收", "销售额"))
        cost = extract_named_number(content, ("成本", "支出", "费用"))
        if revenue is None and cost is None:
            continue
        rows.append(
            {
                "item": evidence_actor(item) or f"证据 {index}",
                "revenue": revenue,
                "cost": cost,
            }
        )
    sheets: list[dict[str, Any]] = []
    if rows:
        sheets.append(
            {
                "name": "核算明细",
                "columns": [
                    {"key": "item", "label": "项目/来源", "type": "text"},
                    {"key": "revenue", "label": "收入/报价", "type": "currency"},
                    {"key": "cost", "label": "成本", "type": "currency"},
                    {"key": "profit", "label": "利润", "type": "formula", "formula": "=B{row}-C{row}"},
                    {"key": "margin", "label": "利润率", "type": "formula", "formula": "=IFERROR(D{row}/B{row},0)"},
                ],
                "rows": rows,
                "source_evidence_ids": source_ids,
            }
        )
    evidence_rows = [
        {
            "source": str(item.get("source") or item.get("source_type") or "资料"),
            "actor": evidence_actor(item),
            "content": evidence_content(item),
        }
        for item in evidence_items
        if evidence_content(item)
    ]
    if not evidence_rows:
        evidence_rows = [{"source": "用户请求", "actor": "用户", "content": goal}]
    sheets.append(
        {
            "name": "资料明细",
            "columns": [
                {"key": "source", "label": "来源", "type": "text"},
                {"key": "actor", "label": "提供者", "type": "text"},
                {"key": "content", "label": "可追溯内容", "type": "text"},
            ],
            "rows": evidence_rows,
            "source_evidence_ids": source_ids,
        }
    )
    filename = safe_filename(title_from_goal(goal)) + ".xlsx"
    output = safe_workspace_path(workspace, filename)
    return render_spreadsheet(
        {
            "title": title_from_goal(goal),
            "purpose": goal,
            "sheets": sheets,
        },
        output,
        task_packet=packet,
    )


def build_docx(packet: dict[str, Any], workspace: Path) -> dict[str, Any]:
    checkpoint = load_checkpoint_spec(workspace, "nomi_document_spec.json")
    if checkpoint is not None:
        title = str(checkpoint.get("title") or requirements_contract_from_packet(packet).get("topic") or "document")
        output = safe_workspace_path(workspace, safe_filename(title) + ".docx")
        return render_document(checkpoint, output, task_packet=packet)
    goal = goal_from_packet(packet)
    evidence_items = extract_evidence_items(packet)
    source_ids = evidence_ids(evidence_items)
    evidence_lines = [
        f"{str(item.get('source') or item.get('source_type') or '资料')} / {evidence_actor(item) or '未知提供者'}：{evidence_content(item)}"
        for item in evidence_items
        if evidence_content(item)
    ]
    if not evidence_lines:
        evidence_lines = [f"当前仅有用户原始目标：{goal}。报告未补充任何未经用户提供或可追溯来源支持的事实。"]
    title = title_from_goal(goal)
    sections = [
        {
            "kind": "executive_summary",
            "heading": "执行摘要",
            "level": 1,
            "paragraphs": [
                f"本文档用于落实以下目标：{goal}。内容严格区分已提供事实、分析框架与待确认事项，不把模板性建议写成已经发生的事实。",
                "当前版本先整理可追溯资料并建立复核结构；任何金额、日期、责任人或结论如没有来源支持，均不得自行补写。",
            ],
            "source_evidence_ids": source_ids,
        },
        {
            "kind": "analysis",
            "heading": "事实与证据",
            "level": 1,
            "paragraphs": ["以下内容按来源逐条保留，便于读者核对原始事实并识别仍需补充的信息。"],
            "bullets": evidence_lines,
            "source_evidence_ids": source_ids,
        },
        {
            "kind": "recommendations",
            "heading": "下一步行动",
            "level": 1,
            "paragraphs": ["后续工作以证据完整、责任清晰和结果可验收为标准，避免仅以任务状态判断交付质量。"],
            "numbered_items": [
                "逐项确认关键事实、数据口径与缺失信息。",
                "为每个行动补充负责人、截止日期和可检查的验收标准。",
                "在最终交付前重新对照用户目标与来源证据，删除无依据结论。",
            ],
            "source_evidence_ids": source_ids,
        },
    ]
    filename = safe_filename(title) + ".docx"
    output = safe_workspace_path(workspace, filename)
    return render_document(
        {
            "title": title,
            "audience": audience_from_goal(goal),
            "purpose": goal,
            "sections": sections,
        },
        output,
        task_packet=packet,
    )


def build_image(packet: dict[str, Any], workspace: Path) -> dict[str, Any]:
    checkpoint = load_checkpoint_spec(workspace, "nomi_image_spec.json")
    if checkpoint is not None:
        title = str(checkpoint.get("title") or requirements_contract_from_packet(packet).get("topic") or "image")
        output = safe_workspace_path(workspace, safe_filename(title) + ".png")
        return render_image(checkpoint, output, task_packet=packet)
    goal = goal_from_packet(packet)
    evidence_items = extract_evidence_items(packet)
    source_ids = evidence_ids(evidence_items)
    blocks: list[dict[str, Any]] = []
    for index, item in enumerate(evidence_items[:4], start=1):
        content = evidence_content(item)
        if not content:
            continue
        blocks.append(
            {
                "kind": "callout",
                "heading": evidence_actor(item) or f"证据 {index}",
                "body": content[:220],
                "items": [f"来源：{str(item.get('source') or item.get('source_type') or '资料')}"],
                "source_evidence_ids": [str(item.get("evidence_id") or item.get("event_id") or "").strip()],
            }
        )
    if not blocks:
        blocks = [
            {
                "kind": "process",
                "heading": "目标",
                "body": goal[:220],
                "items": ["明确输入", "组织重点", "复核输出"],
                "source_evidence_ids": [],
            },
            {
                "kind": "callout",
                "heading": "证据边界",
                "body": "当前没有额外可追溯资料，因此图中不补充具体数据或未经证实的事实。",
                "items": ["保留来源", "标注缺口"],
                "source_evidence_ids": [],
            },
        ]
    title = title_from_goal(goal)
    filename = safe_filename(title) + ".png"
    output = safe_workspace_path(workspace, filename)
    return render_image(
        {
            "title": title,
            "subtitle": "基于用户目标与可追溯资料整理",
            "purpose": goal,
            "audience": audience_from_goal(goal),
            "visual_kind": "infographic",
            "canvas": {"width": 1600, "height": 1000, "background": "#F4F7F9"},
            "palette": {"ink": "#172033", "muted": "#475467", "accent": "#008A7A", "highlight": "#F2B84B"},
            "blocks": blocks,
            "footer": "Nomi · evidence-grounded visual",
        },
        output,
        task_packet=packet,
    )


def evidence_ids(evidence_items: list[dict[str, Any]]) -> list[str]:
    return list(
        dict.fromkeys(
            str(item.get("evidence_id") or item.get("event_id") or "").strip()
            for item in evidence_items
            if str(item.get("evidence_id") or item.get("event_id") or "").strip()
        )
    )


def evidence_content(item: dict[str, Any]) -> str:
    return re.sub(
        r"\s+",
        " ",
        str(item.get("content") or item.get("excerpt") or item.get("summary") or item.get("text") or ""),
    ).strip()


def evidence_actor(item: dict[str, Any]) -> str:
    return str(item.get("actor") or item.get("contact") or item.get("sender") or "").strip()


def extract_named_number(text: str, names: tuple[str, ...]) -> float | None:
    match = re.search(
        rf"(?:{'|'.join(re.escape(name) for name in names)})[^0-9]{{0,12}}([0-9]+(?:\.[0-9]+)?)",
        text,
        flags=re.IGNORECASE,
    )
    return float(match.group(1)) if match else None


def safe_workspace_path(workspace: Path, filename: str) -> Path:
    output = (workspace / filename).resolve()
    output.relative_to(workspace.resolve())
    return output


def load_checkpoint_spec(workspace: Path, filename: str) -> dict[str, Any] | None:
    path = safe_workspace_path(workspace, filename)
    if not path.is_file():
        return None
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"OpenCode checkpoint must be an object: {filename}")
    return payload


def repair_presentation_checkpoint(
    checkpoint: dict[str, Any],
    packet: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Apply narrow, auditable repairs before deterministic fallback rendering.

    OpenCode persists a structured checkpoint before calling the renderer. If a
    model call times out after a quality rejection, the fallback may recover
    that checkpoint, but it must not replay a known-invalid assertion. Repairs
    are deliberately limited to neutralizing invented provenance and removing
    information-gap clauses the user explicitly forbade.
    """

    requirements = requirements_contract_from_packet(packet)
    strict_private_evidence = (
        str(requirements.get("source_policy") or "").strip()
        == "must_use_private_evidence"
    )
    if not strict_private_evidence:
        return json.loads(json.dumps(checkpoint, ensure_ascii=False)), []

    repaired = json.loads(json.dumps(checkpoint, ensure_ascii=False))
    goal = "\n".join(
        str(value or "")
        for value in (
            packet.get("original_goal"),
            packet.get("original_goal_summary"),
            packet.get("user_request"),
            (packet.get("plan_input") or {}).get("user_request")
            if isinstance(packet.get("plan_input"), dict)
            else "",
        )
        if str(value or "").strip()
    )
    forbids_information_gaps = bool(
        CHECKPOINT_INFORMATION_GAP_FORBIDDANCE.search(goal)
    )
    changes: list[dict[str, str]] = []

    def repair_value(value: Any, path: str) -> Any:
        if isinstance(value, str):
            before = value
            after = CHECKPOINT_UNSUPPORTED_PROVENANCE.sub(
                "来自用户提供的证据",
                before,
            )
            if forbids_information_gaps:
                for pattern in CHECKPOINT_INFORMATION_GAP_PATTERNS:
                    after = pattern.sub("", after)
                after = re.sub(
                    r"(?:[—–-]+)?待评估(?:[：:]*)$",
                    "",
                    after.strip(),
                )
                after = after.strip(" \t\r\n，,。；;：:—–-")
            if after != before:
                changes.append({"path": path, "before": before, "after": after})
            return after
        if isinstance(value, list):
            repaired_items = [
                repair_value(item, f"{path}[{index}]")
                for index, item in enumerate(value)
            ]
            return [item for item in repaired_items if item not in ("", None)]
        if isinstance(value, dict):
            return {
                key: repair_value(item, f"{path}.{key}")
                for key, item in value.items()
            }
        return value

    slides = repaired.get("slides")
    if isinstance(slides, list):
        repaired["slides"] = [
            repair_value(slide, f"slides[{index}]")
            for index, slide in enumerate(slides)
        ]
    return repaired, changes


def presentation_specification_from_outline(
    goal: str,
    slides: list[dict[str, Any]],
    *,
    source_evidence_ids: list[str],
    task_packet: dict[str, Any] | None = None,
) -> dict[str, Any]:
    semantic_goal = semantic_goal_text(goal)
    requirements_contract = requirements_contract_from_packet(task_packet or {})
    audience = str(requirements_contract.get("audience") or "").strip() or audience_from_goal(semantic_goal)
    converted: list[dict[str, Any]] = []
    for index, slide in enumerate(slides):
        title = str(slide.get("title") or f"第 {index + 1} 页").strip()
        bullets = [str(item).strip() for item in slide.get("bullets") or [] if str(item).strip()]
        role = slide_role(index, len(slides))
        base = {
            "role": role,
            "title": title,
            "source_evidence_ids": source_evidence_ids,
        }
        if role == "title":
            base["subtitle"] = bullets[0] if bullets else semantic_goal
        elif role == "concept":
            base["takeaway"] = bullets[0] if bullets else "先建立核心概念。"
            base["points"] = bullets[1:] or bullets
            base["visual"] = {
                "type": "flow",
                "nodes": concept_visual_nodes(title, bullets),
            }
        elif role == "process":
            base["takeaway"] = bullets[0] if bullets else "按顺序理解这一过程。"
            base["steps"] = bullets or ["输入", "处理", "输出"]
        elif role == "comparison":
            midpoint = max(1, (len(bullets) + 1) // 2)
            base["left"] = {"label": "核心能力", "items": bullets[:midpoint] or ["待补充"]}
            base["right"] = {"label": "边界与判断", "items": bullets[midpoint:] or bullets[:midpoint] or ["待补充"]}
        elif role == "closing":
            base["takeaway"] = bullets[0] if bullets else "将理解转化为下一步行动。"
            base["actions"] = bullets[1:] or bullets
        else:
            base["takeaway"] = bullets[0] if bullets else "这一页提供必要证据。"
            base["points"] = bullets[1:] or bullets
        converted.append(base)
    return {
        "title": title_from_goal(goal),
        "subtitle": "由 Nomi 基于目标与可追溯资料生成",
        "audience": audience,
        "purpose": semantic_goal or "清晰解释主题并给出可执行结论",
        "desired_action": "理解核心概念，并能据此采取下一步行动",
        "theme": "industrial-light",
        "slides": converted,
    }


def audience_from_goal(goal: str) -> str:
    lowered = goal.lower()
    if "founder" in lowered or "创始人" in goal or "创业者" in goal:
        return "创业者和业务负责人"
    if "普通人" in goal or "non-technical" in lowered:
        return "没有专业技术背景的普通听众"
    if "产品" in goal:
        return "产品与业务团队"
    return "用户指定的目标听众"


def slide_role(index: int, total: int) -> str:
    if index == 0:
        return "title"
    if total >= 5 and index == total - 1:
        return "closing"
    sequence = ("concept", "process", "comparison", "evidence")
    return sequence[(index - 1) % len(sequence)]


def concept_visual_nodes(title: str, bullets: list[str]) -> list[str]:
    candidates = [title, *bullets]
    nodes = [re.sub(r"[，。,.：:]", "", item)[:18] for item in candidates if item]
    return (nodes + ["输入", "关键机制", "结果"])[:3]


def extract_evidence_items(packet: dict[str, Any]) -> list[dict[str, Any]]:
    candidates: list[Any] = []
    context = packet.get("context") if isinstance(packet.get("context"), dict) else {}
    evidence_pack = context.get("evidence_pack") if isinstance(context.get("evidence_pack"), dict) else {}
    candidates.extend(evidence_pack.get("items") or [])
    plan_input = packet.get("plan_input") if isinstance(packet.get("plan_input"), dict) else {}
    plan_evidence_pack = plan_input.get("evidence_pack") if isinstance(plan_input.get("evidence_pack"), dict) else {}
    candidates.extend(plan_evidence_pack.get("items") or [])
    return [dict(item) for item in candidates if isinstance(item, dict)]


def requested_slide_count_from_goal(goal: str) -> int | None:
    text = semantic_goal_text(goal)
    match = re.search(r"(?i)(?<!\d)(\d{1,2})\s*(?:页|張|张|p|pages?|slides?|slide|deck pages?)", text)
    if match:
        return clamp_slide_count(int(match.group(1)))
    chinese_digits = {
        "一": 1,
        "二": 2,
        "两": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
    }
    match = re.search(r"([一二两三四五六七八九十])\s*(?:页|張|张)", text)
    if match:
        return clamp_slide_count(chinese_digits[match.group(1)])
    english_digits = {
        "one": 1,
        "two": 2,
        "three": 3,
        "four": 4,
        "five": 5,
        "six": 6,
        "seven": 7,
        "eight": 8,
        "nine": 9,
        "ten": 10,
    }
    match = re.search(
        r"(?i)\b(one|two|three|four|five|six|seven|eight|nine|ten)\s+(?:pages?|slides?)\b",
        text,
    )
    if match:
        return clamp_slide_count(english_digits[match.group(1).lower()])
    return None


def requirements_contract_from_packet(packet: dict[str, Any]) -> dict[str, Any]:
    plan_input = packet.get("plan_input") if isinstance(packet.get("plan_input"), dict) else {}
    context = packet.get("context") if isinstance(packet.get("context"), dict) else {}
    route_decision = packet.get("route_decision") if isinstance(packet.get("route_decision"), dict) else {}
    candidates = (
        packet.get("requirements_contract"),
        plan_input.get("requirements_contract"),
        context.get("requirements_contract"),
        route_decision.get("requirements_contract"),
    )
    for candidate in candidates:
        if isinstance(candidate, dict):
            return dict(candidate)
    return {}


def requested_slide_count_from_contract(contract: dict[str, Any]) -> int | None:
    raw_value = contract.get("page_count") or contract.get("slide_count") or contract.get("pages")
    if raw_value is None:
        return None
    try:
        return clamp_slide_count(int(raw_value))
    except (TypeError, ValueError):
        return requested_slide_count_from_goal(str(raw_value))


def clamp_slide_count(value: int) -> int:
    return max(1, min(int(value), 20))


def semantic_goal_text(goal: str) -> str:
    text = normalize_goal_text(goal)
    text = re.sub(r"(?i)\bNOMI[_-]REAL[_-]PPT[_-]\d+[A-Z]?\b", " ", text)
    text = text.replace("_", " ")
    return re.sub(r"\s+", " ", text).strip()


def outline_for_goal(
    goal: str,
    evidence_items: list[dict[str, Any]],
    *,
    requested_slide_count: int | None = None,
) -> list[dict[str, Any]]:
    evidence_summary = summarize_evidence(evidence_items)
    semantic_goal = semantic_goal_text(goal)
    explicit_outline = explicit_slide_outline_from_goal(semantic_goal)
    if explicit_outline:
        return fit_slide_count(explicit_outline, requested_slide_count)
    lowered_goal = semantic_goal.lower()
    if "agent memory" in lowered_goal or ("agent" in lowered_goal and "memory" in lowered_goal) or "长期记忆" in semantic_goal:
        return fit_slide_count(
            [
                {
                    "title": "Agent Memory 是什么",
                    "bullets": [
                        "它让 agent 不只处理当前一句话，而能保留目标、偏好、关系和任务进展。",
                        "好的记忆系统会区分事实、对话上下文、任务状态和可检索资料。",
                        "对创始人来说，价值不在“记住”，而在持续推动客户、招聘和运营目标。",
                    ],
                },
                {
                    "title": "三类核心记忆",
                    "bullets": [
                        "KV 记忆适合稳定事实，例如身份、偏好、公司信息。",
                        "关系图谱适合人、公司、岗位、机会之间的关系和强弱变化。",
                        "RAG 适合邮件、聊天、文档等长文本证据，回答时按需检索。",
                    ],
                },
                {
                    "title": "记忆生命周期",
                    "bullets": [
                        "新消息先进入原始事件，再并行做分类、实体抽取和任务判断。",
                        "高置信事实写入 KV 或图谱，长文本进入向量索引并保留来源。",
                        "召回时必须带作用域，避免把 A 联系人的隐私错误带给 B 场景。",
                    ],
                },
                {
                    "title": "落地建议",
                    "bullets": [
                        "先围绕一个高价值目标验证，例如求职、销售或关系推进。",
                        "每条建议都要能追溯到来源，并允许用户纠错。",
                        "把记忆写入、检索和主动建议拆开，分别度量准确率和延迟。",
                    ],
                },
                {
                    "title": "本次生成依据",
                    "bullets": evidence_summary
                    or [
                        "当前没有额外私有材料，因此按用户目标生成通用说明。",
                        "后续可补充目标听众、行业背景和演示时长来继续细化。",
                    ],
                },
            ],
            requested_slide_count,
        )
    if "llm" in lowered_goal or "大模型" in semantic_goal or "语言模型" in semantic_goal:
        return fit_slide_count(
            [
            {
                "title": "普通人也能理解 LLM",
                "bullets": [
                    "LLM 可以理解成一个会根据上下文预测下一段话的文字引擎。",
                    "它不是数据库，也不是有意识的人，而是从大量文本中学会语言模式。",
                    "这份 PPT 用生活类比解释它如何读问题、找线索、生成回答。",
                ],
            },
            {
                "title": "它先把文字切成小块",
                "bullets": [
                    "模型不会直接看整句话，而是把文字拆成 token。",
                    "token 可以是字、词或词的一部分，随后会被转换成数字向量。",
                    "向量让模型能计算词与词之间的相似度和关系。",
                ],
            },
            {
                "title": "它用注意力找重点",
                "bullets": [
                    "注意力机制会判断当前生成内容最该参考哪些上下文。",
                    "例如回答“他是谁”时，模型会回看最近提到的人名和关系。",
                    "上下文越清楚，模型越容易给出稳定回答。",
                ],
            },
            {
                "title": "它一步一步生成答案",
                "bullets": [
                    "模型每次预测下一个 token，再把新 token 放回上下文继续预测。",
                    "看起来像一次性回答，实际是连续生成。",
                    "参数会影响回答速度、创造性和稳定性。",
                ],
            },
            {
                "title": "为什么需要 Nomi 的私有上下文",
                "bullets": [
                    "公开模型知道通用知识，但不知道用户的聊天、邮件、日程和关系。",
                    "Nomi 先从用户授权数据中取相关上下文，再让模型回答或执行任务。",
                    "这样才能回答“我明天去哪见谁”这类只存在于私有数据里的问题。",
                ],
            },
            {
                "title": "本次生成依据",
                "bullets": evidence_summary
                or [
                    "未发现额外私有材料，因此内容基于用户当前请求生成。",
                    "如果后续提供受众、时长或风格，PPT 可以继续迭代。",
                ],
            },
            ],
            requested_slide_count,
        )
    return fit_slide_count(
        [
        {"title": title_from_goal(goal), "bullets": ["根据用户目标生成初版结构。", *evidence_summary[:2]]},
        {"title": "目标和受众", "bullets": ["明确这份材料要解决的问题。", "说明面向谁、希望对方看完后采取什么行动。"]},
        {"title": "关键信息", "bullets": evidence_summary or ["当前没有额外证据，需后续补充业务背景。"]},
        {"title": "建议结构", "bullets": ["先讲背景，再讲方案，最后给出下一步行动。", "每页只表达一个核心观点。"]},
        {"title": "下一步", "bullets": ["补充真实数据。", "根据目标受众调整语言。", "确认视觉风格和页数。"]},
        ],
        requested_slide_count,
    )


def explicit_slide_outline_from_goal(goal: str) -> list[dict[str, Any]]:
    directives = re.findall(
        r"(?is)\bslide\s+(\d+)\s*:\s*(.+?)(?=(?:\.\s*)?\bslide\s+\d+\s*:|(?:\.\s*)?(?:do\s+not|don't|audience|purpose|exactly)\b|$)",
        goal or "",
    )
    if not directives:
        return []

    topic_match = re.search(
        r"(?i)\btitled\s+(.+?)(?=\.|,|;|$)",
        goal or "",
    )
    topic = topic_match.group(1).strip().strip("\"'") if topic_match else title_from_goal(goal)
    slides: list[dict[str, Any]] = []
    for slide_number, raw_content in sorted(directives, key=lambda item: int(item[0])):
        content = raw_content.strip().rstrip(".")
        if content.casefold() in {"title", "title slide", "cover", "cover slide"}:
            slides.append({"title": topic, "bullets": [topic]})
            continue
        slides.append(
            {
                "title": f"Slide {slide_number}",
                "bullets": [content],
            }
        )
    return slides


def fit_slide_count(slides: list[dict[str, Any]], requested_slide_count: int | None) -> list[dict[str, Any]]:
    if requested_slide_count is None:
        return slides
    target = clamp_slide_count(requested_slide_count)
    if len(slides) >= target:
        return slides[:target]
    expanded = list(slides)
    while len(expanded) < target:
        expanded.append(
            {
                "title": f"补充说明 {len(expanded) + 1}",
                "bullets": [
                    "围绕当前主题补充一个独立观点。",
                    "后续可根据用户提供的资料继续细化这一页。",
                ],
            }
        )
    return expanded


def summarize_evidence(evidence_items: list[dict[str, Any]]) -> list[str]:
    summaries: list[str] = []
    for item in evidence_items[:5]:
        source = str(item.get("source") or item.get("source_type") or "source").strip()
        actor = str(item.get("actor") or item.get("contact") or "").strip()
        content = re.sub(r"\s+", " ", str(item.get("content") or item.get("summary") or "")).strip()
        if content:
            prefix = f"{source}"
            if actor:
                prefix += f" / {actor}"
            summaries.append(f"{prefix}：{content[:110]}")
    return summaries


def add_slide(presentation: Presentation, title: str, bullets: list[str], *, index: int) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(246, 248, 251)

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.8), Inches(0.85))
    title_frame = title_box.text_frame
    title_frame.clear()
    title_para = title_frame.paragraphs[0]
    title_para.alignment = PP_ALIGN.LEFT
    title_run = title_para.add_run()
    title_run.text = title
    title_run.font.size = Pt(32)
    title_run.font.bold = True
    title_run.font.name = "Arial"
    title_run.font.color.rgb = rgb(16, 24, 39)

    accent = slide.shapes.add_shape(1, Inches(0.75), Inches(1.55), Inches(0.72), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = rgb(0, 135, 120)
    accent.line.color.rgb = rgb(0, 135, 120)

    body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.6), Inches(4.6))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    body_frame.margin_left = Inches(0.05)
    body_frame.margin_right = Inches(0.05)
    body_frame.clear()
    for bullet_index, bullet in enumerate(bullets[:6]):
        paragraph = body_frame.paragraphs[0] if bullet_index == 0 else body_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.space_after = Pt(10)
        paragraph.font.size = Pt(19)
        paragraph.font.name = "Arial"
        paragraph.font.color.rgb = rgb(44, 54, 74)

    footer = slide.shapes.add_textbox(Inches(0.78), Inches(6.88), Inches(11.8), Inches(0.3))
    footer_frame = footer.text_frame
    footer_frame.clear()
    footer_run = footer_frame.paragraphs[0].add_run()
    footer_run.text = f"Nomi artifact worker · slide {index + 1}"
    footer_run.font.size = Pt(9)
    footer_run.font.name = "Arial"
    footer_run.font.color.rgb = rgb(107, 119, 140)


def title_from_goal(goal: str) -> str:
    text = semantic_goal_text(goal) or "Nomi 生成材料"
    text = re.sub(r"^(那你)?帮我(做|写|生成|制作)?(一个|一份)?", "", text).strip()
    text = re.sub(
        r"(?i)^(please\s+)?(make|create|build|generate|write)\s+(me\s+)?(an?\s+)?(pptx?|slides?|slide\s+deck|deck|presentation)\s+(for|about|on)?\s*",
        "",
        text,
    ).strip()
    parts = [part.strip() for part in re.split(r"[，。,.]", text, maxsplit=1) if part.strip()]
    if len(parts) == 2 and re.fullmatch(r"(?i)pptx?|slides?|deck|幻灯片|文档|报告|表格", parts[0]):
        text = parts[1]
    elif parts:
        text = parts[0]
    text = re.sub(r"(?i)^(pptx?|slides?|deck|幻灯片|文档|报告|表格)\s*", "", text).strip()
    text = re.sub(r"^(让|给|使|关于|围绕|面向)\s*", "", text).strip()
    return text[:48] or "Nomi 生成材料"


def safe_filename(title: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "_", title).strip("._ ")
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned[:80] or "nomi_artifact"


def rgb(red: int, green: int, blue: int):
    from pptx.dml.color import RGBColor

    return RGBColor(red, green, blue)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"nomi_opencode_artifact_command failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        raise
