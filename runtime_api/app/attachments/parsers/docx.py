from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.attachments.models import AttachmentErrorCode, AttachmentRejected
from app.attachments.parsers.common import ParseResult, ParsedChunk


def _paragraph_text(paragraph) -> str:
    text = "".join(node.text or "" for node in paragraph._p.xpath(".//w:t")).strip()
    links: list[str] = []
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        relationship_id = hyperlink.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        relationship = paragraph.part.rels.get(relationship_id)
        if relationship is not None and relationship.is_external:
            links.append(str(relationship.target_ref))
    if links:
        text = f"{text} ({', '.join(links)})".strip()
    return text


def parse_docx(path: Path) -> ParseResult:
    try:
        document = Document(path)
        chunks: list[ParsedChunk] = []
        current_heading = ""
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = _paragraph_text(paragraph)
            if not text:
                continue
            style = str(paragraph.style.name or "")
            if style.lower().startswith("heading") or style.startswith("标题"):
                current_heading = text
            chunks.append(
                ParsedChunk(
                    text=text,
                    locator={
                        "paragraph": index,
                        "heading": current_heading,
                        "style": style,
                    },
                )
            )
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    chunks.append(
                        ParsedChunk(
                            text=text,
                            locator={
                                "table": table_index,
                                "row": row_index,
                                "heading": current_heading,
                            },
                        )
                    )
        media_count = sum(1 for relationship in document.part.rels.values() if "image" in relationship.reltype)
    except AttachmentRejected:
        raise
    except (PackageNotFoundError, KeyError, OSError, ValueError, TypeError) as exc:
        raise AttachmentRejected(AttachmentErrorCode.CORRUPT, internal_detail=type(exc).__name__) from exc

    return ParseResult(
        manifest={
            "parser_kind": "docx",
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "media_count": media_count,
        },
        chunks=tuple(chunks),
    )


__all__ = ["parse_docx"]
