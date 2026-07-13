from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Iterable

from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.attachments.models import AttachmentErrorCode, AttachmentRejected
from app.attachments.parsers.common import ParseResult, ParsedChunk


MAX_NONEMPTY_CELLS = 100_000
CSV_CHUNK_ROWS = 100
MAX_CSV_FIELD_CHARS = 16_000
MAX_CSV_ROW_CHARS = 32_000


def _decode(payload: bytes) -> tuple[str, str]:
    try:
        return payload.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        try:
            return payload.decode("gb18030"), "gb18030"
        except UnicodeDecodeError:
            match = from_bytes(payload).best()
            if match is None:
                raise AttachmentRejected(AttachmentErrorCode.CORRUPT)
            return str(match), str(match.encoding or "unknown")


def _cell_text(coordinate: str, formula: object, value: object) -> str:
    if isinstance(formula, str) and formula.startswith("="):
        return f"{coordinate}=公式:{formula};计算值:{'' if value is None else value}"
    return f"{coordinate}={formula}"


def parse_xlsx(path: Path) -> ParseResult:
    formula_book = None
    value_book = None
    try:
        formula_book = load_workbook(path, read_only=True, data_only=False)
        value_book = load_workbook(path, read_only=True, data_only=True)
        chunks: list[ParsedChunk] = []
        nonempty_count = 0
        sheets: list[dict[str, object]] = []
        for formula_sheet in formula_book.worksheets:
            value_sheet = value_book[formula_sheet.title]
            entries: list[str] = []
            min_row: int | None = None
            max_row = 0
            min_col: int | None = None
            max_col = 0
            value_rows = value_sheet.iter_rows()
            for formula_row, value_row in zip(formula_sheet.iter_rows(), value_rows):
                for formula_cell, value_cell in zip(formula_row, value_row):
                    if formula_cell.value is None:
                        continue
                    nonempty_count += 1
                    if nonempty_count > MAX_NONEMPTY_CELLS:
                        raise AttachmentRejected(AttachmentErrorCode.COMPLEXITY_LIMIT)
                    min_row = formula_cell.row if min_row is None else min(min_row, formula_cell.row)
                    max_row = max(max_row, formula_cell.row)
                    min_col = formula_cell.column if min_col is None else min(min_col, formula_cell.column)
                    max_col = max(max_col, formula_cell.column)
                    entries.append(_cell_text(formula_cell.coordinate, formula_cell.value, value_cell.value))
            if not entries or min_row is None or min_col is None:
                sheets.append({"name": formula_sheet.title, "range": None, "nonempty_cells": 0})
                continue
            cell_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
            chunks.append(
                ParsedChunk(
                    text=" | ".join(entries)[:64_000],
                    locator={"sheet": formula_sheet.title, "range": cell_range},
                )
            )
            sheets.append({"name": formula_sheet.title, "range": cell_range, "nonempty_cells": len(entries)})
    except AttachmentRejected:
        raise
    except Exception as exc:
        raise AttachmentRejected(AttachmentErrorCode.CORRUPT, internal_detail=type(exc).__name__) from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if value_book is not None:
            value_book.close()

    return ParseResult(
        manifest={
            "parser_kind": "xlsx",
            "sheet_count": len(sheets),
            "sheets": sheets,
            "nonempty_cell_count": nonempty_count,
        },
        chunks=tuple(chunks),
    )


def _chunk_csv_rows(rows: Iterable[tuple[int, list[str]]]) -> list[ParsedChunk]:
    chunks: list[ParsedChunk] = []
    batch: list[tuple[int, list[str]]] = []
    for item in rows:
        batch.append(item)
        if len(batch) >= CSV_CHUNK_ROWS:
            chunks.append(_csv_chunk(batch))
            batch = []
    if batch:
        chunks.append(_csv_chunk(batch))
    return chunks


def _csv_chunk(rows: list[tuple[int, list[str]]]) -> ParsedChunk:
    text = "\n".join(",".join(values) for _, values in rows)
    return ParsedChunk(
        text=text[:MAX_CSV_ROW_CHARS],
        locator={"row_start": rows[0][0], "row_end": rows[-1][0]},
    )


def parse_csv(path: Path) -> ParseResult:
    try:
        text, encoding = _decode(path.read_bytes())
        numbered_rows: list[tuple[int, list[str]]] = []
        max_columns = 0
        for row_number, row in enumerate(csv.reader(io.StringIO(text)), start=1):
            if any(len(field) > MAX_CSV_FIELD_CHARS for field in row):
                raise AttachmentRejected(AttachmentErrorCode.COMPLEXITY_LIMIT)
            if sum(len(field) for field in row) > MAX_CSV_ROW_CHARS:
                raise AttachmentRejected(AttachmentErrorCode.COMPLEXITY_LIMIT)
            numbered_rows.append((row_number, row))
            max_columns = max(max_columns, len(row))
    except AttachmentRejected:
        raise
    except (csv.Error, OSError, UnicodeError) as exc:
        raise AttachmentRejected(AttachmentErrorCode.CORRUPT, internal_detail=type(exc).__name__) from exc

    return ParseResult(
        manifest={
            "parser_kind": "csv",
            "encoding": encoding,
            "row_count": len(numbered_rows),
            "max_column_count": max_columns,
        },
        chunks=tuple(_chunk_csv_rows(numbered_rows)),
    )


__all__ = ["parse_csv", "parse_xlsx"]
