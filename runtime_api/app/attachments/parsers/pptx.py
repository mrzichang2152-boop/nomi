from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.attachments.models import AttachmentErrorCode, AttachmentRejected
from app.attachments.parsers.common import ParseResult, ParsedChunk


MAX_PPTX_SLIDES = 150


def parse_pptx(path: Path) -> ParseResult:
    try:
        presentation = Presentation(path)
        slide_count = len(presentation.slides)
        if slide_count > MAX_PPTX_SLIDES:
            raise AttachmentRejected(AttachmentErrorCode.COMPLEXITY_LIMIT)
        chunks: list[ParsedChunk] = []
        table_count = 0
        media_count = 0
        for slide_number, slide in enumerate(presentation.slides, start=1):
            content: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = str(shape.text or "").strip()
                    if text:
                        content.append(text)
                if getattr(shape, "has_table", False):
                    table_count += 1
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip(" |"):
                            content.append(row_text)
                if getattr(shape, "shape_type", None) == 13:
                    media_count += 1
            if content:
                chunks.append(
                    ParsedChunk(
                        text="\n".join(content),
                        locator={"slide": slide_number, "section": "content"},
                    )
                )
            try:
                notes = str(slide.notes_slide.notes_text_frame.text or "").strip()
            except (AttributeError, KeyError):
                notes = ""
            meaningful_notes = "\n".join(
                line for line in notes.splitlines() if line.strip() and line.strip() != str(slide_number)
            ).strip()
            if meaningful_notes:
                chunks.append(
                    ParsedChunk(
                        text=meaningful_notes,
                        locator={"slide": slide_number, "section": "notes"},
                    )
                )
    except AttachmentRejected:
        raise
    except (PackageNotFoundError, KeyError, OSError, ValueError, TypeError) as exc:
        raise AttachmentRejected(AttachmentErrorCode.CORRUPT, internal_detail=type(exc).__name__) from exc

    return ParseResult(
        manifest={
            "parser_kind": "pptx",
            "slide_count": slide_count,
            "table_count": table_count,
            "media_count": media_count,
        },
        chunks=tuple(chunks),
        requires_default_visual_sweep=False,
    )


__all__ = ["parse_pptx"]
