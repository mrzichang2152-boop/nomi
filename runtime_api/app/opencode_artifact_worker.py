from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable, Optional, Protocol

from psycopg.types.json import Jsonb

from app.capability_packs import select_capability_pack
from app.long_tail_agent import LongTailGraphRunner


OpenCodeExecutor = Callable[[dict[str, Any], str], dict[str, Any]]
ArtifactLoader = Callable[[Any, str], Optional[dict[str, Any]]]


class OpenCodeExecutionError(RuntimeError):
    pass


class SubprocessOpenCodeExecutor:
    def __init__(
        self,
        command: str,
        *,
        timeout_seconds: int = 600,
        fallback_command: str | None = None,
    ) -> None:
        self.command = str(command or "").strip()
        self.timeout_seconds = int(timeout_seconds)
        self.fallback_command = str(
            fallback_command if fallback_command is not None else os.getenv("OPENCODE_ARTIFACT_FALLBACK_COMMAND", "")
        ).strip()

    def __call__(self, step_packet: dict[str, Any], workspace_dir: str) -> dict[str, Any]:
        if not self.command:
            raise OpenCodeExecutionError("OPENCODE_ARTIFACT_COMMAND is not configured.")
        workspace = Path(workspace_dir)
        workspace.mkdir(parents=True, exist_ok=True)
        attachment_report = stage_evidence_attachments(step_packet, workspace)
        packet_path = workspace / "nomi_opencode_step_packet.json"
        manifest_path = workspace / "nomi_artifact_manifest.json"
        packet_path.write_text(json.dumps(step_packet, ensure_ascii=False, indent=2), encoding="utf-8")
        env = {
            **os.environ,
            "NOMI_OPENCODE_STEP_PACKET": str(packet_path),
            "NOMI_ARTIFACT_WORKSPACE": str(workspace),
            "NOMI_ARTIFACT_MANIFEST": str(manifest_path),
        }
        if manifest_path.exists():
            manifest_path.unlink()

        used_fallback = False
        primary_error: dict[str, Any] | None = None
        try:
            completed = self._run_command(self.command, workspace=workspace, env=env)
            if completed.returncode != 0:
                primary_error = {
                    "error_type": "CalledProcessError",
                    "message": f"OpenCode command failed with exit code {completed.returncode}",
                    "stdout": completed.stdout[-2000:],
                    "stderr": completed.stderr[-2000:],
                }
        except subprocess.TimeoutExpired as exc:
            primary_error = {
                "error_type": "TimeoutExpired",
                "message": f"OpenCode command timed out after {self.timeout_seconds} seconds",
                "stdout": str(exc.stdout or "")[-2000:],
                "stderr": str(exc.stderr or "")[-2000:],
            }
            completed = None

        if primary_error is not None:
            if not self.fallback_command:
                raise OpenCodeExecutionError(
                    f"{primary_error['message']}: {primary_error.get('stderr') or primary_error.get('stdout') or ''}"
                )
            if manifest_path.exists():
                manifest_path.unlink()
            completed = self._run_command(self.fallback_command, workspace=workspace, env=env)
            used_fallback = True
            if completed.returncode != 0:
                primary_detail = str(
                    primary_error.get("stderr")
                    or primary_error.get("stdout")
                    or primary_error.get("message")
                    or "unknown primary error"
                )[-1000:]
                fallback_detail = str(completed.stderr or completed.stdout or "unknown fallback error")[-1000:]
                raise OpenCodeExecutionError(
                    "OpenCode command failed and fallback command also failed. "
                    f"Primary: {primary_detail}; Fallback: {fallback_detail}"
                )

        if not manifest_path.exists():
            raise OpenCodeExecutionError("OpenCode command did not write NOMI_ARTIFACT_MANIFEST.")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        result = {
            "status": "passed",
            "summary": "OpenCode fallback command completed." if used_fallback else "OpenCode command completed.",
            "used_fallback": used_fallback,
            "stdout": completed.stdout[-4000:],
            "stderr": completed.stderr[-4000:],
            "attachment_staging": attachment_report,
            "artifact_manifest": manifest,
        }
        if primary_error is not None:
            result["primary_error"] = primary_error
        return result

    def _run_command(
        self,
        command: str,
        *,
        workspace: Path,
        env: dict[str, str],
    ) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            command,
            shell=True,
            cwd=str(workspace),
            env=env,
            text=True,
            capture_output=True,
            timeout=self.timeout_seconds,
            check=False,
        )


class ArtifactDbConnection(Protocol):
    def execute(self, sql: str, params: tuple[Any, ...] = ()) -> Any:
        ...


class OpenCodeArtifactWorker:
    def __init__(
        self,
        *,
        runner: LongTailGraphRunner,
        executor: OpenCodeExecutor,
        artifact_storage_dir: str | os.PathLike[str],
        artifact_loader: ArtifactLoader | None = None,
    ) -> None:
        self.runner = runner
        self.executor = executor
        self.artifact_storage_dir = Path(artifact_storage_dir)
        self.artifact_loader = artifact_loader or load_persisted_opencode_artifact

    def run_task(self, task_id: str, *, conn: ArtifactDbConnection | None = None) -> dict[str, Any]:
        task_id = str(task_id)
        self.artifact_storage_dir.mkdir(parents=True, exist_ok=True)
        try:
            self._complete_gather_step(task_id)
            state = self.runner.get_task_state(task_id)
            completed_step_ids = {
                str(step.get("step_id") or "")
                for step in state.get("completed_steps") or []
                if isinstance(step, dict)
            }
            verification_ready = (
                str(state.get("current_step_id") or "") == "verify_artifact_delivery"
                or (
                    "opencode_execute_artifact" in completed_step_ids
                    and "verify_artifact_delivery" not in completed_step_ids
                )
            )
            if verification_ready:
                if conn is None:
                    raise OpenCodeExecutionError(
                        "Cannot resume artifact verification without persistent artifact storage."
                    )
                artifact = self.artifact_loader(conn, task_id)
                if not artifact:
                    raise OpenCodeExecutionError(
                        "Generated artifact metadata is missing while resuming delivery verification."
                    )
            else:
                artifact = self._run_opencode_step(task_id, conn=conn)
                if artifact.get("status") == "awaiting_user_input":
                    return artifact
            self._complete_verification_step(task_id, artifact, conn=conn)
            final = self.runner.evaluate_final(task_id)
        except Exception as exc:
            self._block_task(task_id, reason=str(exc))
            return {
                "status": "blocked",
                "reason": str(exc),
                "error_type": type(exc).__name__,
            }
        return {
            "status": "completed" if final.get("status") == "passed" else str(final.get("status") or "partial"),
            "artifact": artifact,
            "final": final,
        }

    def _complete_gather_step(self, task_id: str) -> None:
        packet = self._ensure_step_packet(task_id)
        if packet["step_id"] != "gather_artifact_context":
            return
        task = self.runner.get_task_state(task_id)
        plan = self.runner.plans_by_task[task_id]
        plan_input = dict(plan.get("input") or {})
        context_plan = dict(plan_input.get("context_plan") or {})
        evidence_pack = dict(plan_input.get("evidence_pack") or {})
        missing_report = {
            "missing_evidence": list(evidence_pack.get("missing_evidence") or []),
            "coverage": dict(evidence_pack.get("coverage") or {}),
        }
        self.runner.complete_current_step(
            task_id,
            executor_result={
                "status": "passed",
                "summary": "已收集开放任务产物所需的上下文和证据缺口。",
                "outputs": {
                    "scoped_context_pack": {
                        "context_plan": context_plan,
                        "route_decision": task.get("route_decision") or {},
                    },
                    "evidence_pack": evidence_pack or {"items": [], "coverage": {}},
                    "missing_evidence_report": missing_report,
                },
                "evidence": [
                    {
                        "type": "source_ids",
                        "source_ids": list((task.get("route_decision") or {}).get("source_event_ids") or []),
                        "summary": "Context plan and evidence pack were built by Nomi runtime before OpenCode execution.",
                    }
                ],
                "action_events": [],
                "next_risk": "none",
            },
        )

    def _run_opencode_step(
        self,
        task_id: str,
        *,
        conn: ArtifactDbConnection | None,
    ) -> dict[str, Any]:
        packet = self._ensure_step_packet(task_id)
        if packet["step_id"] != "opencode_execute_artifact":
            raise OpenCodeExecutionError(f"Expected opencode_execute_artifact, got {packet['step_id']}.")
        workspace = self.artifact_storage_dir / task_id / "opencode_workspace"
        workspace.mkdir(parents=True, exist_ok=True)
        executor_result = dict(self.executor(packet, str(workspace)) or {})
        if str(executor_result.get("status") or "") == "awaiting_user_input":
            return self._pause_for_user_input(task_id, executor_result)
        manifest = dict(executor_result.get("artifact_manifest") or {})
        artifact = self._store_artifact_from_manifest(task_id, manifest, workspace, conn=conn)
        self.runner.complete_current_step(
            task_id,
            executor_result={
                "status": "passed",
                "summary": str(executor_result.get("summary") or "OpenCode 已生成产物。"),
                "outputs": {
                    "artifact_file_path": artifact["storage_path"],
                    "artifact_manifest": artifact,
                    "evidence_to_content_map": manifest.get("evidence_to_content_map") or {"artifact": "unmapped"},
                },
                "evidence": [
                    {
                        "type": "tool_result",
                        "tool": "opencode",
                        "artifact_id": artifact["artifact_id"],
                        "summary": str(executor_result.get("summary") or "OpenCode returned artifact manifest."),
                    }
                ],
                "action_events": [],
                "next_risk": "none",
            },
        )
        return artifact

    def _pause_for_user_input(self, task_id: str, executor_result: dict[str, Any]) -> dict[str, Any]:
        step = self.runner._current_step(task_id)
        step_id = str(step.get("step_id") or "opencode_execute_artifact")
        partial_outputs = dict(executor_result.get("partial_outputs") or {})
        self.runner.event_store.append_event(
            task_id=task_id,
            event_type="memory.patch_applied",
            step_id=step_id,
            payload={
                "memory_patch": {
                    "opencode_partial_outputs": partial_outputs,
                    "opencode_missing_fields": [
                        str(item) for item in executor_result.get("missing_fields") or []
                    ],
                }
            },
            idempotency_key=f"{task_id}:{step_id}:opencode_partial:{len(self.runner.event_store.task_events(task_id))}",
        )
        self.runner.request_human_input(
            task_id,
            step_id=step_id,
            input_type=str(executor_result.get("input_type") or "opencode_clarification"),
            question=str(executor_result.get("question") or "OpenCode 需要你补充信息后才能继续。"),
            options=[dict(item) for item in executor_result.get("options") or []],
        )
        return {
            "status": "awaiting_user_input",
            "input_type": str(executor_result.get("input_type") or "opencode_clarification"),
            "question": str(executor_result.get("question") or ""),
            "missing_fields": [str(item) for item in executor_result.get("missing_fields") or []],
            "options": [dict(item) for item in executor_result.get("options") or []],
            "partial_outputs": partial_outputs,
        }

    def _complete_verification_step(
        self,
        task_id: str,
        artifact: dict[str, Any],
        *,
        conn: ArtifactDbConnection | None,
    ) -> None:
        packet = self._ensure_step_packet(task_id)
        if packet["step_id"] != "verify_artifact_delivery":
            raise OpenCodeExecutionError(f"Expected verify_artifact_delivery, got {packet['step_id']}.")
        storage_path = Path(str(artifact.get("storage_path") or ""))
        if not storage_path.exists() or not storage_path.is_file():
            raise OpenCodeExecutionError("Stored artifact is missing before verification.")
        verification_report = verify_artifact_file(
            storage_path,
            artifact_type=str(artifact.get("artifact_type") or ""),
        )
        request_report = verify_artifact_against_request(packet, verification_report=verification_report)
        traceability_report = verify_artifact_traceability(
            artifact,
            verification_report=verification_report,
            source_evidence_text_by_id=source_evidence_text_by_id_from_packet(packet),
        )
        capability_report = verify_artifact_capability_contract(
            artifact,
            verification_report=verification_report,
        )
        artifact["verification_status"] = "verified"
        if conn is not None:
            update_opencode_artifact_verification(conn, artifact["artifact_id"], "verified")
        download_url = f"/api/artifacts/{artifact['artifact_id']}/download"
        self.runner.complete_current_step(
            task_id,
            executor_result={
                "status": "passed",
                "summary": "已校验 OpenCode 产物并生成下载入口。",
                "outputs": {
                    "verification_report": {
                        **verification_report,
                        "request": request_report,
                        "traceability": traceability_report,
                        "capability_pack": capability_report,
                        "artifact_id": artifact["artifact_id"],
                        "filename": artifact["filename"],
                        "mime_type": artifact["mime_type"],
                    },
                    "download_url": download_url,
                    "user_delivery_message": f"产物已生成并通过基础校验：{artifact['filename']}",
                },
                "evidence": [
                    {
                        "type": "tool_result",
                        "tool": "artifact.verify",
                        "artifact_id": artifact["artifact_id"],
                        "summary": "Artifact file exists in storage and can be downloaded.",
                    }
                ],
                "action_events": [],
                "next_risk": "none",
            },
        )

    def _ensure_step_packet(self, task_id: str) -> dict[str, Any]:
        state = self.runner.get_task_state(task_id)
        if state.get("current_node") == "select_step":
            packet = self.runner.run_next(task_id)
            return self._enrich_step_packet(task_id, packet)
        if state.get("current_node") != "awaiting_executor":
            raise OpenCodeExecutionError(f"Task is not awaiting executor: {state.get('current_node')}.")
        step = self.runner._current_step(task_id)
        packet = self.runner.build_step_packet(state, step)
        return self._enrich_step_packet(task_id, packet)

    def _enrich_step_packet(self, task_id: str, packet: dict[str, Any]) -> dict[str, Any]:
        state = self.runner.get_task_state(task_id)
        plan = dict(self.runner.plans_by_task.get(task_id) or {})
        plan_input = dict(plan.get("input") or {})
        route_decision = dict(state.get("route_decision") or {})
        requirements_contract = dict(
            plan_input.get("requirements_contract")
            or route_decision.get("requirements_contract")
            or {}
        )
        context = {
            **dict(packet.get("minimal_context") or {}),
            "context_plan": dict(plan_input.get("context_plan") or {}),
            "evidence_pack": dict(plan_input.get("evidence_pack") or {}),
            "requirements_contract": requirements_contract,
        }
        artifact_type = str(
            plan_input.get("artifact_type")
            or route_decision.get("artifact_type")
            or packet.get("artifact_type")
            or ""
        ).strip()
        return {
            **packet,
            "original_goal": str(state.get("original_goal") or packet.get("original_goal_summary") or ""),
            "plan_input": plan_input,
            "route_decision": route_decision,
            "task_route": route_decision,
            "artifact_type": artifact_type,
            "requirements_contract": requirements_contract,
            "context": context,
            "minimal_context": context,
        }

    def _store_artifact_from_manifest(
        self,
        task_id: str,
        manifest: dict[str, Any],
        workspace: Path,
        *,
        conn: ArtifactDbConnection | None,
    ) -> dict[str, Any]:
        source_path = Path(str(manifest.get("file_path") or manifest.get("storage_path") or ""))
        if not source_path.is_absolute():
            source_path = workspace / source_path
        source_path = source_path.resolve()
        workspace_root = workspace.resolve()
        try:
            source_path.relative_to(workspace_root)
        except ValueError as exc:
            raise OpenCodeExecutionError("OpenCode artifact path is outside workspace.") from exc
        if not source_path.exists() or not source_path.is_file():
            raise OpenCodeExecutionError("OpenCode artifact file does not exist.")

        artifact_type = str(manifest.get("artifact_type") or infer_artifact_type(source_path.name))
        filename = safe_artifact_filename(str(manifest.get("filename") or source_path.name), artifact_type)
        artifact_id = str(manifest.get("artifact_id") or f"artifact_{uuid.uuid4().hex}")
        target_dir = self.artifact_storage_dir / task_id / "artifacts"
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = (target_dir / filename).resolve()
        target_path.relative_to(self.artifact_storage_dir.resolve())
        if source_path != target_path:
            shutil.copy2(source_path, target_path)
        artifact = {
            "artifact_id": artifact_id,
            "task_run_id": task_id,
            "artifact_type": artifact_type,
            "filename": filename,
            "mime_type": str(manifest.get("mime_type") or mime_type_for_artifact(artifact_type)),
            "storage_path": str(target_path),
            "version": int(manifest.get("version") or 1),
            "source_evidence_ids": [str(item) for item in manifest.get("source_evidence_ids") or []],
            "verification_status": "generated",
            "evidence_to_content_map": manifest.get("evidence_to_content_map") or {},
            "capability_pack_id": str(manifest.get("capability_pack_id") or ""),
            "capability_pack_version": str(manifest.get("capability_pack_version") or ""),
            "slide_roles": [str(item) for item in manifest.get("slide_roles") or []],
            "quality_report": dict(manifest.get("quality_report") or {}),
        }
        if conn is not None:
            persist_opencode_artifact(conn, task_id, artifact)
        return artifact

    def _block_task(self, task_id: str, *, reason: str) -> None:
        try:
            state = self.runner.state_by_task[task_id]
        except KeyError:
            return
        step_id = str(state.get("current_step_id") or "")
        self.runner.event_store.append_event(
            task_id=task_id,
            event_type="fallback.decided",
            step_id=step_id or None,
            payload={
                "fallback_decision": {
                    "action": "stop",
                    "reason": str(reason or "OpenCode artifact task failed."),
                    "checkpoint_id": "",
                }
            },
            idempotency_key=(
                f"{task_id}:{step_id or 'artifact'}:worker_blocked:"
                f"{len(self.runner.event_store.task_events(task_id))}"
            ),
        )
        state["status"] = "blocked"
        state["current_node"] = "blocked"
        state["failure_reason"] = str(reason or "OpenCode artifact task failed.")


def persist_opencode_artifact(conn: ArtifactDbConnection, task_id: str, artifact: dict[str, Any]) -> None:
    conn.execute(
        """
        INSERT INTO task_runs (
          task_run_id, task_type, source_event_ids, pipeline_id, route_type,
          status, idempotency_key, risk_permission, requires_user_confirmation,
          final_user_visible_summary, payload
        )
        VALUES (%s, %s, %s::TEXT[], %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (task_run_id) DO UPDATE
        SET status = EXCLUDED.status,
            final_user_visible_summary = EXCLUDED.final_user_visible_summary,
            payload = task_runs.payload || EXCLUDED.payload,
            updated_at = now()
        """,
        (
            task_id,
            "artifact_creation",
            [str(item) for item in artifact.get("source_evidence_ids") or []],
            "open_task_opencode_artifact_pipeline",
            "long_tail_agent",
            "running",
            f"opencode_artifact:{task_id}",
            "local_artifact_only",
            False,
            f"OpenCode generated {artifact.get('filename')}",
            Jsonb({"artifact": artifact}),
        ),
    )
    conn.execute(
        """
        INSERT INTO task_artifacts (
          artifact_id, task_run_id, artifact_type, filename, mime_type, storage_path,
          version, source_evidence_ids, verification_status
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s::TEXT[], %s)
        ON CONFLICT (artifact_id) DO UPDATE
        SET filename = EXCLUDED.filename,
            mime_type = EXCLUDED.mime_type,
            storage_path = EXCLUDED.storage_path,
            verification_status = EXCLUDED.verification_status
        """,
        (
            artifact["artifact_id"],
            task_id,
            artifact["artifact_type"],
            artifact["filename"],
            artifact["mime_type"],
            artifact["storage_path"],
            artifact.get("version") or 1,
            [str(item) for item in artifact.get("source_evidence_ids") or []],
            artifact.get("verification_status") or "generated",
        ),
    )


def load_persisted_opencode_artifact(
    conn: ArtifactDbConnection,
    task_id: str,
) -> dict[str, Any] | None:
    row = conn.execute(
        """
        SELECT payload
        FROM task_runs
        WHERE task_run_id = %s
        """,
        (task_id,),
    ).fetchone()
    if not row:
        return None
    payload = row[0]
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except json.JSONDecodeError:
            return None
    if not isinstance(payload, dict):
        return None
    artifact = payload.get("artifact")
    return dict(artifact) if isinstance(artifact, dict) else None


def update_opencode_artifact_verification(
    conn: ArtifactDbConnection,
    artifact_id: str,
    verification_status: str,
) -> None:
    conn.execute(
        """
        UPDATE task_artifacts
        SET verification_status = %s
        WHERE artifact_id = %s
        """,
        (verification_status, artifact_id),
    )


def verify_artifact_file(path: Path, *, artifact_type: str) -> dict[str, Any]:
    artifact_type = str(artifact_type or "").lower()
    size = path.stat().st_size
    if size <= 0:
        raise OpenCodeExecutionError("Artifact file is empty.")
    if artifact_type == "pptx":
        return verify_pptx_content(
            path,
        )
    if artifact_type == "docx":
        return verify_docx_content(path)
    if artifact_type == "xlsx":
        return verify_xlsx_content(path)
    if artifact_type in {"image", "png", "jpg", "jpeg"}:
        return verify_image_content(path, artifact_type=artifact_type)
    if artifact_type == "markdown":
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            raise OpenCodeExecutionError("Markdown artifact is empty.")
        return {
            "status": "passed",
            "artifact_type": artifact_type,
            "checks": ["non_empty_text"],
            "file_size": size,
        }
    return {
        "status": "passed",
        "artifact_type": artifact_type or "artifact",
        "checks": ["non_empty_file"],
        "file_size": size,
    }


def verify_artifact_capability_contract(
    artifact: dict[str, Any],
    *,
    verification_report: dict[str, Any],
) -> dict[str, Any]:
    pack_id = str(artifact.get("capability_pack_id") or "").strip()
    if not pack_id:
        return {
            "status": "skipped",
            "checks": ["legacy_artifact_without_capability_pack"],
        }

    artifact_type = str(
        artifact.get("artifact_type") or verification_report.get("artifact_type") or ""
    ).strip().lower()
    pack = select_capability_pack(artifact_type)
    if pack is None or pack.pack_id != pack_id:
        raise OpenCodeExecutionError(
            f"Artifact references unavailable capability pack '{pack_id}' for '{artifact_type}'."
        )
    version = str(artifact.get("capability_pack_version") or "").strip()
    if version != pack.version:
        raise OpenCodeExecutionError(
            f"Artifact capability pack version mismatch: expected {pack.version}, got {version or 'missing'}."
        )

    quality = artifact.get("quality_report")
    if not isinstance(quality, dict) or str(quality.get("status") or "") != "passed":
        raise OpenCodeExecutionError("Capability-pack artifact is missing a passed quality report.")

    profile = dict(pack.quality_profile)
    generated_checks = quality.get("checks") if isinstance(quality.get("checks"), dict) else {}
    source_ids = [str(item).strip() for item in artifact.get("source_evidence_ids") or [] if str(item).strip()]
    evidence_map = artifact.get("evidence_to_content_map") if isinstance(artifact.get("evidence_to_content_map"), dict) else {}
    if profile.get("require_source_mapping") and any(not evidence_map.get(source_id) for source_id in source_ids):
        raise OpenCodeExecutionError("Capability-pack artifact has incomplete source evidence mapping.")

    common = [
        "capability_pack_resolved",
        "capability_pack_version_matched",
        "capability_quality_report_passed",
    ]
    if artifact_type == "xlsx":
        sheet_count = int(verification_report.get("sheet_count") or 0)
        formula_count = int(verification_report.get("formula_count") or 0)
        formula_preview_count = int(quality.get("formula_preview_count") or 0)
        if sheet_count <= 0:
            raise OpenCodeExecutionError("Spreadsheet capability requires at least one readable worksheet.")
        if int(quality.get("sheet_count") or sheet_count) != sheet_count:
            raise OpenCodeExecutionError("Spreadsheet sheet count differs between generation and independent verification.")
        if formula_count and formula_preview_count != formula_count:
            raise OpenCodeExecutionError("Spreadsheet formulas are missing deterministic preview values.")
        if profile.get("require_filters") and not (
            generated_checks.get("filters_and_freeze_panes")
            and int(verification_report.get("filtered_sheet_count") or 0) == sheet_count
        ):
            raise OpenCodeExecutionError("Spreadsheet capability requires filters on every worksheet.")
        if profile.get("require_freeze_panes") and int(verification_report.get("frozen_sheet_count") or 0) != sheet_count:
            raise OpenCodeExecutionError("Spreadsheet capability requires frozen headers on every worksheet.")
        format_issues = [
            str(item).strip()
            for item in verification_report.get("semantic_number_format_issues") or []
            if str(item).strip()
        ]
        if format_issues:
            raise OpenCodeExecutionError(
                "Spreadsheet capability has invalid business number formats: " + "; ".join(format_issues)
            )
        return {
            "status": "passed",
            "capability_pack_id": pack.pack_id,
            "capability_pack_version": pack.version,
            "checks": [*common, "spreadsheet_structure", "spreadsheet_formula_previews", "spreadsheet_navigation", "spreadsheet_business_number_formats"],
            "sheet_count": sheet_count,
            "formula_count": formula_count,
        }
    if artifact_type == "docx":
        body_characters = int(verification_report.get("body_character_count") or 0)
        minimum_body = int(profile.get("minimum_body_characters") or 120)
        if body_characters < minimum_body:
            raise OpenCodeExecutionError(
                f"Document body is too short: expected at least {minimum_body} characters, got {body_characters}."
            )
        if int(verification_report.get("heading_count") or 0) <= 0:
            raise OpenCodeExecutionError("Document capability requires a readable heading hierarchy.")
        if profile.get("require_heading_hierarchy") and not generated_checks.get("heading_hierarchy_valid"):
            raise OpenCodeExecutionError("Document capability is missing a passed heading hierarchy check.")
        if profile.get("reject_placeholder_content") and not generated_checks.get("no_placeholder_content"):
            raise OpenCodeExecutionError("Document capability did not prove placeholder content was removed.")
        return {
            "status": "passed",
            "capability_pack_id": pack.pack_id,
            "capability_pack_version": pack.version,
            "checks": [*common, "document_structure", "document_substantive_body", "document_native_content"],
            "heading_count": int(verification_report.get("heading_count") or 0),
            "body_character_count": body_characters,
            "table_count": int(verification_report.get("table_count") or 0),
        }
    if artifact_type in {"image", "png", "jpg", "jpeg"}:
        width = int(verification_report.get("width") or 0)
        height = int(verification_report.get("height") or 0)
        minimum = int(profile.get("minimum_dimension") or 512)
        maximum = int(profile.get("maximum_dimension") or 4096)
        if not (minimum <= width <= maximum and minimum <= height <= maximum):
            raise OpenCodeExecutionError(
                f"Image dimensions must be between {minimum} and {maximum}; got {width}x{height}."
            )
        non_background_ratio = float(verification_report.get("non_background_ratio") or 0.0)
        minimum_non_background = float(profile.get("minimum_non_background_ratio") or 0.08)
        if non_background_ratio < minimum_non_background:
            raise OpenCodeExecutionError("Image capability produced a visually blank artifact.")
        if profile.get("require_no_text_overflow") and (
            int(quality.get("overflow_count") or 0) != 0 or not generated_checks.get("text_fits_safe_bounds")
        ):
            raise OpenCodeExecutionError("Image capability detected text outside safe canvas bounds.")
        if not generated_checks.get("minimum_contrast_met"):
            raise OpenCodeExecutionError("Image capability did not meet minimum text contrast.")
        return {
            "status": "passed",
            "capability_pack_id": pack.pack_id,
            "capability_pack_version": pack.version,
            "checks": [*common, "image_visual_quality", "image_canvas_bounds", "image_contrast"],
            "width": width,
            "height": height,
            "non_background_ratio": non_background_ratio,
        }

    slide_count = int(verification_report.get("slide_count") or 0)
    non_empty_count = int(verification_report.get("non_empty_slide_count") or slide_count)
    if artifact_type == "pptx" and (slide_count <= 0 or non_empty_count != slide_count):
        raise OpenCodeExecutionError("Presentation capability requires every slide to contain meaningful text.")

    roles = [str(item).strip() for item in artifact.get("slide_roles") or [] if str(item).strip()]
    if artifact_type == "pptx" and len(roles) != slide_count:
        raise OpenCodeExecutionError(
            f"Presentation capability slide role count mismatch: expected {slide_count}, got {len(roles)}."
        )

    maximum_chars = int(profile.get("maximum_text_chars_per_slide") or 700)
    generated_maximum = int(quality.get("maximum_slide_text_chars") or 0)
    verified_maximum = int(
        (verification_report.get("layout_quality") or {}).get("max_slide_text_chars") or 0
    )
    observed_maximum = max(generated_maximum, verified_maximum)
    if observed_maximum > maximum_chars:
        raise OpenCodeExecutionError(
            f"Presentation slide text density exceeds {maximum_chars} characters ({observed_maximum})."
        )

    minimum_layout_kinds = min(
        int(profile.get("minimum_layout_kinds") or 1),
        max(slide_count, 1),
    )
    layout_kind_count = int(quality.get("layout_kind_count") or len(set(roles)))
    if layout_kind_count < minimum_layout_kinds:
        raise OpenCodeExecutionError(
            "Presentation capability requires more layout diversity: "
            f"expected {minimum_layout_kinds}, got {layout_kind_count}."
        )

    require_visual_at = int(profile.get("require_visual_for_slide_count") or 4)
    visual_slide_count = int(quality.get("visual_slide_count") or 0)
    if slide_count >= require_visual_at and visual_slide_count <= 0:
        raise OpenCodeExecutionError(
            f"Presentation capability requires a visual explanation for decks with {require_visual_at}+ slides."
        )

    return {
        "status": "passed",
        "capability_pack_id": pack.pack_id,
        "capability_pack_version": pack.version,
        "checks": [
            *common,
            "presentation_text_density",
            "presentation_layout_diversity",
            "presentation_visual_explanation",
        ],
        "layout_kind_count": layout_kind_count,
        "visual_slide_count": visual_slide_count,
        "maximum_slide_text_chars": observed_maximum,
    }


def verify_artifact_against_request(
    packet: dict[str, Any],
    *,
    verification_report: dict[str, Any],
) -> dict[str, Any]:
    expected_slide_count = requested_slide_count_from_packet(packet)
    if expected_slide_count is None:
        return {"status": "passed", "checks": ["no_explicit_slide_count_requested"]}
    artifact_type = str(verification_report.get("artifact_type") or "").lower()
    if artifact_type != "pptx":
        return {"status": "skipped", "checks": ["explicit_slide_count_not_applicable"], "expected_slide_count": expected_slide_count}
    actual_slide_count = int(verification_report.get("slide_count") or 0)
    if actual_slide_count != expected_slide_count:
        raise OpenCodeExecutionError(
            "PPTX slide count does not match request: "
            f"expected {expected_slide_count}, got {actual_slide_count}."
        )
    return {
        "status": "passed",
        "checks": ["explicit_slide_count_matched"],
        "expected_slide_count": expected_slide_count,
        "actual_slide_count": actual_slide_count,
    }


def requested_slide_count_from_packet(packet: dict[str, Any]) -> int | None:
    plan_input = packet.get("plan_input") if isinstance(packet.get("plan_input"), dict) else {}
    contract_candidates = [
        packet.get("requirements_contract"),
        plan_input.get("requirements_contract"),
        (packet.get("context") or {}).get("requirements_contract") if isinstance(packet.get("context"), dict) else None,
        (packet.get("task_route") or {}).get("requirements_contract") if isinstance(packet.get("task_route"), dict) else None,
        (packet.get("route_decision") or {}).get("requirements_contract") if isinstance(packet.get("route_decision"), dict) else None,
    ]
    for contract in contract_candidates:
        if not isinstance(contract, dict):
            continue
        raw_count = contract.get("page_count") or contract.get("slide_count") or contract.get("pages")
        if raw_count is None:
            continue
        try:
            return max(1, min(int(raw_count), 20))
        except (TypeError, ValueError):
            count = requested_slide_count_from_text(str(raw_count))
            if count is not None:
                return count
    candidates = [
        plan_input.get("user_request"),
        plan_input.get("original_goal"),
        packet.get("original_goal"),
        packet.get("original_goal_summary"),
        packet.get("user_request"),
        packet.get("goal"),
    ]
    for candidate in candidates:
        count = requested_slide_count_from_text(str(candidate or ""))
        if count is not None:
            return count
    return None


def requested_slide_count_from_text(text: str) -> int | None:
    clean = normalize_artifact_text(str(text or "").replace("_", " "))
    match = re.search(r"(?i)(?<!\d)(\d{1,2})\s*(?:页|張|张|p|pages?|slides?|slide|deck pages?)", clean)
    if match:
        return max(1, min(int(match.group(1)), 20))
    chinese_digits = {
        "一": 1,
        "二": 2,
        "两": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
    }
    match = re.search(r"([一二两三四五六七八九十])\s*(?:页|張|张)", clean)
    if match:
        return chinese_digits[match.group(1)]
    return None


def stage_evidence_attachments(packet: dict[str, Any], workspace: Path) -> dict[str, Any]:
    attachment_dir = workspace / "evidence_attachments"
    staged_count = 0
    skipped: list[dict[str, str]] = []
    for attachment in iter_packet_attachments(packet):
        local_path = str(attachment.get("local_path") or "").strip()
        if not local_path:
            continue
        source_path = Path(local_path).expanduser()
        if not source_path.exists() or not source_path.is_file():
            skipped.append({"local_path": local_path, "reason": "missing_file"})
            continue
        attachment_dir.mkdir(parents=True, exist_ok=True)
        filename = safe_staged_attachment_filename(
            str(attachment.get("filename") or source_path.name or "attachment")
        )
        target_path = (attachment_dir / f"{staged_count + 1:03d}_{filename}").resolve()
        target_path.relative_to(workspace.resolve())
        shutil.copy2(source_path, target_path)
        attachment["staged_path"] = str(target_path)
        attachment["staged_filename"] = target_path.name
        staged_count += 1
    return {
        "status": "passed",
        "staged_count": staged_count,
        "skipped": skipped,
    }


def iter_packet_attachments(packet: dict[str, Any]):
    for evidence_pack in iter_packet_evidence_packs(packet):
        for item in evidence_pack.get("items") or []:
            if not isinstance(item, dict):
                continue
            for attachment in item.get("attachments") or []:
                if isinstance(attachment, dict):
                    yield attachment


def iter_packet_evidence_packs(packet: dict[str, Any]):
    context = packet.get("context")
    if isinstance(context, dict) and isinstance(context.get("evidence_pack"), dict):
        yield context["evidence_pack"]
    plan_input = packet.get("plan_input")
    if isinstance(plan_input, dict) and isinstance(plan_input.get("evidence_pack"), dict):
        yield plan_input["evidence_pack"]


def safe_staged_attachment_filename(filename: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "_", str(filename or "").strip())
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    return (cleaned or "attachment")[:140]


def verify_artifact_traceability(
    artifact: dict[str, Any],
    *,
    verification_report: dict[str, Any] | None = None,
    source_evidence_text_by_id: dict[str, str] | None = None,
) -> dict[str, Any]:
    source_ids = [str(item) for item in artifact.get("source_evidence_ids") or [] if str(item).strip()]
    mapping = artifact.get("evidence_to_content_map") or {}
    if not source_ids:
        return {
            "status": "passed",
            "checks": ["no_source_evidence_required"],
            "source_evidence_ids": [],
            "mapped_source_evidence_ids": [],
        }
    if not isinstance(mapping, dict):
        raise OpenCodeExecutionError("Artifact manifest evidence_to_content_map must be an object.")
    missing = [
        source_id
        for source_id in source_ids
        if not evidence_mapping_value_present(mapping.get(source_id))
    ]
    if missing:
        raise OpenCodeExecutionError(
            "Artifact manifest evidence_to_content_map is missing entries for source evidence: "
            + ", ".join(missing)
        )
    grounding_report = verify_artifact_grounding_against_evidence(
        source_ids=source_ids,
        source_evidence_text_by_id=source_evidence_text_by_id or {},
        verification_report=verification_report or {},
    )
    return {
        "status": "passed",
        "checks": ["source_evidence_mapped", *grounding_report.get("checks", [])],
        "source_evidence_ids": source_ids,
        "mapped_source_evidence_ids": source_ids,
        "grounding": grounding_report,
    }


def evidence_mapping_value_present(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, list):
        return any(evidence_mapping_value_present(item) for item in value)
    if isinstance(value, dict):
        return any(evidence_mapping_value_present(item) for item in value.values())
    return True


def verify_artifact_grounding_against_evidence(
    *,
    source_ids: list[str],
    source_evidence_text_by_id: dict[str, str],
    verification_report: dict[str, Any],
) -> dict[str, Any]:
    artifact_text = grounding_text_from_verification_report(verification_report)
    if not artifact_text:
        return {"status": "skipped", "checks": ["no_artifact_text_for_grounding"]}
    evidence_text = "\n".join(
        str(source_evidence_text_by_id.get(source_id) or "") for source_id in source_ids
    ).strip()
    if not evidence_text:
        return {
            "status": "skipped",
            "checks": ["source_evidence_text_unavailable"],
            "source_evidence_ids": source_ids,
        }
    artifact_keywords = grounding_keywords(artifact_text)
    evidence_keywords = grounding_keywords(evidence_text)
    if len(artifact_keywords) < 3 or len(evidence_keywords) < 3:
        return {
            "status": "skipped",
            "checks": ["insufficient_keywords_for_grounding"],
            "artifact_keyword_count": len(artifact_keywords),
            "evidence_keyword_count": len(evidence_keywords),
        }
    overlap = sorted(artifact_keywords & evidence_keywords)
    required_overlap = min(5, max(2, int(len(artifact_keywords) * 0.08)))
    if len(overlap) < required_overlap:
        raise OpenCodeExecutionError(
            "Artifact content is not sufficiently grounded in source evidence; "
            f"keyword_overlap={len(overlap)}, required={required_overlap}."
        )
    return {
        "status": "passed",
        "checks": ["artifact_text_keyword_grounding"],
        "artifact_keyword_count": len(artifact_keywords),
        "evidence_keyword_count": len(evidence_keywords),
        "keyword_overlap_count": len(overlap),
        "keyword_overlap_samples": overlap[:12],
        "required_overlap": required_overlap,
    }


def source_evidence_text_by_id_from_packet(packet: dict[str, Any]) -> dict[str, str]:
    text_by_id: dict[str, list[str]] = {}
    seen_pack_ids: set[int] = set()
    for evidence_pack in iter_packet_evidence_packs(packet):
        pack_identity = id(evidence_pack)
        if pack_identity in seen_pack_ids:
            continue
        seen_pack_ids.add(pack_identity)
        for item in evidence_pack.get("items") or []:
            if not isinstance(item, dict):
                continue
            evidence_id = str(
                item.get("evidence_id")
                or item.get("event_id")
                or item.get("source_event_id")
                or item.get("id")
                or ""
            ).strip()
            if not evidence_id:
                continue
            parts: list[str] = []
            for key in ("content", "text", "excerpt", "summary", "title"):
                value = item.get(key)
                if isinstance(value, str) and value.strip():
                    parts.append(value.strip())
            if parts:
                text_by_id.setdefault(evidence_id, []).extend(parts)
    return {
        evidence_id: "\n".join(dict.fromkeys(parts))
        for evidence_id, parts in text_by_id.items()
        if parts
    }


def grounding_text_from_verification_report(report: dict[str, Any]) -> str:
    parts: list[str] = []
    for key in ("title_candidates", "text_samples"):
        value = report.get(key)
        if isinstance(value, list):
            parts.extend(str(item) for item in value if str(item).strip())
    return normalize_artifact_text("\n".join(parts))


def flatten_mapping_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join(flatten_mapping_text(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(flatten_mapping_text(item) for item in value.values())
    return str(value)


def grounding_keywords(text: str) -> set[str]:
    lowered = normalize_artifact_text(text).lower()
    stopwords = {
        "the",
        "and",
        "for",
        "with",
        "this",
        "that",
        "from",
        "into",
        "about",
        "一个",
        "这个",
        "那个",
        "需要",
        "可以",
        "用户",
        "内容",
        "信息",
    }
    keywords = {
        token
        for token in re.findall(r"[a-z0-9][a-z0-9_.-]{1,}", lowered)
        if token not in stopwords and len(token) >= 2
    }
    for chunk in re.findall(r"[\u4e00-\u9fff]{2,}", lowered):
        if chunk in stopwords:
            continue
        if 2 <= len(chunk) <= 12:
            keywords.add(chunk)
        for size in (2, 3):
            for index in range(0, max(len(chunk) - size + 1, 0)):
                token = chunk[index : index + size]
                if token not in stopwords:
                    keywords.add(token)
    return keywords


def verify_office_archive(path: Path, *, artifact_type: str, required_members: list[str]) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            bad_file = archive.testzip()
    except zipfile.BadZipFile as exc:
        raise OpenCodeExecutionError(f"Artifact is not a valid {artifact_type.upper()} Office archive.") from exc
    if bad_file:
        raise OpenCodeExecutionError(f"Artifact is not a valid {artifact_type.upper()} archive; bad member: {bad_file}.")
    missing = [member for member in required_members if member not in names]
    if missing:
        raise OpenCodeExecutionError(
            f"Artifact is not a valid {artifact_type.upper()} file; missing members: {', '.join(missing)}."
        )
    return {
        "status": "passed",
        "artifact_type": artifact_type,
        "checks": ["zip_archive", "office_required_members"],
        "required_members": required_members,
        "file_size": path.stat().st_size,
    }


def verify_xlsx_content(path: Path) -> dict[str, Any]:
    archive_report = verify_office_archive(
        path,
        artifact_type="xlsx",
        required_members=["[Content_Types].xml", "xl/workbook.xml"],
    )
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise OpenCodeExecutionError("Artifact is not readable as an XLSX workbook.") from exc
    sheet_names = list(workbook.sheetnames)
    if not sheet_names:
        raise OpenCodeExecutionError("XLSX artifact has no worksheets.")
    formula_count = 0
    data_row_count = 0
    text_samples: list[str] = []
    frozen_sheet_count = 0
    filtered_sheet_count = 0
    for sheet in workbook.worksheets:
        if sheet.freeze_panes:
            frozen_sheet_count += 1
        if str(sheet.auto_filter.ref or "").strip():
            filtered_sheet_count += 1
        for row_index, row in enumerate(sheet.iter_rows(values_only=False), start=1):
            values = [cell.value for cell in row]
            if any(value not in (None, "") for value in values):
                sample = " | ".join(str(value) for value in values if value not in (None, ""))
                if sample and len(text_samples) < 12:
                    text_samples.append(f"{sheet.title}: {sample}")
                if row_index > 1:
                    data_row_count += 1
            formula_count += sum(1 for value in values if isinstance(value, str) and value.startswith("="))
    from .spreadsheet_capability import semantic_number_format_issues

    format_issues = semantic_number_format_issues(workbook)
    workbook.close()
    if not text_samples:
        raise OpenCodeExecutionError("XLSX artifact has no extractable cell content.")
    return {
        **archive_report,
        "checks": [*archive_report["checks"], "xlsx_readable_workbook", "xlsx_non_empty_cells"],
        "sheet_count": len(sheet_names),
        "sheet_names": sheet_names,
        "data_row_count": data_row_count,
        "formula_count": formula_count,
        "frozen_sheet_count": frozen_sheet_count,
        "filtered_sheet_count": filtered_sheet_count,
        "semantic_number_format_issues": format_issues,
        "text_samples": text_samples,
    }


def verify_docx_content(path: Path) -> dict[str, Any]:
    archive_report = verify_office_archive(
        path,
        artifact_type="docx",
        required_members=["[Content_Types].xml", "word/document.xml"],
    )
    try:
        from docx import Document

        document = Document(str(path))
    except Exception as exc:
        raise OpenCodeExecutionError("Artifact is not readable as a DOCX document.") from exc
    paragraphs = [normalize_artifact_text(paragraph.text) for paragraph in document.paragraphs]
    paragraphs = [text for text in paragraphs if text]
    heading_candidates = [
        normalize_artifact_text(paragraph.text)
        for paragraph in document.paragraphs
        if normalize_artifact_text(paragraph.text)
        and str(getattr(paragraph.style, "name", "")).lower().startswith("heading ")
    ]
    table_texts: list[str] = []
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(normalize_artifact_text(cell.text) for cell in row.cells if normalize_artifact_text(cell.text))
            if row_text:
                table_texts.append(row_text)
    joined = normalize_artifact_text("\n".join([*paragraphs, *table_texts]))
    if not joined:
        raise OpenCodeExecutionError("DOCX artifact has no extractable text.")
    return {
        **archive_report,
        "checks": [*archive_report["checks"], "docx_readable_document", "docx_non_empty_text"],
        "paragraph_count": len(paragraphs),
        "heading_count": len(heading_candidates),
        "heading_candidates": heading_candidates[:20],
        "table_count": len(document.tables),
        "body_character_count": len(re.sub(r"\s+", "", joined)),
        "text_samples": [*paragraphs[:5], *table_texts[:5]],
    }


def verify_image_content(path: Path, *, artifact_type: str) -> dict[str, Any]:
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.load()
            width, height = image.size
            mode = image.mode
            probe = image.convert("RGB")
            probe.thumbnail((240, 240))
            colors = probe.getcolors(maxcolors=probe.width * probe.height) or []
            background = probe.getpixel((0, 0))
            background_pixels = sum(count for count, color in colors if color == background)
            total_pixels = max(probe.width * probe.height, 1)
            non_background_ratio = 1.0 - background_pixels / total_pixels
            image_format = str(image.format or "").upper()
    except Exception as exc:
        raise OpenCodeExecutionError("Artifact is not readable as an image.") from exc
    if width <= 0 or height <= 0:
        raise OpenCodeExecutionError("Image artifact has invalid dimensions.")
    if len(colors) <= 1 or non_background_ratio <= 0.001:
        raise OpenCodeExecutionError("Image artifact is visually blank.")
    return {
        "status": "passed",
        "artifact_type": artifact_type,
        "checks": ["image_readable", "image_dimensions", "image_non_blank_visual"],
        "file_size": path.stat().st_size,
        "format": image_format,
        "mode": mode,
        "width": width,
        "height": height,
        "distinct_color_count": len(colors),
        "non_background_ratio": round(non_background_ratio, 4),
    }


def verify_pptx_content(path: Path) -> dict[str, Any]:
    archive_report = verify_office_archive(
        path,
        artifact_type="pptx",
        required_members=["[Content_Types].xml", "ppt/presentation.xml"],
    )
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception as exc:
        raise OpenCodeExecutionError("Artifact is not readable as a PPTX presentation.") from exc

    slide_texts: list[str] = []
    slide_text_char_counts: list[int] = []
    slide_visible_text_char_counts: list[int] = []
    title_candidates: list[str] = []
    for slide in presentation.slides:
        title_shape = getattr(slide.shapes, "title", None)
        title = ""
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title = normalize_artifact_text(title_shape.text_frame.text)
        parts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = normalize_artifact_text(shape.text_frame.text)
            if text:
                parts.append(text)
        if not title and parts:
            title = parts[0]
        if title:
            title_candidates.append(title)
        slide_text = normalize_artifact_text("\n".join(parts))
        if slide_text:
            slide_texts.append(slide_text)
            slide_text_char_counts.append(len(re.sub(r"\s+", "", slide_text)))
            slide_visible_text_char_counts.append(len(slide_text))

    joined_text = normalize_artifact_text("\n".join(slide_texts))
    if not joined_text:
        raise OpenCodeExecutionError("PPTX artifact has no extractable text.")

    return {
        **archive_report,
        "checks": [
            *archive_report["checks"],
            "pptx_slide_count",
            "pptx_non_empty_text",
        ],
        "slide_count": len(presentation.slides),
        "text_character_count": len(joined_text),
        "non_empty_slide_count": len(slide_texts),
        "empty_slide_count": max(len(presentation.slides) - len(slide_texts), 0),
        "title_candidates": title_candidates[:10],
        "text_samples": slide_texts[:5],
        "layout_quality": {
            "max_slide_text_chars": max(slide_text_char_counts) if slide_text_char_counts else 0,
            "max_slide_visible_text_chars": (
                max(slide_visible_text_char_counts) if slide_visible_text_char_counts else 0
            ),
            "average_slide_text_chars": int(sum(slide_text_char_counts) / len(slide_text_char_counts))
            if slide_text_char_counts
            else 0,
            "overfull_slide_count": sum(1 for count in slide_text_char_counts if count > 900),
            "overfull_threshold_chars": 900,
        },
    }


def normalize_artifact_text(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def safe_artifact_filename(filename: str, artifact_type: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "_", str(filename or "").strip())
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    extension = extension_for_artifact(artifact_type)
    if not cleaned:
        cleaned = f"nomi_artifact{extension}"
    if extension and not cleaned.lower().endswith(extension):
        cleaned = f"{cleaned}{extension}"
    return cleaned[:180]


def infer_artifact_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return {
        ".pptx": "pptx",
        ".docx": "docx",
        ".xlsx": "xlsx",
        ".md": "markdown",
        ".markdown": "markdown",
        ".png": "png",
        ".jpg": "jpg",
        ".jpeg": "jpeg",
    }.get(suffix, "artifact")


def extension_for_artifact(artifact_type: str) -> str:
    return {
        "pptx": ".pptx",
        "docx": ".docx",
        "xlsx": ".xlsx",
        "markdown": ".md",
        "image": ".png",
        "png": ".png",
        "jpg": ".jpg",
        "jpeg": ".jpeg",
    }.get(str(artifact_type or "").lower(), "")


def mime_type_for_artifact(artifact_type: str) -> str:
    return {
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "markdown": "text/markdown",
        "image": "image/png",
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
    }.get(str(artifact_type or "").lower(), "application/octet-stream")
