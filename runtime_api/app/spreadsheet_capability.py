from __future__ import annotations

import json
import math
import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .capability_packs import select_capability_pack
from .font_support import preferred_cjk_font_name
from .office_metadata import office_core_property


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAME_RE = re.compile(r"^[^\\/*?:\[\]]{1,31}$")
FORMULA_FUNCTION_RE = re.compile(r"\b([A-Z][A-Z0-9_.]*)\s*\(", re.IGNORECASE)
ALLOWED_FORMULA_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "IFERROR", "ROUND"}
UNSAFE_FORMULA_FUNCTIONS = {"CALL", "DDE", "EXEC", "FILTERXML", "HYPERLINK", "REGISTER.ID", "RTD", "WEBSERVICE"}
SUPPORTED_COLUMN_TYPES = {"text", "integer", "number", "currency", "percentage", "date", "datetime", "boolean", "formula"}
FORMULA_RESULT_TYPES = {"integer", "number", "currency", "percentage"}
PERCENTAGE_LABEL_RE = re.compile(r"(率|百分比|占比|比例|margin|rate|percent|ratio)", re.IGNORECASE)
NUMERIC_LABEL_RE = re.compile(r"(收入|成本|利润|金额|价格|费用|预算|营收|revenue|cost|profit|amount|price|budget)", re.IGNORECASE)


class SpreadsheetSpecError(ValueError):
    pass


def _pack():
    pack = select_capability_pack("xlsx")
    if pack is None:
        raise SpreadsheetSpecError("spreadsheet capability pack is not installed")
    return pack


def validate_spreadsheet_spec(spec: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise SpreadsheetSpecError("spreadsheet specification must be an object")
    normalized = deepcopy(spec)
    for key in ("title", "purpose"):
        value = str(normalized.get(key) or "").strip()
        if not value:
            raise SpreadsheetSpecError(f"spreadsheet {key} is required")
        normalized[key] = value
    sheets = normalized.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise SpreadsheetSpecError("spreadsheet sheets must be a non-empty list")
    profile = dict(_pack().quality_profile)
    if len(sheets) > int(profile.get("maximum_sheet_count") or 20):
        raise SpreadsheetSpecError("spreadsheet has too many sheets")

    seen_names: set[str] = set()
    total_rows = 0
    for sheet_index, sheet in enumerate(sheets, start=1):
        if not isinstance(sheet, dict):
            raise SpreadsheetSpecError(f"sheet {sheet_index} must be an object")
        name = str(sheet.get("name") or "").strip()
        if not SHEET_NAME_RE.fullmatch(name):
            raise SpreadsheetSpecError(f"sheet {sheet_index} has invalid name")
        if name.casefold() in seen_names:
            raise SpreadsheetSpecError(f"duplicate sheet name: {name}")
        seen_names.add(name.casefold())
        sheet["name"] = name
        columns = sheet.get("columns")
        if not isinstance(columns, list) or not columns:
            raise SpreadsheetSpecError(f"sheet {name} columns must be a non-empty list")
        if len(columns) > int(profile.get("maximum_column_count") or 50):
            raise SpreadsheetSpecError(f"sheet {name} has too many columns")
        seen_keys: set[str] = set()
        for column_index, column in enumerate(columns, start=1):
            if not isinstance(column, dict):
                raise SpreadsheetSpecError(f"sheet {name} column {column_index} must be an object")
            key = str(column.get("key") or "").strip()
            label = str(column.get("label") or "").strip()
            kind = str(column.get("type") or "text").strip().lower()
            if not key or not label:
                raise SpreadsheetSpecError(f"sheet {name} column {column_index} needs key and label")
            if key in seen_keys:
                raise SpreadsheetSpecError(f"sheet {name} has duplicate column key: {key}")
            if kind not in SUPPORTED_COLUMN_TYPES:
                raise SpreadsheetSpecError(f"sheet {name} column {key} has unsupported type: {kind}")
            seen_keys.add(key)
            column["key"] = key
            column["label"] = label
            column["type"] = kind
            if kind == "formula":
                formula = str(column.get("formula") or "").strip()
                if not formula.startswith("="):
                    raise SpreadsheetSpecError(f"sheet {name} formula column {key} needs formula")
                validate_formula(formula)
                result_type = str(column.get("result_type") or "").strip().lower()
                if result_type and result_type not in FORMULA_RESULT_TYPES:
                    raise SpreadsheetSpecError(
                        f"sheet {name} formula column {key} has unsupported result type: {result_type}"
                    )
                if result_type:
                    column["result_type"] = result_type
        rows = sheet.get("rows")
        if not isinstance(rows, list):
            raise SpreadsheetSpecError(f"sheet {name} rows must be a list")
        total_rows += len(rows)
        known_keys = {str(column["key"]) for column in columns}
        for row_index, row in enumerate(rows, start=2):
            if not isinstance(row, dict):
                raise SpreadsheetSpecError(f"sheet {name} row {row_index} must be an object")
            unknown = sorted(set(row) - known_keys)
            if unknown:
                raise SpreadsheetSpecError(f"sheet {name} row {row_index} has unknown fields: {', '.join(unknown)}")
            for value in row.values():
                if isinstance(value, str) and value.startswith("="):
                    validate_formula(value)
        evidence_ids = [str(value).strip() for value in sheet.get("source_evidence_ids") or [] if str(value).strip()]
        sheet["source_evidence_ids"] = list(dict.fromkeys(evidence_ids))
    if total_rows > int(profile.get("maximum_total_rows") or 100000):
        raise SpreadsheetSpecError("spreadsheet has too many total rows")
    return normalized


def validate_formula(formula: str) -> None:
    upper = str(formula or "").upper()
    functions = {match.upper() for match in FORMULA_FUNCTION_RE.findall(upper)}
    unsafe = sorted(functions & UNSAFE_FORMULA_FUNCTIONS)
    if unsafe:
        raise SpreadsheetSpecError(f"unsafe formula function: {unsafe[0]}")
    unknown = sorted(functions - ALLOWED_FORMULA_FUNCTIONS)
    if unknown:
        raise SpreadsheetSpecError(f"unsupported formula function: {unknown[0]}")
    if any(token in upper for token in ("[HTTP:", "[HTTPS:", "'HTTP:", "'HTTPS:", "CMD|", "POWERSHELL")):
        raise SpreadsheetSpecError("unsafe external formula reference")


def render_spreadsheet(
    spec: dict[str, Any],
    output_path: str | Path,
    *,
    manifest_path: str | Path | None = None,
    task_packet: dict[str, Any] | None = None,
) -> dict[str, Any]:
    normalized = validate_spreadsheet_spec(spec)
    pack = _pack()
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise SpreadsheetSpecError("spreadsheet output filename must end in .xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    evidence_map: dict[str, list[str]] = {}
    formula_cells: list[dict[str, Any]] = []
    total_data_rows = 0
    for sheet_spec in normalized["sheets"]:
        sheet = workbook.create_sheet(sheet_spec["name"])
        columns = sheet_spec["columns"]
        rows = sheet_spec["rows"]
        total_data_rows += len(rows)
        for column_index, column in enumerate(columns, start=1):
            cell = sheet.cell(1, column_index, column["label"])
            _style_header(cell)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column in enumerate(columns, start=1):
                key = column["key"]
                value = row.get(key)
                if column["type"] == "formula":
                    value = str(column["formula"]).format(row=row_index)
                cell = sheet.cell(row_index, column_index, value)
                _style_data_cell(cell, column, row_index=row_index)
                if isinstance(value, str) and value.startswith("="):
                    formula_cells.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": value})
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"
        sheet.sheet_view.showGridLines = False
        for column_index, column in enumerate(columns, start=1):
            values = [str(column["label"]), *[str(row.get(column["key"]) or "") for row in rows[:200]]]
            width = min(max(max((len(value) for value in values), default=8) + 3, 11), 36)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for evidence_id in sheet_spec.get("source_evidence_ids") or []:
            evidence_map.setdefault(evidence_id, []).append(f"sheet:{sheet.title}")

    preview_values = calculate_formula_previews(workbook, formula_cells)
    format_issues = semantic_number_format_issues(workbook)
    if format_issues:
        raise SpreadsheetSpecError("spreadsheet business number formats are invalid: " + "; ".join(format_issues))
    workbook.properties.title = office_core_property(normalized["title"])
    workbook.properties.subject = office_core_property(normalized["purpose"])
    workbook.properties.creator = office_core_property("Nomi Spreadsheet Studio")
    workbook.properties.description = office_core_property(
        f"Capability pack {pack.pack_id}@{pack.version}"
    )
    workbook.save(output)

    source_ids = list(evidence_map)
    formula_count = len(formula_cells)
    quality = {
        "status": "passed",
        "checks": {
            "has_sheets": bool(normalized["sheets"]),
            "has_headers": all(bool(sheet["columns"]) for sheet in normalized["sheets"]),
            "formulas_are_safe": True,
            "formula_previews_complete": len(preview_values) == formula_count,
            "filters_and_freeze_panes": True,
            "source_mapping_complete": all(evidence_map.get(source_id) for source_id in source_ids),
            "semantic_number_formats": not format_issues,
        },
        "sheet_count": len(normalized["sheets"]),
        "data_row_count": total_data_rows,
        "formula_count": formula_count,
        "formula_preview_count": len(preview_values),
        "formula_previews": preview_values,
        "format_issues": format_issues,
    }
    if formula_count and len(preview_values) != formula_count:
        raise SpreadsheetSpecError("could not calculate previews for every spreadsheet formula")
    manifest = {
        "artifact_type": "xlsx",
        "filename": output.name,
        "file_path": str(output.resolve()),
        "mime_type": XLSX_MIME,
        "capability_pack_id": pack.pack_id,
        "capability_pack_version": pack.version,
        "title": normalized["title"],
        "sheet_names": [sheet["name"] for sheet in normalized["sheets"]],
        "source_evidence_ids": source_ids,
        "evidence_to_content_map": evidence_map,
        "quality_report": quality,
        "task_contract_present": bool(task_packet),
    }
    if manifest_path is not None:
        path = Path(manifest_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def _style_header(cell) -> None:
    cell.font = Font(name=preferred_cjk_font_name(), size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1D2939")
    cell.alignment = Alignment(vertical="center", horizontal="left")
    cell.border = Border(bottom=Side(style="medium", color="10B981"))


def _style_data_cell(cell, column: dict[str, Any], *, row_index: int) -> None:
    kind = str(column.get("type") or "text")
    if kind == "formula":
        kind = str(column.get("result_type") or classify_business_number_format(column.get("label")) or "number")
    cell.font = Font(name=preferred_cjk_font_name(), size=10, color="172033")
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if row_index % 2 == 1:
        cell.fill = PatternFill("solid", fgColor="F3F6F8")
    cell.border = Border(bottom=Side(style="thin", color="D8DEE8"))
    if kind == "currency":
        cell.number_format = '¥#,##0.00;[Red]-¥#,##0.00'
    elif kind == "percentage":
        cell.number_format = "0.0%"
    elif kind == "integer":
        cell.number_format = "0"
    elif kind == "number":
        cell.number_format = "0.00"
    elif kind == "date":
        cell.number_format = "yyyy-mm-dd"
    elif kind == "datetime":
        cell.number_format = "yyyy-mm-dd hh:mm"


def classify_business_number_format(label: Any) -> str | None:
    text = str(label or "").strip()
    if not text:
        return None
    if PERCENTAGE_LABEL_RE.search(text):
        return "percentage"
    if NUMERIC_LABEL_RE.search(text):
        return "currency"
    return None


def semantic_number_format_issues(workbook) -> list[str]:
    issues: list[str] = []
    for sheet in workbook.worksheets:
        for column_index in range(1, sheet.max_column + 1):
            label = str(sheet.cell(1, column_index).value or "").strip()
            semantic_type = classify_business_number_format(label)
            if semantic_type is None:
                continue
            for row_index in range(2, sheet.max_row + 1):
                cell = sheet.cell(row_index, column_index)
                if cell.value in (None, ""):
                    continue
                number_format = str(cell.number_format or "General")
                if semantic_type == "percentage" and "%" not in number_format:
                    issues.append(
                        f"{sheet.title}!{cell.coordinate} '{label}' is not formatted as a percentage"
                    )
                elif semantic_type == "currency" and number_format in {"General", "@"}:
                    issues.append(
                        f"{sheet.title}!{cell.coordinate} '{label}' uses General instead of a numeric format"
                    )
    return issues


def calculate_formula_previews(workbook: Workbook, formula_cells: list[dict[str, Any]]) -> list[dict[str, Any]]:
    cache: dict[tuple[str, str], float] = {}

    def value_for(sheet_name: str, coordinate: str, stack: set[tuple[str, str]]) -> float:
        key = (sheet_name, coordinate)
        if key in cache:
            return cache[key]
        if key in stack:
            raise SpreadsheetSpecError(f"circular formula reference: {sheet_name}!{coordinate}")
        stack.add(key)
        raw = workbook[sheet_name][coordinate].value
        if isinstance(raw, bool):
            result = float(raw)
        elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
            result = float(raw)
        elif isinstance(raw, str) and raw.startswith("="):
            result = evaluate_formula(raw, workbook=workbook, current_sheet=sheet_name, value_for=value_for, stack=stack)
        else:
            result = 0.0
        stack.remove(key)
        cache[key] = result
        return result

    previews: list[dict[str, Any]] = []
    for item in formula_cells:
        value = value_for(item["sheet"], item["cell"], set())
        previews.append({**item, "calculated_value": round(value, 8)})
    return previews


def evaluate_formula(formula: str, *, workbook: Workbook, current_sheet: str, value_for, stack: set[tuple[str, str]]) -> float:
    expression = formula[1:].strip()
    if expression.upper().startswith("IFERROR(") and expression.endswith(")"):
        inside = expression[8:-1]
        primary, fallback = _split_args(inside)
        try:
            return _evaluate_arithmetic(primary, current_sheet=current_sheet, value_for=value_for, stack=stack)
        except (ZeroDivisionError, SpreadsheetSpecError, ValueError):
            return float(fallback or 0)
    sum_match = re.fullmatch(r"SUM\((?:(?:'([^']+)'|([^'!]+))!)?([A-Z]+\d+):([A-Z]+\d+)\)", expression, re.IGNORECASE)
    if sum_match:
        sheet_name = (sum_match.group(1) or sum_match.group(2) or current_sheet).strip()
        total = 0.0
        for row in workbook[sheet_name][f"{sum_match.group(3)}:{sum_match.group(4)}"]:
            for cell in row:
                total += value_for(sheet_name, cell.coordinate, stack)
        return total
    return _evaluate_arithmetic(expression, current_sheet=current_sheet, value_for=value_for, stack=stack)


def _split_args(value: str) -> tuple[str, str]:
    depth = 0
    for index, char in enumerate(value):
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif char == "," and depth == 0:
            return value[:index], value[index + 1 :]
    raise SpreadsheetSpecError("formula function arguments are invalid")


def _evaluate_arithmetic(expression: str, *, current_sheet: str, value_for, stack: set[tuple[str, str]]) -> float:
    token_re = re.compile(r"(?:(?:'([^']+)'|([\w\u4e00-\u9fff -]+))!)?([A-Z]+\d+)", re.IGNORECASE)

    def replace(match: re.Match[str]) -> str:
        sheet_name = (match.group(1) or match.group(2) or current_sheet).strip()
        return str(value_for(sheet_name, match.group(3).upper(), stack))

    substituted = token_re.sub(replace, expression)
    if not re.fullmatch(r"[0-9eE+\-*/().\s]+", substituted):
        raise SpreadsheetSpecError(f"unsupported formula expression: {expression}")
    return float(eval(substituted, {"__builtins__": {}}, {}))
