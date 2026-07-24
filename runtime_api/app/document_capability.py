from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .capability_packs import select_capability_pack
from .font_support import preferred_cjk_font_name
from .office_metadata import office_core_property


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_SECTION_KINDS = {"executive_summary", "analysis", "narrative", "table", "recommendations", "appendix"}
PLACEHOLDER_RE = re.compile(r"(?i)\b(?:todo|tbd|lorem ipsum|placeholder)\b|待补充|此处填写|示例内容|测试内容")
PERCENT_METRIC_RE = re.compile(r"(?<![\d.])\d+(?:\.\d+)?%")
TARGET_CUE_RE = re.compile(
    r"(?i)(?:确保|维持|保持|目标|不得低于|至少达到|控制在|"
    r"\b(?:target|maintain|keep|ensure|at\s+least|no\s+less\s+than)\b)"
)
CAUSAL_INFERENCE_RE = re.compile(
    r"可能与|可能由于|可能因为|原因是|源于|归因于|导致|因此|反映了|反映出"
)


class DocumentSpecError(ValueError):
    pass


def _pack():
    pack = select_capability_pack("docx")
    if pack is None:
        raise DocumentSpecError("document capability pack is not installed")
    return pack


def validate_document_spec(spec: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise DocumentSpecError("document specification must be an object")
    normalized = deepcopy(spec)
    for key in ("title", "audience", "purpose"):
        value = str(normalized.get(key) or "").strip()
        if not value:
            raise DocumentSpecError(f"document {key} is required")
        normalized[key] = value
    sections = normalized.get("sections")
    if not isinstance(sections, list) or not sections:
        raise DocumentSpecError("document sections must be a non-empty list")
    profile = dict(_pack().quality_profile)
    if len(sections) > int(profile.get("maximum_section_count") or 40):
        raise DocumentSpecError("document has too many sections")
    seen_headings: set[str] = set()
    previous_level = 1
    body_chars = 0
    for index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            raise DocumentSpecError(f"section {index} must be an object")
        kind = str(section.get("kind") or "narrative").strip().lower()
        heading = str(section.get("heading") or "").strip()
        level = int(section.get("level") or 1)
        if kind not in SUPPORTED_SECTION_KINDS:
            raise DocumentSpecError(f"section {index} has unsupported kind: {kind}")
        if not heading:
            raise DocumentSpecError(f"section {index} heading is required")
        if heading.casefold() in seen_headings:
            raise DocumentSpecError(f"duplicate section heading: {heading}")
        if level < 1 or level > 3:
            raise DocumentSpecError(f"section {heading} heading level must be 1-3")
        if index > 1 and level > previous_level + 1:
            raise DocumentSpecError(f"section {heading} skips heading hierarchy")
        seen_headings.add(heading.casefold())
        previous_level = level
        section["kind"] = kind
        section["heading"] = heading
        section["level"] = level
        content_strings: list[str] = []
        for field in ("paragraphs", "bullets", "numbered_items"):
            values = section.get(field) or []
            if not isinstance(values, list):
                raise DocumentSpecError(f"section {heading} {field} must be a list")
            if any(not isinstance(value, str) for value in values):
                raise DocumentSpecError(f"section {heading} {field} items must be strings")
            cleaned = [value.strip() for value in values if value.strip()]
            section[field] = cleaned
            content_strings.extend(cleaned)
        table = section.get("table")
        if table is not None:
            if not isinstance(table, dict):
                raise DocumentSpecError(f"section {heading} table must be an object")
            headers = [str(value).strip() for value in table.get("headers") or []]
            rows = table.get("rows") or []
            if not headers or not isinstance(rows, list):
                raise DocumentSpecError(f"section {heading} table needs headers and rows")
            cleaned_rows: list[list[str]] = []
            for row_index, row in enumerate(rows, start=1):
                if not isinstance(row, list) or len(row) != len(headers):
                    raise DocumentSpecError(f"section {heading} table row width mismatch at row {row_index}")
                cleaned = [str(value).strip() for value in row]
                cleaned_rows.append(cleaned)
                content_strings.extend(cleaned)
            table["headers"] = headers
            table["rows"] = cleaned_rows
            content_strings.extend(headers)
        if not content_strings:
            raise DocumentSpecError(f"section {heading} has no content")
        joined = "\n".join(content_strings)
        if PLACEHOLDER_RE.search(joined):
            raise DocumentSpecError(f"section {heading} contains placeholder content")
        body_chars += len(re.sub(r"\s+", "", joined))
        evidence_ids = [str(value).strip() for value in section.get("source_evidence_ids") or [] if str(value).strip()]
        section["source_evidence_ids"] = list(dict.fromkeys(evidence_ids))
    if body_chars < int(profile.get("minimum_body_characters") or 120):
        raise DocumentSpecError("document body is too short to satisfy the requested report")
    normalized["body_character_count"] = body_chars
    return normalized


def render_document(
    spec: dict[str, Any],
    output_path: str | Path,
    *,
    manifest_path: str | Path | None = None,
    task_packet: dict[str, Any] | None = None,
) -> dict[str, Any]:
    normalized = validate_document_spec(spec)
    _validate_document_request_contract(normalized, task_packet=task_packet)
    _validate_recommendation_metrics(normalized, task_packet=task_packet)
    _validate_causal_inferences(normalized, task_packet=task_packet)
    pack = _pack()
    output = Path(output_path)
    if output.suffix.lower() != ".docx":
        raise DocumentSpecError("document output filename must end in .docx")
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document, normalized)
    evidence_map: dict[str, list[str]] = {}
    action_item_count = 0
    table_count = 0
    for section in normalized["sections"]:
        heading = document.add_heading(section["heading"], level=section["level"])
        heading.paragraph_format.keep_with_next = True
        for paragraph_text in section.get("paragraphs") or []:
            paragraph = document.add_paragraph(paragraph_text, style="Body Text")
            paragraph.paragraph_format.space_after = Pt(7)
        for bullet in section.get("bullets") or []:
            document.add_paragraph(bullet, style="List Bullet")
        for number, item in enumerate(section.get("numbered_items") or [], start=1):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(item)
            action_item_count += 1
        table_spec = section.get("table")
        if isinstance(table_spec, dict):
            table_count += 1
            table = document.add_table(rows=1, cols=len(table_spec["headers"]))
            table.style = "Light Shading Accent 1"
            table.autofit = True
            for index, header in enumerate(table_spec["headers"]):
                cell = table.rows[0].cells[index]
                cell.text = header
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
            for values in table_spec["rows"]:
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    cells[index].text = value
                    cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            document.add_paragraph()
        source_ids = section.get("source_evidence_ids") or []
        if source_ids:
            source_note = document.add_paragraph("来源：" + "、".join(source_ids), style="Caption")
            source_note.paragraph_format.space_after = Pt(9)
        for evidence_id in source_ids:
            evidence_map.setdefault(evidence_id, []).append(f"section:{section['heading']}")

    core = document.core_properties
    core.title = office_core_property(normalized["title"])
    core.subject = office_core_property(normalized["purpose"])
    core.author = office_core_property(normalized.get("author") or "Nomi")
    core.comments = office_core_property(f"Capability pack {pack.pack_id}@{pack.version}")
    _apply_direct_cjk_fonts(document, preferred_cjk_font_name())
    document.save(output)

    section_headings = [section["heading"] for section in normalized["sections"]]
    source_ids = list(evidence_map)
    quality = {
        "status": "passed",
        "checks": {
            "heading_hierarchy_valid": True,
            "body_is_substantive": normalized["body_character_count"] >= int(pack.quality_profile.get("minimum_body_characters") or 120),
            "no_placeholder_content": True,
            "source_mapping_complete": all(evidence_map.get(source_id) for source_id in source_ids),
            "tables_are_rectangular": True,
            "editable_native_content": True,
        },
        "section_count": len(section_headings),
        "body_character_count": normalized["body_character_count"],
        "table_count": table_count,
        "action_item_count": action_item_count,
    }
    manifest = {
        "artifact_type": "docx",
        "filename": output.name,
        "file_path": str(output.resolve()),
        "mime_type": DOCX_MIME,
        "capability_pack_id": pack.pack_id,
        "capability_pack_version": pack.version,
        "title": normalized["title"],
        "section_headings": section_headings,
        "source_evidence_ids": source_ids,
        "evidence_to_content_map": evidence_map,
        "quality_report": quality,
        "task_contract_present": bool(task_packet),
    }
    if manifest_path is not None:
        path = Path(manifest_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def _validate_document_request_contract(
    spec: dict[str, Any],
    *,
    task_packet: dict[str, Any] | None,
) -> None:
    contract = _task_requirements_contract(task_packet)
    required_sections = [
        str(value).strip()
        for value in contract.get("sections") or []
        if str(value).strip()
    ]
    if not required_sections:
        return
    actual_headings = [
        re.sub(r"\s+", "", str(section.get("heading") or "")).casefold()
        for section in spec.get("sections") or []
        if isinstance(section, dict)
    ]
    for required in required_sections:
        normalized_required = re.sub(r"\s+", "", required).casefold()
        if not any(
            normalized_required in heading or heading in normalized_required
            for heading in actual_headings
        ):
            raise DocumentSpecError(f"missing required section: {required}")


def _task_requirements_contract(task_packet: dict[str, Any] | None) -> dict[str, Any]:
    if not isinstance(task_packet, dict):
        return {}
    plan_input = (
        task_packet.get("plan_input")
        if isinstance(task_packet.get("plan_input"), dict)
        else {}
    )
    for candidate in (
        task_packet.get("requirements_contract"),
        plan_input.get("requirements_contract"),
    ):
        if isinstance(candidate, dict):
            return dict(candidate)
    return {}


def _validate_recommendation_metrics(
    spec: dict[str, Any],
    *,
    task_packet: dict[str, Any] | None,
) -> None:
    evidence_texts = _task_evidence_texts(task_packet)
    if not evidence_texts:
        return
    for section in spec.get("sections") or []:
        if section.get("kind") != "recommendations":
            continue
        recommendation_texts = [
            *(section.get("paragraphs") or []),
            *(section.get("bullets") or []),
            *(section.get("numbered_items") or []),
        ]
        for recommendation in recommendation_texts:
            if not TARGET_CUE_RE.search(recommendation):
                continue
            for metric in PERCENT_METRIC_RE.findall(recommendation):
                supporting = [text for text in evidence_texts if metric in text]
                if supporting and not any(TARGET_CUE_RE.search(text) for text in supporting):
                    raise DocumentSpecError(f"observed metric {metric} is recast as a target")


def _validate_causal_inferences(
    spec: dict[str, Any],
    *,
    task_packet: dict[str, Any] | None,
) -> None:
    evidence = _task_evidence_by_id(task_packet)
    if not evidence:
        return
    for section in spec.get("sections") or []:
        if section.get("kind") == "recommendations":
            continue
        source_ids = section.get("source_evidence_ids") or []
        unknown = [source_id for source_id in source_ids if source_id not in evidence]
        if unknown:
            raise DocumentSpecError(
                f"section {section.get('heading') or '?'} has unknown evidence id: {unknown[0]}"
            )
        supporting_text = "\n".join(evidence[source_id] for source_id in source_ids)
        if not supporting_text:
            continue
        content = [
            *(section.get("paragraphs") or []),
            *(section.get("bullets") or []),
            *(section.get("numbered_items") or []),
        ]
        for text in content:
            for match in CAUSAL_INFERENCE_RE.finditer(text):
                cue = match.group(0)
                if cue not in supporting_text:
                    raise DocumentSpecError(
                        "unsupported causal inference in section "
                        f"{section.get('heading') or '?'}: {cue}"
                    )


def _task_evidence_texts(task_packet: dict[str, Any] | None) -> list[str]:
    return list(_task_evidence_by_id(task_packet).values())


def _task_evidence_by_id(task_packet: dict[str, Any] | None) -> dict[str, str]:
    if not isinstance(task_packet, dict):
        return {}
    packs: list[Any] = []
    context = task_packet.get("context")
    if isinstance(context, dict):
        packs.append(context.get("evidence_pack"))
    plan_input = task_packet.get("plan_input")
    if isinstance(plan_input, dict):
        packs.append(plan_input.get("evidence_pack"))
    evidence: dict[str, str] = {}
    for pack in packs:
        if not isinstance(pack, dict):
            continue
        for item in pack.get("items") or []:
            if not isinstance(item, dict):
                continue
            evidence_id = str(item.get("evidence_id") or item.get("source_id") or "").strip()
            text = str(item.get("content") or item.get("excerpt") or "").strip()
            if evidence_id and text:
                evidence[evidence_id] = text
    return evidence


def _configure_document(document: Document, spec: dict[str, Any]) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    styles = document.styles
    _configure_east_asian_language(document)
    cjk_font = preferred_cjk_font_name()
    styles["Normal"].font.name = cjk_font
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = cjk_font
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(23, 32, 51)
    for name, size, color in (("Heading 1", 18, "0F766E"), ("Heading 2", 14, "172033"), ("Heading 3", 12, "344054")):
        styles[name].font.name = cjk_font
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor.from_string(color)
    for name in ("Body Text", "Subtitle", "Caption", "List Bullet", "List Number"):
        styles[name].font.name = cjk_font
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(spec["title"])
    subtitle = str(spec.get("subtitle") or "").strip()
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.style = "Subtitle"
    meta = document.add_paragraph()
    meta.add_run(f"面向：{spec['audience']}\n").bold = True
    meta.add_run(f"目的：{spec['purpose']}")
    document.add_paragraph()

    header = section.header.paragraphs[0]
    header.text = spec["title"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(102, 112, 133)
    footer = section.footer.paragraphs[0]
    footer.text = f"{spec.get('author') or 'Nomi'} · 可追溯文档"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(102, 112, 133)


def _configure_east_asian_language(document: Document) -> None:
    default_run_properties = document.styles.element.xpath(
        ".//w:docDefaults/w:rPrDefault/w:rPr"
    )[0]
    languages = default_run_properties.find(qn("w:lang"))
    if languages is None:
        languages = OxmlElement("w:lang")
        default_run_properties.append(languages)
    languages.set(qn("w:val"), "zh-CN")
    languages.set(qn("w:eastAsia"), "zh-CN")

    settings = document.settings.element
    theme_languages = settings.find(qn("w:themeFontLang"))
    if theme_languages is None:
        theme_languages = OxmlElement("w:themeFontLang")
        settings.append(theme_languages)
    theme_languages.set(qn("w:val"), "zh-CN")
    theme_languages.set(qn("w:eastAsia"), "zh-CN")


def _apply_direct_cjk_fonts(document: Document, font_name: str) -> None:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            if not run.text:
                continue
            run.font.name = font_name
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:ascii"), font_name)
            fonts.set(qn("w:hAnsi"), font_name)
            fonts.set(qn("w:eastAsia"), font_name)
